"""Shared context for the RFS workbook builder (book.py): styles, period model, cell writers and
formula helpers. The layout follows the engagement's final RFS file (KMC format):
MasterData -> TB sheets (Masters in G) -> Notes (SUMIF on Masters + restatement entries) -> BS_1 / PL_2.
"""
import datetime, re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.worksheet.properties import PageSetupProperties

TN = "Times New Roman"
fN = Font(name=TN, size=10)
fB = Font(name=TN, size=10, bold=True)
fI = Font(name=TN, size=10, italic=True)
fBI = Font(name=TN, size=10, bold=True, italic=True)
fChk = Font(name=TN, size=8, color="808080")           # reviewer check cells
fChkI = Font(name=TN, size=8, italic=True, color="808080")
fGrey = Font(name=TN, size=10, color="808080")
fGreyI = Font(name=TN, size=10, italic=True, color="808080")
fGreyB = Font(name=TN, size=10, bold=True, color="808080")
fBlue = Font(name=TN, size=10, color="0000FF")
fRed = Font(name=TN, size=10, color="FF0000")
fRedI = Font(name=TN, size=10, italic=True, color="FF0000")
fSmallI = Font(name=TN, size=9, italic=True)
fTB = Font(name="Arial", size=9)
fTBB = Font(name="Arial", size=9, bold=True)
fTBI = Font(name="Arial", size=8, italic=True, color="808080")

ACC = '_(* #,##0.00_);_(* \\(#,##0.00\\);_(* \\-??_);_(@_)'
ACC0 = '_(* #,##0_);_(* \\(#,##0\\);_(* \\-??_);_(@_)'
PCT = '0.00%;\\(0.00%\\);\\-'
PCT3 = '0.000%;\\(0.000%\\);\\-'
EPS = '#,##0.00;\\(#,##0.00\\)'
NUM0 = '#,##0;\\(#,##0\\)'
RATIO = '0.00;\\(0.00\\);\\-'
DMY = 'dd\\-mm\\-yyyy'

thin = Side(style="thin")
dbl = Side(style="double")
B_TB = Border(top=thin, bottom=thin)
B_T = Border(top=thin)
B_B = Border(bottom=thin)
B_TB2 = Border(top=thin, bottom=dbl)
B_DBL = Border(top=dbl, bottom=dbl)
B_ALL = Border(top=thin, bottom=thin, left=thin, right=thin)
B_LR = Border(left=thin, right=thin)
B_BLR = Border(bottom=thin, left=thin, right=thin)
A_C = Alignment(horizontal="center")
A_CW = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_L = Alignment(horizontal="left")
A_LW = Alignment(horizontal="left", vertical="top", wrap_text=True)
A_J = Alignment(horizontal="justify", vertical="top", wrap_text=True)
A_W = Alignment(wrap_text=True, vertical="top")
Y_IN = PatternFill("solid", fgColor="FFF2CC")         # input
Y_URG = PatternFill("solid", fgColor="FFFF00")        # urgent input

MD = "MasterData"
ADJ = "Adj Entry RFS"


def q(s):
    return "'" + s.replace("'", "''") + "'" if re.search(r"[^A-Za-z0-9_]", s) else s


def qs(s):
    """string literal for a formula"""
    return '"' + str(s).replace('"', '""') + '"'


def fy_label(d):
    y = d.year if d.month > 3 else d.year - 1
    return f"{y}-{str(y + 1)[2:]}"


def ddate(d):
    if isinstance(d, datetime.datetime):
        return d.date()
    if isinstance(d, str):
        return datetime.date.fromisoformat(d[:10])
    return d


def long_date(d):
    return f"{d:%B} {d.day:02d}, {d.year}"


class Ctx:
    """Period model and helpers shared by all sheets."""

    def __init__(self, I):
        self.I = I
        D = I["dates"]
        self.stub = ddate(D["stub"]) if D.get("stub") else None
        fy = [ddate(x) for x in D["fy"]]                 # latest, previous, earliest year ends
        self.opening = ddate(D.get("opening") or fy[2].replace(year=fy[2].year - 1))
        self.PK = (["S"] if self.stub else []) + ["Y1", "Y2", "Y3"]
        self.N = len(self.PK)
        self.DATE = {"S": self.stub, "Y1": fy[0], "Y2": fy[1], "Y3": fy[2], "OPEN": self.opening}
        self.MDC = {"S": f"{MD}!$B$9", "Y1": f"{MD}!$B$11", "Y2": f"{MD}!$B$12", "Y3": f"{MD}!$B$13", "OPEN": f"{MD}!$B$14"}
        self.PREV = {"S": "Y1", "Y1": "Y2", "Y2": "Y3", "Y3": "OPEN"}
        self.SCALE = f"{MD}!$B$15"
        self.TAX = f"{MD}!$B$17"
        self.TBN = {p: "TB " + fy_label(self.DATE[p]) for p in self.PK}
        self.EARLY = "Y3"
        # columns: notes D.. (+ opening), BS_1 D.. gap opening, CFS C.., SOA E..
        self.NC = {p: CL(4 + i) for i, p in enumerate(self.PK)}
        self.NC["OPEN"] = CL(4 + self.N)
        self.BC = dict(self.NC)
        self.BC["OPEN"] = CL(5 + self.N)
        self.BGAP = CL(4 + self.N)
        self.CC = {p: CL(3 + i) for i, p in enumerate(self.PK)}
        self.SC = {p: CL(5 + i) for i, p in enumerate(self.PK)}
        self.SC["OPEN"] = CL(5 + self.N)
        self.LASTN = self.NC["OPEN"]                     # last column of a note table (opening)
        self.LASTP = self.NC[self.PK[-1]]                # last period column
        self.ADJ_END = 150
        self.refs = {}                                   # name -> row (resolved at the end: {{name}})
        self.noteno = {}                                 # note key -> number (texts may say {NOTE:key})
        self.noteno_counter = 2                          # Annexure IV carries notes 1 and 2
        self.adj_masters = set()
        self.adj_pl_masters = set()
        self.tbval = {}                                  # (pk, master) -> lakhs (python-side, for nil-line removal)
        self.custom = I.get("custom_masters", [])
        self.texts = dict(I.get("texts") or {})
        self.text_used = {}
        self.wb = Workbook()
        self.wb.remove(self.wb.active)

    # ---------------------------------------------------------------- periods
    def is_stub(self, p):
        return p == "S"

    def dtxt(self, p, kind="bs", nl=True):
        sep = "\n" if nl else " "
        if kind == "bs":
            pre = "As at " + sep
        elif kind == "pl":
            pre = ("For the period ended" if p == "S" else "For the year ended") + sep
        else:
            pre = kind
        return f'="{pre}"&TEXT({self.MDC[p]},"mmmm dd, yyyy")'

    def per_word(self, p):
        return "period" if p == "S" else "year"

    def dates_text(self, keys=None):
        keys = keys or self.PK
        ds = [long_date(self.DATE[k]) for k in keys]
        return ", ".join(ds[:-1]) + " and " + ds[-1] if len(ds) > 1 else ds[0]

    # ---------------------------------------------------------------- formula helpers
    def tb(self, p, ref, col="I", mcol="G"):
        if p == "OPEN":
            s = q(self.TBN[self.EARLY])
            return f"SUMIF({s}!$J:$J,{ref},{s}!$L:$L)"
        s = q(self.TBN[p])
        return f"SUMIF({s}!${mcol}:${mcol},{ref},{s}!${col}:${col})"

    def tbl(self, p, ledger, col="I"):
        """SUMIF on the ledger name (column A)"""
        if p == "OPEN":
            s = q(self.TBN[self.EARLY])
            return f"SUMIF({s}!$A:$A,{qs(ledger)},{s}!$L:$L)"
        s = q(self.TBN[p])
        return f"SUMIF({s}!$A:$A,{qs(ledger)},{s}!${col}:${col})"

    def _rng(self, c):
        return f"{q(ADJ)}!${c}$3:${c}${self.ADJ_END}"

    def adj_bs(self, ref, p, nat="Dr", col="C"):
        """cumulative restatement entries up to the date of period p (Dr - Cr for Dr nature)"""
        d = self.MDC[p]
        a = f"SUMIFS({self._rng('E')},{self._rng(col)},{ref},{self._rng('A')},\"<=\"&{d})"
        b = f"SUMIFS({self._rng('F')},{self._rng(col)},{ref},{self._rng('A')},\"<=\"&{d})"
        return f"({a}-{b})" if nat == "Dr" else f"({b}-{a})"

    def adj_pl(self, ref, p, nat="Dr", col="C"):
        s, e = self.MDC[self.PREV[p]], self.MDC[p]
        a = f"SUMIFS({self._rng('E')},{self._rng(col)},{ref},{self._rng('A')},\">\"&{s},{self._rng('A')},\"<=\"&{e})"
        b = f"SUMIFS({self._rng('F')},{self._rng(col)},{ref},{self._rng('A')},\">\"&{s},{self._rng('A')},\"<=\"&{e})"
        return f"({a}-{b})" if nat == "Dr" else f"({b}-{a})"

    def master_val(self, p, master, ref=None, grp="BS", nat="Dr"):
        """TB balance of a master (+ restatement entries if the master is used in 'Adj Entry RFS')"""
        ref = ref or qs(master)
        f = self.tb(p, ref)
        if master in self.adj_masters:
            f += "+" + (self.adj_pl(ref, p, nat) if grp == "P&L" else self.adj_bs(ref, p, nat))
        return f

    def nonzero(self, master, keys=None):
        keys = keys or (self.PK + ["OPEN"])
        if master in self.adj_masters:
            return True
        return any(abs(self.tbval.get((k, master), 0)) > 0.000001 for k in keys)

    # ---------------------------------------------------------------- refs
    def R(self, name):
        return "{{" + name + "}}"

    def setref(self, name, row):
        self.refs[name] = row

    def resolve(self):
        pat = re.compile(r"\{\{([^}]+)\}\}")
        npat = re.compile(r"\{NOTE:([A-Za-z_]+)\}")
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if isinstance(v, str) and "{NOTE:" in v:
                        v = npat.sub(lambda m: str(self.noteno.get(m.group(1), "[ ]")), v)
                        c.value = v
                    if isinstance(v, str) and "{{" in v:
                        def rep(m):
                            k = m.group(1)
                            if k not in self.refs:
                                raise KeyError(f"unresolved reference {k} in {ws.title}!{c.coordinate}")
                            return str(self.refs[k])
                        c.value = pat.sub(rep, v)

    # ---------------------------------------------------------------- Annexure V groups
    PL_GROUPS = [("dep", "Depreciation and loss on write off of property, plant and equipment", ["Depreciation and amortisation", "Miscellaneous expenses"]),
                 ("grat", "Provision for gratuity", ["Gratuity"]),
                 ("msme", "Interest on delayed payments to micro and small enterprises", ["Interest on delayed payments to MSME"]),
                 ("curtax", "Current tax", ["Current tax"]), ("deftax", "Deferred tax", ["Deferred tax"])]
    BS_GROUPS = [("ppe", "Property, plant and equipment (depreciation as per Fixed Asset Register)", ["Property, Plant and Equipment"]),
                 ("grat", "Provision for gratuity", ["Provision for gratuity"]),
                 ("msme", "Interest accrued on dues to micro and small enterprises", ["Interest accrued on dues to MSME"]),
                 ("taxprov", "Provision for income tax", ["Provision for income tax (net of advance tax)"]),
                 ("dta", "Deferred tax asset (net)", ["Deferred tax asset (net)"])]

    def _groups(self, base, pool):
        out, seen = [], set()
        for key, label, ms in base:
            use = [m for m in ms if m in pool]
            seen |= set(ms)
            if use:
                out.append((label, use, key))
        for m in sorted(pool - seen):
            out.append((m, [m], re.sub(r"[^a-z0-9]+", "_", m.lower()).strip("_")))
        return out

    def soa_pl_groups(self):
        return self._groups(self.PL_GROUPS, set(self.adj_pl_masters))

    def soa_bs_groups(self):
        return self._groups(self.BS_GROUPS, set(m for m in self.adj_masters if m not in self.adj_pl_masters and m != "Reserves and surplus"))

    # ---------------------------------------------------------------- texts
    def text(self, key, default, label=None):
        """editable paragraph: the user's text if given, else the default (recorded so the app can show every default)"""
        self.text_used[key] = {"default": default, "label": label or key}
        v = self.texts.get(key)
        return default if v in (None, "") else v

    def text_list(self, key, default, label=None):
        """list of paragraphs, edited as one text (one paragraph per line)"""
        t = self.text(key, "\n".join(x for x in default if x), label)
        return [x for x in str(t).split("\n") if x.strip()]

    # ---------------------------------------------------------------- sheets / writers
    def sheet(self, name, widths, print_cols=None, portrait=True):
        ws = self.wb.create_sheet(name)
        for k, v in widths.items():
            ws.column_dimensions[k].width = v
        ws.sheet_view.showGridLines = False
        ws.page_setup.paperSize = 9
        ws.page_setup.orientation = "portrait" if portrait else "landscape"
        if print_cols:
            ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
        ws._rfs_pc = print_cols
        return ws

    def finish_print(self, ws, last_row):
        if getattr(ws, "_rfs_pc", None):
            ws.print_area = f"A1:{ws._rfs_pc}{last_row}"


def w(ws, ref, v, font=fN, fmt=None, al=None, bd=None, fill=None):
    c = ws[ref]
    c.value = v
    c.font = font
    if fmt:
        c.number_format = fmt
    if al:
        c.alignment = al
    if bd:
        c.border = bd
    if fill:
        c.fill = fill
    return c


def merge(ws, rng, v=None, font=fN, al=None, height=None):
    a = rng.split(":")[0]
    if v is not None:
        w(ws, a, v, font, al=al)
    ws.merge_cells(rng)
    if height:
        ws.row_dimensions[int(re.sub(r"[A-Z]", "", a))].height = height


def est_height(text, width_chars, line=13.5, minimum=15.0):
    """row height for a wrapped paragraph in a merged range of total width_chars"""
    if not text:
        return minimum
    n = 0
    for part in str(text).split("\n"):
        n += max(1, -(-len(part) // max(1, int(width_chars * 1.05))))
    return max(minimum, round(n * line + 1.5, 2))
