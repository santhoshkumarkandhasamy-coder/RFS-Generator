"""Read text and tables from any input document: Excel (.xlsx/.xlsm/.xls), CSV, Word (.docx/.doc),
PowerPoint (.pptx/.ppt) and PDF. Legacy binary formats are converted with LibreOffice
(or Microsoft Office on Windows) first.
"""
import csv
import os
import shutil
import subprocess
import sys

SUPPORTED = ["xlsx", "xlsm", "xls", "csv", "docx", "doc", "pptx", "ppt", "pdf", "odt", "ods", "odp"]
LEGACY = {"xls": "xlsx", "ods": "xlsx", "doc": "docx", "odt": "docx", "ppt": "pptx", "odp": "pptx"}


def _soffice():
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(p):
            return p
    return None


def _office_convert(path, target, out_dir):
    """Windows fallback using Microsoft Office COM."""
    import win32com.client
    src = os.path.abspath(path)
    out = os.path.abspath(os.path.join(out_dir, os.path.splitext(os.path.basename(path))[0] + "." + target))
    if target == "xlsx":
        app = win32com.client.DispatchEx("Excel.Application"); app.DisplayAlerts = False
        try:
            wb = app.Workbooks.Open(src); wb.SaveAs(out, 51); wb.Close(False)
        finally:
            app.Quit()
    elif target == "docx":
        app = win32com.client.DispatchEx("Word.Application")
        try:
            d = app.Documents.Open(src); d.SaveAs(out, 16); d.Close(False)
        finally:
            app.Quit()
    else:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        try:
            p = app.Presentations.Open(src, WithWindow=False); p.SaveAs(out, 24); p.Close()
        finally:
            app.Quit()
    return out


def to_modern(path, out_dir):
    """Convert .xls/.doc/.ppt/.od* to the matching Office Open XML format. Returns new path (or the original)."""
    ext = os.path.splitext(path)[1].lower().lstrip(".")
    if ext not in LEGACY:
        return path
    target = LEGACY[ext]
    os.makedirs(out_dir, exist_ok=True)
    so = _soffice()
    if so:
        subprocess.run([so, "--headless", "--convert-to", target, "--outdir", out_dir, path], check=True, timeout=300, capture_output=True)
        return os.path.join(out_dir, os.path.splitext(os.path.basename(path))[0] + "." + target)
    if sys.platform.startswith("win"):
        return _office_convert(path, target, out_dir)
    raise RuntimeError(f".{ext} files need LibreOffice (or Microsoft Office on Windows) to be read. Save the file as .{target} and upload again.")


def _cell(v):
    if v is None:
        return ""
    return str(v).strip()


def extract_tables(path, work_dir=None):
    """Returns a list of tables; each table is a list of rows (list of str)."""
    work_dir = work_dir or os.path.dirname(path)
    path = to_modern(path, work_dir)
    ext = os.path.splitext(path)[1].lower()
    tables = []
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            rows = [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(r)]
            if rows:
                tables.append(rows)
    elif ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
            tables.append([r for r in csv.reader(fh) if any(c.strip() for c in r)])
    elif ext == ".docx":
        from docx import Document
        d = Document(path)
        for t in d.tables:
            rows = []
            for r in t.rows:
                cells, prev = [], None
                for c in r.cells:            # merged cells repeat - keep once
                    if c._tc is prev:
                        continue
                    prev = c._tc
                    cells.append(c.text.strip())
                rows.append(cells)
            tables.append(rows)
    elif ext == ".pptx":
        from pptx import Presentation
        for slide in Presentation(path).slides:
            for sh in slide.shapes:
                if sh.has_table:
                    tables.append([[c.text.strip() for c in r.cells] for r in sh.table.rows])
    elif ext == ".pdf":
        t = pdf_ledger_table(path)
        if t:
            return [t]
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            merged = []
            for page in pdf.pages:
                found = page.extract_tables()
                if found:
                    for t in found:
                        merged.extend([[_cell(c) for c in r] for r in t if r and any(r)])
                else:   # text-only page: split lines on runs of spaces
                    import re
                    for line in (page.extract_text() or "").splitlines():
                        parts = re.split(r"\s{2,}", line.strip())
                        if len(parts) > 1:
                            merged.append(parts)
            if merged:
                tables.append(merged)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
    return tables


def extract_text(path, work_dir=None, max_pages=None):
    work_dir = work_dir or os.path.dirname(path)
    path = to_modern(path, work_dir)
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        rd = PdfReader(path)
        pages = rd.pages if not max_pages else rd.pages[:max_pages]
        return "\n".join((p.extract_text() or "") for p in pages)
    if ext == ".docx":
        from docx import Document
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text.strip() for c in r.cells))
        return "\n".join(parts)
    if ext == ".pptx":
        from pptx import Presentation
        parts = []
        for i, slide in enumerate(Presentation(path).slides, 1):
            parts.append(f"--- Slide {i} ---")
            for sh in slide.shapes:
                if sh.has_text_frame:
                    parts.append(sh.text_frame.text)
                if sh.has_table:
                    for r in sh.table.rows:
                        parts.append(" | ".join(c.text for c in r.cells))
        return "\n".join(parts)
    return "\n".join(" | ".join(r) for t in extract_tables(path, work_dir) for r in t)


def page_texts(pdf_path):
    from pypdf import PdfReader
    return [(p.extract_text() or "") for p in PdfReader(pdf_path).pages]


_AMT = __import__("re").compile(r"^\(?-?[\d,]+(\.\d+)?\)?$")


def pdf_ledger_table(path):
    """Statement-style PDFs (trial balance printed from Tally / Busy / Excel): place every amount under the
    nearest column heading. Returns rows [name, opening, debit, credit, closing] or None."""
    import pdfplumber
    heads = ("opening", "debit", "credit", "closing")
    anchors, out = None, [["Particulars", "Opening Balance", "Debit", "Credit", "Closing Balance"]]
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            words = sorted(page.extract_words(keep_blank_chars=False, use_text_flow=False), key=lambda z: (z["top"], z["x0"]))
            bands = []                               # rows = words whose vertical extents overlap (handles wrapped names)
            for w_ in words:
                if bands and w_["top"] < bands[-1]["bottom"] - 1:
                    bands[-1]["words"].append(w_); bands[-1]["bottom"] = max(bands[-1]["bottom"], w_["bottom"])
                else:
                    bands.append({"bottom": w_["bottom"], "words": [w_]})
            for b_ in bands:
                b_["top"] = min(z["top"] for z in b_["words"])
            gaps = sorted(bands[k + 1]["top"] - bands[k]["bottom"] for k in range(len(bands) - 1))
            med = gaps[len(gaps) // 2] if gaps else 0
            def _has_amt(b_):
                return any(_AMT.match(z["text"]) and ("," in z["text"] or "." in z["text"]) for z in b_["words"])
            merged = []
            for b_ in bands:                          # glue wrapped name lines to the row that carries the amounts
                if merged and med > 0 and (b_["top"] - merged[-1]["bottom"]) < 0.6 * med and (not _has_amt(b_) or not _has_amt(merged[-1])) \
                        and not (_has_amt(b_) and _has_amt(merged[-1])):
                    merged[-1]["words"] += b_["words"]; merged[-1]["bottom"] = max(merged[-1]["bottom"], b_["bottom"])
                else:
                    merged.append(b_)
            for band in merged:
                ws_ = band["words"]
                if anchors:
                    xs = sorted(anchors.values())
                    gap = (xs[1] - xs[0]) if len(xs) > 1 else 80
                    zone = xs[0] - 0.75 * gap          # left edge of the amount columns
                else:
                    zone = 0
                near = (xs[0] - 0.3 * gap) if anchors else 0
                isamt = lambda z: bool(_AMT.match(z["text"]) and any(ch.isdigit() for ch in z["text"])
                                       and ((("," in z["text"] or "." in z["text"]) and z["x1"] > zone) or z["x0"] > near))
                amt_x = [z["x0"] for z in ws_ if isamt(z)]
                cut = min(amt_x) if amt_x else 1e9
                names = sorted([z for z in ws_ if z["x0"] < cut], key=lambda z: (round(z["top"]), z["x0"]))
                rest = sorted([z for z in ws_ if z["x0"] >= cut], key=lambda z: z["x0"])
                ln = names + rest
                low = [z["text"].lower() for z in ln]
                found = {h: z for h in heads for z in ln if z["text"].lower().startswith(h)}
                if len(found) >= 3:                       # header line
                    anchors = {h: (z["x0"] + z["x1"]) / 2 for h, z in found.items()}
                    continue
                if not anchors or "particulars" in low:
                    continue
                cells = {h: "" for h in heads}
                name_parts, i = [], 0
                first_amt_x = None
                while i < len(ln):
                    t = ln[i]["text"]
                    if isamt(ln[i]):
                        side = ""
                        if i + 1 < len(ln) and ln[i + 1]["text"].rstrip(".").lower() in ("dr", "cr"):
                            side = " " + ln[i + 1]["text"].rstrip(".").capitalize(); i += 1
                        cx = (ln[i]["x1"] + ln[i - (1 if side else 0)]["x0"]) / 2 if side else (ln[i]["x0"] + ln[i]["x1"]) / 2
                        h = min(anchors, key=lambda k: abs(anchors[k] - cx))
                        cells[h] = t + side
                        first_amt_x = first_amt_x if first_amt_x is not None else ln[i]["x0"]
                    elif first_amt_x is None:
                        name_parts.append(t)
                    i += 1
                name = " ".join(name_parts).strip()
                if name and any(cells.values()):
                    out.append([name, cells["opening"], cells["debit"], cells["credit"], cells["closing"]])
                elif name and not any(cells.values()) and name.lower() not in ("grand total",):
                    out.append([name, "", "", "", ""])
    return out if len(out) > 5 else None
