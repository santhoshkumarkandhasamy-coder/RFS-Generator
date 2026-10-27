"""Export the RFS workbook to Word, PowerPoint and HTML.

The workbook is formula-driven, so it is first recalculated (LibreOffice, or Microsoft Excel on
Windows) and the computed values are then rendered annexure by annexure.
"""
import html
import os
import shutil
import subprocess
import sys

import openpyxl

# sheets that make up the RFS document, in order (TB / source / working sheets are left out)
REPORT_SHEETS = ["BS_1", "PL_2", "CFS_3", "SAP_4", "SOA_5", "Notes_Equity & Reserves", "Notes_Liabilities", "Notes_PPE",
                 "Notes_Assets", "Notes_P&L", "RPT", "Gratuity and contingencies", "Tax Computation", "Accounting ratios",
                 "Analytical Ratios", "Capitalisation", "Additional Disclosures"]


# ------------------------------------------------------------------ recalculation
def _soffice():
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(p):
            return p
    return None


def recalc_engine():
    if _soffice():
        return "LibreOffice"
    if sys.platform.startswith("win"):
        try:
            import win32com.client  # noqa: F401
            return "Microsoft Excel"
        except ImportError:
            pass
    return None


def recalculate(xlsx_path, out_dir):
    """Returns path of a copy of the workbook with calculated values stored."""
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, os.path.basename(xlsx_path))
    so = _soffice()
    if so:
        subprocess.run([so, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out_dir, xlsx_path],
                       check=True, timeout=600, capture_output=True)
        return out
    if sys.platform.startswith("win"):
        import win32com.client
        xl = win32com.client.DispatchEx("Excel.Application")
        xl.Visible = False; xl.DisplayAlerts = False
        try:
            wbk = xl.Workbooks.Open(os.path.abspath(xlsx_path))
            xl.CalculateFull()
            wbk.SaveAs(os.path.abspath(out), 51)
            wbk.Close(False)
        finally:
            xl.Quit()
        return out
    raise RuntimeError("Install LibreOffice (or use Windows with Microsoft Excel and pywin32) to export Word / PowerPoint / HTML.")


def docx_to_pdf(docx_path, out_dir):
    so = _soffice()
    if so:
        subprocess.run([so, "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path], check=True, timeout=600, capture_output=True)
        return os.path.join(out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
    if sys.platform.startswith("win"):
        import win32com.client
        wd = win32com.client.DispatchEx("Word.Application")
        try:
            d = wd.Documents.Open(os.path.abspath(docx_path))
            out = os.path.abspath(os.path.join(out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf"))
            d.SaveAs(out, 17); d.Close(False)
        finally:
            wd.Quit()
        return out
    raise RuntimeError("Reading a Word offer document needs LibreOffice or Microsoft Word - upload the PDF instead.")


def add_reference_sheet(xlsx_path, regcheck, docs, tb_sources):
    """Adds 'Reference & Documents': ICDR check, DRHP / RHP used as format reference, documents considered."""
    from openpyxl.styles import Font
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.create_sheet("Reference & Documents")
    b, n, i = Font(name="Times New Roman", size=10, bold=True), Font(name="Times New Roman", size=10), Font(name="Times New Roman", size=9, italic=True)
    ws.column_dimensions["A"].width = 14; ws.column_dimensions["B"].width = 90; ws.column_dimensions["C"].width = 22; ws.column_dimensions["D"].width = 50
    r = 1
    def put(row, vals, font=n):
        for j, v in enumerate(vals):
            c = ws.cell(row, 1 + j, v); c.font = font
    put(r, ["Regulatory reference and documents considered in preparing the Restated Financial Statements"], b); r += 2
    put(r, ["1. SEBI (ICDR) Regulations - latest updates checked"], b); r += 1
    ic = (regcheck or {}).get("icdr")
    if ic:
        put(r, [f"Checked on {ic['checked']}"], i); r += 1
        put(r, ["Date", "Title", "Source", "Link"], b); r += 1
        for it in ic["items"]:
            put(r, [str(it["date"] or ""), it["title"], it.get("source", ""), it["url"]]); r += 1
    else:
        put(r, ["", "Not checked in the app - confirm the latest ICDR amendments before issue."], i); r += 1
    r += 1
    put(r, ["2. Offer document used as format reference"], b); r += 1
    rf = (regcheck or {}).get("reference")
    if rf:
        put(r, ["", f"{rf['name']} ({rf['type']}, {rf.get('date', '')})", "", rf.get("url", "")]); r += 1
        for c in rf["analysis"]["coverage"]:
            put(r, ["", c["Item produced by RFS Generator"], c["Seen in DRHP / RHP"]]); r += 1
    else:
        put(r, ["", "No DRHP / RHP compared in the app."], i); r += 1
    r += 1
    put(r, ["3. Trial balances"], b); r += 1
    for k, v in (tb_sources or {}).items():
        put(r, [k, v]); r += 1
    r += 1
    put(r, ["4. Supporting documents"], b); r += 1
    put(r, ["Type", "Document", "Category", "Added"], b); r += 1
    for k, v in (docs or {}).items():
        put(r, [v.get("type", ""), k, v.get("category", ""), v.get("added", "")]); r += 1
    wb.save(xlsx_path)
    return xlsx_path


def docx_to_doc(docx_path, out_dir):
    so = _soffice()
    if so:
        subprocess.run([so, "--headless", "--convert-to", "doc", "--outdir", out_dir, docx_path], check=True, timeout=600, capture_output=True)
        return os.path.join(out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".doc")
    if sys.platform.startswith("win"):
        import win32com.client
        wd = win32com.client.DispatchEx("Word.Application")
        try:
            d = wd.Documents.Open(os.path.abspath(docx_path))
            out = os.path.abspath(os.path.join(out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".doc"))
            d.SaveAs(out, 0); d.Close(False)
        finally:
            wd.Quit()
        return out
    raise RuntimeError("Converting to .doc needs LibreOffice or Microsoft Word.")


# ------------------------------------------------------------------ read calculated sheets
def _fmt(v, number_format):
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        nf = number_format or ""
        if nf == "General" and float(v).is_integer():
            return f"{int(v)}"
        if "%" in nf:
            return f"{v * 100:,.2f}%" if abs(v) > 1e-12 else "-"
        if abs(v) < 0.005:
            return "-"
        if nf.startswith("0.00") and "x" not in nf and "#" not in nf:
            return f"{v:,.2f}"
        if "#,##0_" in nf or nf in ("#,##0", "0"):
            return f"{v:,.0f}"
        return f"{v:,.2f}"
    if hasattr(v, "strftime"):
        return v.strftime("%d-%m-%Y")
    return str(v).strip()


def read_tables(calc_path):
    """{sheet: {"title": [...header lines], "rows": [[(text, bold, is_number), ...], ...]}} from the calculated workbook."""
    wb = openpyxl.load_workbook(calc_path, data_only=True)
    out = {}
    for name in REPORT_SHEETS:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        cols = []
        for c in range(2, min(ws.max_column, 13) + 1):  # B..M (a stub period adds a column)
            letter = openpyxl.utils.get_column_letter(c)
            if ws.column_dimensions[letter].hidden:
                continue
            cols.append(c)
        title = [str(ws.cell(r, 2).value).strip() for r in range(2, 7) if ws.cell(r, 2).value not in (None, "", 0)]
        rows = []
        for r in range(7, ws.max_row + 1):
            cells, any_val = [], False
            for c in cols:
                cell = ws.cell(r, c)
                f = cell.font
                if f is not None and f.sz is not None and f.sz <= 8:     # reviewer check cells
                    cells.append(("", False, False)); continue
                grey = f is not None and f.color is not None and isinstance(f.color.rgb, str) and f.color.rgb.upper().endswith("808080")
                if grey and c > 3:                                       # working columns (e.g. opening balance sheet)
                    cells.append(("", False, False)); continue
                t = _fmt(cell.value, cell.number_format)
                if t:
                    any_val = True
                hdr = isinstance(cell.value, str) and c > 3 and cell.value.strip().startswith(("As at", "For the year", "No. of shares", "Pre-issue", "Post-issue"))
                cells.append((t, bool(f.b) if f else False, isinstance(cell.value, (int, float)) or hdr))
            if any_val:
                rows.append(cells)
        # drop columns that are empty throughout
        keep = [i for i in range(len(cols)) if any(rw[i][0] for rw in rows)]
        rows = [[rw[i] for i in keep] for rw in rows]
        out[name] = {"title": title, "rows": rows}
    return out


# ------------------------------------------------------------------ HTML
def to_html(tables, path, company):
    css = """body{font-family:'Times New Roman',serif;font-size:13px;margin:24px;color:#111}
h1{font-size:18px;margin:0}h2{font-size:15px;margin:28px 0 4px;border-bottom:1px solid #999;padding-bottom:3px}
.t{font-size:12px;color:#444;margin:2px 0}nav{position:sticky;top:0;background:#fff;padding:6px 0;border-bottom:1px solid #ddd;margin-bottom:8px}
nav a{margin-right:10px;font-size:12px;color:#1a4fa0;text-decoration:none}table{border-collapse:collapse;width:100%;max-width:1100px}
td{padding:3px 6px;vertical-align:top;border-bottom:1px solid #eee}td.n{text-align:right;white-space:nowrap}tr.b td{font-weight:bold}
@media print{nav{display:none}h2{page-break-before:always}}"""
    parts = [f"<!doctype html><html><head><meta charset='utf-8'><meta name='viewport' content='width=device-width, initial-scale=1'>"
             f"<title>{html.escape(company)} - RFS</title><style>{css}</style></head><body>",
             f"<h1>{html.escape(company)}</h1><div class='t'>Restated Financial Statements</div><nav>"]
    for i, name in enumerate(tables):
        parts.append(f"<a href='#s{i}'>{html.escape(name)}</a>")
    parts.append("</nav>")
    for i, (name, t) in enumerate(tables.items()):
        heading = t["title"][3] if len(t["title"]) > 3 else name
        parts.append(f"<h2 id='s{i}'>{html.escape(heading)}</h2>")
        for line in t["title"][4:]:
            parts.append(f"<div class='t'>{html.escape(line)}</div>")
        parts.append("<table>")
        for rw in t["rows"]:
            bold = any(b for _, b, _ in rw[:1]) or (rw and rw[0][1])
            parts.append("<tr class='b'>" if bold else "<tr>")
            for txt, b, num in rw:
                parts.append(f"<td class='n'>{html.escape(txt).replace(chr(10), '<br>')}</td>" if num else f"<td>{html.escape(txt).replace(chr(10), '<br>')}</td>")
            parts.append("</tr>")
        parts.append("</table>")
    parts.append("</body></html>")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("".join(parts))
    return path


# ------------------------------------------------------------------ Word
def to_docx(tables, path, company):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Cm

    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(9)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.5); sec.top_margin = sec.bottom_margin = Cm(1.5)
    p = doc.add_paragraph(); r = p.add_run(company); r.bold = True; r.font.size = Pt(16); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Restated Financial Statements"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first = True
    for name, t in tables.items():
        ncols = max((len(rw) for rw in t["rows"]), default=1)
        if not first:
            new = doc.add_section()
            wide = ncols > 5
            new.orientation = WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT
            if (wide and new.page_width < new.page_height) or (not wide and new.page_width > new.page_height):
                new.page_width, new.page_height = new.page_height, new.page_width
        first = False
        for k, line in enumerate(t["title"]):
            p = doc.add_paragraph(); run = p.add_run(line); run.bold = k < 4; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
        if not t["rows"]:
            continue
        table = doc.add_table(rows=0, cols=ncols)
        table.style = "Table Grid"
        table.autofit = False
        usable = (doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin)
        maxlen = [max((len(rw[j][0]) for rw in t["rows"] if j < len(rw) and not rw[j][2]), default=0) for j in range(ncols)]
        lab = max(range(ncols), key=lambda j: maxlen[j])
        widths = []
        for j in range(ncols):
            if j == lab:
                widths.append(None)
            elif j < lab and maxlen[j] <= 4:
                widths.append(int(usable * 0.05))          # note number column
            else:
                widths.append(int(usable * min(0.16, 0.54 / max(ncols - 1, 1))))
        widths[lab] = usable - sum(x for x in widths if x)
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        for j, gc in enumerate(table._tbl.tblGrid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths[j] / 635)))
        tblW = table._tbl.tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW"); table._tbl.tblPr.append(tblW)
        tblW.set(qn("w:type"), "dxa"); tblW.set(qn("w:w"), str(int(usable / 635)))
        for rw in t["rows"]:
            cells = table.add_row().cells
            for j in range(ncols):
                cells[j].width = widths[j]
            filled = [j for j, c in enumerate(rw) if c[0]]
            if len(filled) == 1 and len(rw[filled[0]][0]) > 70 and not rw[filled[0]][2]:   # long text: span the row
                j0 = filled[0]
                m = cells[j0].merge(cells[ncols - 1]) if j0 < ncols - 1 else cells[j0]
                run = m.paragraphs[0].add_run(rw[j0][0]); run.font.size = Pt(8.5); run.bold = rw[j0][1]
                continue
            for j, (txt, b, num) in enumerate(rw):
                para = cells[j].paragraphs[0]
                run = para.add_run(txt); run.bold = b; run.font.size = Pt(8.5)
                if num:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.save(path)
    return path


# ------------------------------------------------------------------ PowerPoint (summary deck)
def to_pptx(tables, path, company):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def title_box(slide, text, sub=None):
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8)).text_frame
        tb.text = text; tb.paragraphs[0].runs[0].font.size = Pt(24); tb.paragraphs[0].runs[0].font.bold = True
        if sub:
            p = tb.add_paragraph(); p.text = sub; p.runs[0].font.size = Pt(12); p.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.7), Inches(2)).text_frame
    tb.text = company; tb.paragraphs[0].runs[0].font.size = Pt(36); tb.paragraphs[0].runs[0].font.bold = True
    p = tb.add_paragraph(); p.text = "Restated Financial Statements - summary"; p.runs[0].font.size = Pt(20)
    p = tb.add_paragraph(); p.text = "(All amounts in Rs. lakhs, except as otherwise stated)"; p.runs[0].font.size = Pt(14)

    def table_slides(name, heading, max_rows=16):
        t = tables.get(name)
        if not t or not t["rows"]:
            return
        rows = t["rows"]
        for start in range(0, len(rows), max_rows):
            chunk = rows[start:start + max_rows]
            ncols = max(len(rw) for rw in chunk)
            sl = prs.slides.add_slide(blank)
            title_box(sl, heading + (" (contd.)" if start else ""), t["title"][4] if len(t["title"]) > 4 else None)
            shape = sl.shapes.add_table(len(chunk), ncols, Inches(0.5), Inches(1.3), Inches(12.3), Inches(0.3) * len(chunk))
            tbl = shape.table
            tbl.columns[0].width = Inches(12.3 - 1.7 * (ncols - 1)) if ncols > 1 else Inches(12.3)
            for j in range(1, ncols):
                tbl.columns[j].width = Inches(1.7)
            for i, rw in enumerate(chunk):
                for j in range(ncols):
                    txt, b, num = rw[j] if j < len(rw) else ("", False, False)
                    cell = tbl.cell(i, j); cell.text = txt
                    para = cell.text_frame.paragraphs[0]
                    if para.runs:
                        para.runs[0].font.size = Pt(10); para.runs[0].font.bold = b
                    if num:
                        from pptx.enum.text import PP_ALIGN
                        para.alignment = PP_ALIGN.RIGHT

    table_slides("BS_1", "Restated Statement of Assets and Liabilities")
    table_slides("PL_2", "Restated Statement of Profit and Loss")
    table_slides("CFS_3", "Restated Statement of Cash Flows")
    table_slides("SOA_5", "Statement of Adjustments to Audited Financial Statements")
    table_slides("Accounting ratios", "Accounting Ratios")
    table_slides("Analytical Ratios", "Analytical Ratios")
    table_slides("Capitalisation", "Capitalisation Statement")
    prs.save(path)
    return path


def export_all(xlsx_path, formats, out_dir, company):
    """formats: subset of {"xlsx", "docx", "doc", "pptx", "html"}. Returns {format: path}."""
    results = {}
    os.makedirs(out_dir, exist_ok=True)
    calc = xlsx_path
    if any(f in formats for f in ("docx", "doc", "pptx", "html")) or "xlsx" in formats:
        try:
            calc = recalculate(xlsx_path, os.path.join(out_dir, "calc"))
        except Exception:
            calc = None
    if "xlsx" in formats:
        results["xlsx"] = xlsx_path          # original workbook: exact formatting, recalculates when opened in Excel
    if calc:
        results["_calc"] = calc              # calculated copy, used for previews
    others = [f for f in formats if f != "xlsx"]
    if others:
        if not calc:
            raise RuntimeError("Word / PowerPoint / HTML export needs LibreOffice, or Microsoft Excel with pywin32 on Windows.")
        tables = read_tables(calc)
        base = os.path.join(out_dir, os.path.splitext(os.path.basename(xlsx_path))[0])
        if "html" in formats:
            results["html"] = to_html(tables, base + ".html", company)
        if "docx" in formats or "doc" in formats:
            d = to_docx(tables, base + ".docx", company)
            if "docx" in formats:
                results["docx"] = d
            if "doc" in formats:
                results["doc"] = docx_to_doc(d, out_dir)
        if "pptx" in formats:
            results["pptx"] = to_pptx(tables, base + ".pptx", company)
    return results
