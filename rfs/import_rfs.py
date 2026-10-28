"""Import the inputs of an existing RFS workbook in this app's format (MasterData, ageing, shareholder register, Adj Entry RFS,
notes and annexures) into the app's project. Rows are found by their labels, so earlier versions of the file also work."""
import datetime as dt
import re
import openpyxl

PER_COLS = {"S": 0, "Y1": 1, "Y2": 2, "Y3": 3}


def _rs(v):
    """'=87182/MasterData!$B$15' or '=ROUND(2592600/...,4)' -> 87182 (amount in Rs as typed); numbers as is"""
    if isinstance(v, (int, float)):
        return v
    m = re.match(r"^=\s*(?:ROUND\()?\(?([-0-9.+*/() ]+?)\)?/MasterData!\$B\$15", str(v or ""))
    if m:
        try:
            return eval(m.group(1), {"__builtins__": {}})
        except Exception:
            return None
    return None


def _iso(d):
    return d.date().isoformat() if isinstance(d, dt.datetime) else (d.isoformat() if isinstance(d, dt.date) else None)


def _find(ws, text, cols="BC", start=1):
    t = text.lower()
    for r in range(start, ws.max_row + 1):
        for c in cols:
            v = ws[f"{c}{r}"].value
            if isinstance(v, str) and v.strip().lower().startswith(t):
                return r
    return None


def import_workbook(path):
    """returns (project_update, notes_update, messages)"""
    wb = openpyxl.load_workbook(path)
    wv = openpyxl.load_workbook(path, data_only=True)
    pr, N, msg = {}, {}, []
    S = wb.sheetnames
    stub = False
    if "MasterData" in S:
        md, mv = wb["MasterData"], wv["MasterData"]
        lab = {str(md[f"A{r}"].value).strip().lower(): r for r in range(1, 60) if md[f"A{r}"].value}
        g = lambda key: next((mv[f"B{r}"].value for k, r in lab.items() if k.startswith(key)), None)
        pr.update({"cin": g("cin"), "name": g("company name"), "former": g("formerly"), "address": g("registered address"),
                   "face_value": g("face value"), "bonus_factor": g("bonus"), "tax_rate": mv["B17"].value})
        y1 = g("year ended (latest)")
        if isinstance(y1, dt.datetime):
            pr["latest"] = y1.date().isoformat()
        sb = g("period ended (stub")
        if isinstance(sb, dt.datetime):
            pr["stub_end"] = sb.date().isoformat()
            stub = True
        for r in range(1, 60):
            a = str(md[f"A{r}"].value or "")
            if a.startswith("Closing stock") and "September" not in a and isinstance(mv[f"B{r}"].value, (int, float)):
                if stub and r == 19:
                    pr["stub_closing_stock"] = mv[f"B{r}"].value
                else:
                    pr["closing_stock"] = mv[f"B{r}"].value
            if a.startswith("Closing stock") and stub and r != 16 and isinstance(mv[f"B{r}"].value, (int, float)):
                pr["stub_closing_stock"] = mv[f"B{r}"].value
        pr["auditor"] = {"firm": g("audit firm"), "frn": g("firm reg"), "peer": g("peer review"), "partner": g("engagement partner"),
                         "mno": g("membership"), "sign_date": g("date of signing"), "udin": g("udin"), "place": g("place")}
        pr["auditor"] = {k: ("" if v is None else str(v)) for k, v in pr["auditor"].items()}
        people = []
        for r in range(4, 30):
            if md[f"F{r}"].value in (None, "") and md[f"D{r}"].value in (None, ""):
                break
            if md[f"F{r}"].value:
                people.append({"din": str(md[f"E{r}"].value or ""), "name": str(md[f"F{r}"].value or "").title(), "designation": md[f"G{r}"].value or "",
                               "category": md[f"H{r}"].value or "", "signatory": md[f"I{r}"].value or ""})
        pr["directors"] = people
        r0 = next((r for r in range(25, 60) if str(md[f"A{r}"].value or "").lower().startswith("authorised capital history")), None)
        if r0:
            N["auth_history"] = [{"date": _iso(mv[f"A{r}"].value), "amount": mv[f"B{r}"].value, "remarks": md[f"D{r}"].value}
                                 for r in range(r0 + 2, r0 + 12) if isinstance(mv[f"B{r}"].value, (int, float))]
        r0 = next((r for r in range(25, 70) if str(md[f"A{r}"].value or "").lower().startswith("statutory auditor")), None)
        if r0:
            pr["stat_auditor"] = {"firm": md[f"B{r0 + 1}"].value or "", "frn": md[f"B{r0 + 2}"].value or "", "partner": md[f"B{r0 + 3}"].value or "",
                                  "reports": [[md[f"A{r}"].value, md[f"B{r}"].value] for r in range(r0 + 4, r0 + 10) if md[f"A{r}"].value]}
        msg.append(f"Company details, {len(people)} directors / signatories, auditor")
    keys = (["S"] if stub else []) + ["Y1", "Y2", "Y3"]
    dk = {"S": "dS", "Y1": "d26", "Y2": "d25", "Y3": "d24"}
    # ageing (blocks: heading 'As at ...', header row, 4 rows)
    for sh, key, nb in (("Debtor Ageing", "ageing_dr", 5), ("Creditor Ageing", "ageing_cr", 4)):
        if sh not in S:
            continue
        ws = wv[sh]
        blocks = [r for r in range(1, ws.max_row + 1) if str(wb[sh][f"B{r}"].value or "").strip().lower() == "particulars"]
        out = {}
        for k, r in zip(keys, blocks):
            vals = [[ws.cell(r + 1 + i, 3 + j).value or 0 for j in range(nb)] for i in range(4)]
            if any(any(isinstance(x, (int, float)) and x for x in row) for row in vals):
                out[dk[k]] = vals
        pr[key] = out
        msg.append(f"{sh}: {len(out)} date(s)")
    if "Shareholder Register" in S:
        ws = wv["Shareholder Register"]
        sk = {"S": "sS", "Y1": "s26", "Y2": "s25", "Y3": "s24"}
        sh = []
        prom = set()
        if "Notes_Equity & Reserves" in S:
            ne = wb["Notes_Equity & Reserves"]
            r1, r2 = _find(ne, "(b) details of shareholding of promoters"), _find(ne, "(c) details of shareholders")
            if r1 and r2:
                prom = {re.sub(r"^\d+\.\s*", "", str(ne[f"C{r}"].value)).strip().lower() for r in range(r1, r2)
                        if re.match(r"^\d+\.", str(ne[f"C{r}"].value or ""))}
        for r in range(5, ws.max_row + 1):
            n = ws[f"B{r}"].value
            if not n or str(n).strip().lower() == "total":
                break
            sh.append({"name": n, **{sk[k]: ws.cell(r, 3 + i).value or 0 for i, k in enumerate(keys)}, "promoter": str(n).strip().lower() in prom})
        pr["shareholders"] = sh
        msg.append(f"Shareholder register: {len(sh)} shareholders")
    # restatement entries (typed into tab 8 as amounts in lakhs)
    if "Adj Entry RFS" in S:
        ws, wvv = wb["Adj Entry RFS"], wv["Adj Entry RFS"]
        rows, eno, narr = [], 0, ""
        for r in range(3, ws.max_row + 1):
            c = ws[f"C{r}"].value
            a = ws[f"A{r}"].value
            if c:
                if a is not None and not str(a).startswith("=A"):
                    eno += 1
                d = wvv[f"A{r}"].value
                b = str(ws[f"B{r}"].value or "").strip()
                rows.append({"entry": eno, "date": _iso(d), "master": c, "particulars": b[3:].strip() if b.startswith("To ") else b,
                             "dr": float(wvv[f"E{r}"].value or 0), "cr": float(wvv[f"F{r}"].value or 0), "narration": ""})
            elif rows and str(ws[f"B{r}"].value or "").startswith("("):
                for x in rows:
                    if x["entry"] == eno and not x["narration"]:
                        x["narration"] = str(ws[f"B{r}"].value).strip()[1:-1]
                        break
        pr["adj_rows"] = rows
        msg.append(f"Restatement entries: {eno} (tab 8)")
    # notes inputs
    if "Notes_Liabilities" in S:
        ws = wb["Notes_Liabilities"]
        r = _find(ws, "current maturities", "C") or _find(ws, '="current maturities', "C")
        if not r:
            r = next((x for x in range(1, ws.max_row + 1) if "current maturities" in str(ws[f"C{x}"].value or "").lower()), None)
        if r:
            cm = {}
            for i, k in enumerate(keys + ["OPEN"]):
                v = ws.cell(r, 4 + i).value
                rs = _rs(v)
                cm[k] = "all" if isinstance(v, str) and "SUMIF" in v else (str(rs) if rs is not None else None)
            N["current_maturities"] = cm
    if "Borrowing Terms" in S:
        ws = wb["Borrowing Terms"]
        L = []
        for r in range(11, ws.max_row + 1):
            if not ws[f"C{r}"].value or str(ws[f"C{r}"].value).strip().lower() == "total":
                break
            led = re.findall(r'\$A:\$A,"([^"]+)"', str(ws[f"J{r}"].value or ""))
            lim = ws[f"F{r}"].value
            L.append({"lender": ws[f"C{r}"].value, "nature": ws[f"D{r}"].value, "date": _iso(ws[f"E{r}"].value),
                      "limit_rs": lim * 100000 if isinstance(lim, (int, float)) else None, "roi": ws[f"G{r}"].value, "tenor": ws[f"H{r}"].value,
                      "security": ws[f"I{r}"].value, "ledgers": ", ".join(led), "guarantee": "guarantee" in str(ws[f"I{r}"].value or "").lower()})
        N["borrowings"] = L
        msg.append(f"Borrowing terms: {len(L)} facilities")
    if "Notes_P&L" in S:
        ws = wb["Notes_P&L"]
        r0, r1 = _find(ws, "other income", "C"), None
        oi = []
        if r0:
            for r in range(r0 + 1, r0 + 30):
                f = str(ws[f"D{r}"].value or "")
                if str(ws[f"C{r}"].value or "").strip().lower() == "total":
                    break
                m = re.search(r'\$G:\$G,"([^"]+)",[^,]+\$A:\$A,"([^"]+)"', f)
                if m:
                    lbl = str(ws[f"C{r}"].value)
                    oi.append({"label": lbl, "master": m.group(1), "ledgers": m.group(2), "interest": "interest" in lbl.lower(),
                               "fd": "fixed deposit" in lbl.lower(), "ppe_sale": "sale of property" in lbl.lower()})
        if oi:
            N["oi_lines"] = oi
    if "RPT" in S:
        ws = wb["RPT"]
        r0 = _find(ws, "directors' remuneration", "C")
        rem = []
        if r0:
            for r in range(r0 + 1, r0 + 20):
                n = str(ws[f"C{r}"].value or "")
                if not n or n.lower() == "total" or n.lower().startswith("not allocated"):
                    break
                rem.append({"name": n, **{k: _rs(ws.cell(r, 4 + i).value) for i, k in enumerate(keys)}})
        N["remuneration"] = rem
        r1, r2 = _find(ws, "key managerial personnel", "C"), _find(ws, "(b) transactions", "C")
        if r1 and r2:
            nat = [ws[f"D{r}"].value for r in range(r1 + 1, r2) if ws[f"D{r}"].value and str(ws[f"C{r}"].value or "").startswith("=PROPER")]
            N["kmp_nature"] = {p["name"]: nat[i] for i, p in enumerate(pr.get("directors", [])) if i < len(nat)}
    if "Gratuity and contingencies" in S:
        ws = wb["Gratuity and contingencies"]
        fy = [k for k in keys if k != "S"]
        G = {"assumptions": [], "dbo": {}}
        ra, rb = _find(ws, "a) actuarial assumptions", "C"), _find(ws, "b) change in the present value", "C")
        if ra and rb:
            for r in range(ra + 2, rb):
                lbl = ws[f"C{r}"].value
                vals = [ws.cell(r, 4 + PER_COLS[k]).value for k in fy]
                if lbl or any(v not in (None, "") for v in vals):
                    G["assumptions"].append({"label": lbl or "", **{k: v for k, v in zip(fy, vals)}})
        rows = {"present value of obligation as at the beginning": "open", "current service cost": "service", "interest cost": "interest",
                "benefits paid": "paid", "actuarial (gain) / loss on obligations": "actuarial", "present value of obligation as at the end": "closing"}
        if rb:
            for r in range(rb + 2, rb + 9):
                lbl = str(ws[f"C{r}"].value or "").lower()
                for k_, key in rows.items():
                    if lbl.startswith(k_) and key not in G["dbo"]:
                        G["dbo"][key] = [_rs(ws.cell(r, 4 + PER_COLS[k]).value) for k in fy]
        rd = _find(ws, "add: gratuity paid directly", "C")
        if rd:
            G["dbo"]["direct_paid"] = [_rs(ws.cell(rd, 4 + PER_COLS[k]).value) for k in fy]
        re_ = _find(ws, "experience adjustments on plan", "C")
        if re_:
            G["dbo"]["experience"] = [_rs(ws.cell(re_, 4 + PER_COLS[k]).value) for k in fy]
        rc = _find(ws, "current (included", "C")
        if rc:
            G["current"] = {k: _rs(ws.cell(rc, 4 + i).value) for i, k in enumerate(keys + ["OPEN"])}
        rs_ = _find(ws, "discount rate", "C", (_find(ws, "h) sensitivity", "C") or ws.max_row))
        if rs_:
            G["sensitivity"] = [[_rs(ws.cell(rs_ + i, 5).value), _rs(ws.cell(rs_ + i, 6).value)] for i in range(3)]
        m = re.search(r"independent actuary \(([^)]+)\)", str(ws[f"C{(_find(ws, 'the company provides for gratuity', 'C') or 1)}"].value or ""))
        if m:
            G["actuary"] = m.group(1)
        N["gratuity"] = G
        cont = {}
        for key, lbl in (("claims", "(a) claims against"), ("guarantees", "(b) guarantees issued"), ("other", "(c) other money")):
            r = _find(ws, lbl, "C")
            if r:
                cont[key] = {k: _rs(ws.cell(r, 4 + i).value) for i, k in enumerate(keys)}
        N["contingent"] = cont
        msg.append("Gratuity disclosure, contingent liabilities")
    if "Tax Computation" in S:
        ws = wb["Tax Computation"]
        T = {"other_rows": [], "permanent": []}
        r = _find(ws, "depreciation as per the income-tax act", "C")
        if r:
            T["it_dep"] = {k: (abs(_rs(ws.cell(r, 4 + i).value)) if _rs(ws.cell(r, 4 + i).value) else None) for i, k in enumerate(keys)}
        r0, r1 = _find(ws, "other differences", "C"), _find(ws, "total timing differences", "C")
        if r0 and r1:
            for r in range(r0 + 1, r1):
                lbl = str(ws[f"C{r}"].value or "")
                if lbl and "provision for gratuity" not in lbl.lower():
                    T["other_rows"].append({"label": lbl, **{k: _rs(ws.cell(r, 4 + i).value) for i, k in enumerate(keys)}})
        r = _find(ws, "assessment year", "C")
        if r:
            T["itr"] = {k: [ws.cell(r + j, 4 + i).value if not str(ws.cell(r + j, 4 + i).value or "").startswith("=") else None for j in range(7)]
                        for i, k in enumerate(keys)}
        N["tax"] = T
        msg.append("Tax computation inputs and return details")
    if "Additional Disclosures" in S:
        ws = wb["Additional Disclosures"]
        r = _find(ws, "for taxation matters", "C")
        if r:
            N["auditor_fees"] = {"tax_audit": {k: _rs(ws.cell(r, 4 + i).value) for i, k in enumerate(keys)}}
        fx = {}
        for key, lbl in (("imports", "a. cif value"), ("exports", "b. earnings in foreign"), ("expenditure", "c. expenditure in foreign")):
            r = _find(ws, lbl, "C")
            if r:
                fx[key] = {k: (_rs(ws.cell(r, 4 + i).value) if _rs(ws.cell(r, 4 + i).value) is not None else
                               (ws.cell(r, 4 + i).value if isinstance(ws.cell(r, 4 + i).value, str) else None)) for i, k in enumerate(keys)}
                if key == "expenditure":
                    fx["expenditure_label"] = ws[f"C{r}"].value
        N["foreign"] = fx
    return pr, N, msg
