"""Annexure VI note sheets: Notes_Equity & Reserves, Notes_Liabilities, Notes_Assets, Notes_P&L."""
from openpyxl.styles import Font
from .book_core import *          # noqa: F401,F403
from . import book_masters as BM

NOTE_W = {"A": 1.5, "B": 5.0, "C": 55.0}


class NoteSheet:
    """A sheet of consecutive notes (number in B, label in C, periods from D, opening balance column for BS notes)."""

    def __init__(self, cx, name, bs=True, prev_ref=None, title="ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS"):
        self.cx, self.bs = cx, bs
        widths = dict(NOTE_W)
        for i, p in enumerate(cx.PK):
            widths[cx.NC[p]] = 16.0
        widths[cx.NC["OPEN"]] = 15.0
        widths[CL(5 + cx.N)] = 16.0
        for i in range(3):
            widths[CL(6 + cx.N + i)] = 12.0
        self.ws = cx.sheet(name, widths, print_cols=cx.LASTN if bs else cx.LASTP)
        self.name = name
        self.last_note = None
        self.prev_ref = prev_ref
        head(cx, self.ws, title, cx.LASTN if bs else cx.LASTP)
        self.r = 8
        self.notes = {}

    # -------------------------------------------------------------- structure
    def note(self, key, title):
        ws, cx = self.ws, self.cx
        if self.last_note is None:
            w(ws, f"B{self.r}", f"={self.prev_ref}+1" if self.prev_ref else 1, fB, al=A_L)
        else:
            w(ws, f"B{self.r}", f"=B{self.last_note}+1", fB, al=A_L)
        w(ws, f"C{self.r}", title, fB)
        self.last_note = self.r
        cx.noteno_counter += 1
        cx.noteno[key] = cx.noteno_counter
        cx.setref("note:" + key, self.r)
        cx.notes_sheet = getattr(cx, "notes_sheet", {})
        cx.notes_sheet[key] = (self.name, self.r)
        self.notes[key] = self.r
        self.r += 2
        return self.r - 2

    def colhead(self, kind="bs", opening=None, label="Particulars"):
        ws, cx = self.ws, self.cx
        opening = self.bs if opening is None else opening
        w(ws, f"C{self.r}", label, fB, bd=B_TB)
        for p in cx.PK:
            w(ws, f"{cx.NC[p]}{self.r}", cx.dtxt(p, kind), fB, bd=B_TB, al=A_CW)
        if opening:
            w(ws, f"{cx.NC['OPEN']}{self.r}", cx.dtxt("OPEN", "bs"), fGreyI, bd=B_TB, al=A_CW)
        ws.row_dimensions[self.r].height = 30
        self.r += 1

    def sub(self, text, font=fB):
        w(self.ws, f"C{self.r}", text, font)
        self.r += 1

    def keys(self, opening=None):
        opening = self.bs if opening is None else opening
        return self.cx.PK + (["OPEN"] if opening else [])

    def line(self, label, f, font=fN, fmt=ACC, opening=None, label_font=None, bd=None, ref=None):
        """f(p, r) -> formula / value for period p (and 'OPEN')"""
        ws, cx = self.ws, self.cx
        r = self.r
        w(ws, f"C{r}", label, label_font or font, bd=bd)
        for p in self.keys(opening):
            v = f(p, r)
            if v is None:
                continue
            fo = font
            if p == "OPEN":
                fo = Font(name=TN, size=10, italic=True, color="808080", bold=font.b)
            w(ws, f"{cx.NC[p]}{r}", v, fo, fmt, bd=bd)
        if ref:
            cx.setref(ref, r)
        self.r += 1
        return r

    def master_line(self, master, label=None, nat=None, grp="BS", minus=None, opening=None, keep=False, ref=None):
        cx = self.cx
        if not keep and not cx.nonzero(master, self.keys(opening)):
            return None
        info = BM.info(master, cx.custom) or ("Dr", grp, None)
        nat = nat or info[0]
        grp = info[1]
        label = label or master

        def f(p, r):
            mref = f"$C{r}" if label == master else qs(master)
            v = "=" + cx.master_val(p, master, mref, grp, nat)
            if minus:
                v += "-" + minus(p)
            return v
        return self.line(label, f, opening=opening, ref=ref)

    def total(self, first, last, label="Total", fmt=ACC, bd=B_TB2, opening=None, ref=None, rows=None):
        def f(p, r):
            c = self.cx.NC[p]
            if rows is not None:
                return "=" + ("+".join(f"{c}{x}" for x in rows) if rows else "0")
            return f"=SUM({c}{first}:{c}{last})"
        return self.line(label, f, font=fB, fmt=fmt, bd=bd, opening=opening, ref=ref)

    def para(self, text, font=fN, al=A_J, last_col=None, height=None):
        ws = self.ws
        lc = last_col or self.cx.LASTP
        merge(ws, f"C{self.r}:{lc}{self.r}", text, font, al)
        width = 55 + 16 * (CI(lc) - 3)
        ws.row_dimensions[self.r].height = height or est_height(text if not str(text).startswith("=") else "x" * 300, width)
        self.r += 1

    def gap(self, n=1):
        self.r += n

    def checks(self, r, fmls):
        """reviewer checks to the right of the table (grey, small)"""
        ws, cx = self.ws, self.cx
        for i, p in enumerate(cx.PK):
            w(ws, f"{CL(6 + cx.N + i)}{r}", fmls(p), fChk, ACC)


def head(cx, ws, title, last_col, col="B", units=True):
    w(ws, f"{col}2", f"={MD}!$B$3", fB, al=A_C)
    w(ws, f"{col}3", f"={MD}!$B$4", fB, al=A_C)
    w(ws, f"{col}4", f'="CIN: "&{MD}!$B$2', fB, al=A_C)
    for rr in (2, 3, 4):
        ws.merge_cells(f"{col}{rr}:{last_col}{rr}")
    w(ws, f"{col}5", title, fB, al=A_L)
    if units:
        w(ws, f"{col}6", "(All amounts in Rs. lakhs, except as otherwise stated)", fI, al=A_L)


# ====================================================================== Notes_Equity & Reserves
def notes_equity(cx):
    I = cx.I
    ns = NoteSheet(cx, "Notes_Equity & Reserves", True, prev_ref="SAP_4!B12")
    ws = ns.ws
    S = cx.SCALE
    FV = f"{MD}!$B$7"
    ns.note("SC", "Share capital")
    ns.colhead()
    ns.sub("Authorised share capital")
    acr = f"{MD}!$A$31:$A${cx.AC_LAST}"
    acv = f"{MD}!$B$31:$B${cx.AC_LAST}"
    ra = ns.line(f'="Equity shares of Rs. "&{FV}&"/- each"', lambda p, r: f'=SUMIF({acr},"<="&{cx.MDC[p]},{acv})/{S}')
    ns.line("   (Number of equity shares)", lambda p, r: f"={cx.NC[p]}{ra}*{S}/{FV}", font=fI, fmt=ACC0)
    for p in ns.keys():
        ws[f"{cx.NC[p]}{ra + 1}"].font = fI if p != "OPEN" else fGreyI
    ns.sub("Issued, subscribed and fully paid up")
    r1 = ns.master_line("Equity share capital", keep=True)
    rdesc = ns.r
    ns.line(f'="("&TEXT(D{{{{sc_end}}}},"#,##0")&" Equity shares of Rs. "&{FV}&"/- each fully paid up)"', lambda p, r: None, font=fI)
    rt = ns.total(r1, r1, ref="sc_total")
    ns.gap()
    ns.para(cx.text("sc_terms", f'="Terms / rights attached to equity shares: The Company has only one class of equity shares having a par value of Rs. "&{FV}&'
                    '" per share. Each holder of equity shares is entitled to one vote per share. In the event of liquidation, the equity shareholders are eligible '
                    'to receive the remaining assets of the Company after distribution of all preferential amounts, in proportion to their shareholding."'), height=52.5)
    ac = I.get("auth_capital") or []
    if len(ac) <= 1 and ac:
        a0 = ac[0]
        amt = float(a0.get("amount") or 0)
        d0 = ddate(a0.get("date"))
        fv = float(I["company"].get("fv") or 10)
        dflt = (f"Note: There has been no change in the authorised share capital of the Company since its incorporation on {long_date(d0) if d0 else '[date]'} - "
                f"Rs. {inr(amt)} divided into {inr(amt / fv)} equity shares of Rs. {fv:g} each.")
    else:
        dflt = ""
    t = cx.text("sc_auth_note", dflt)
    if t:
        ns.para(t, fI, last_col=cx.LASTN, height=26.25)
    ns.sub("(a) Reconciliation of the number of shares outstanding at the beginning and at the end of the year")
    rb = ns.r
    nxt = {p: (cx.NC[cx.PREV[p]]) for p in cx.PK}
    ns.line("No. of shares at the beginning of the year", lambda p, r: f"={nxt[p]}{rt}*{S}/{FV}", fmt=ACC0, opening=False)
    ns.line("Add: Shares issued during the year", lambda p, r: f"={cx.NC[p]}{rt}*{S}/{FV}-{cx.NC[p]}{rb}", fmt=ACC0, opening=False)
    re_ = ns.total(rb, rb + 1, "No. of shares at the end of the year", fmt=ACC0, opening=False, ref="sc_end")
    ns.gap()
    # promoters
    SH = I.get("shareholders") or []
    prom = [h for h in SH if h.get("promoter")]
    ns.sub("(b) Details of shareholding of promoters")
    ns.colhead("No. of shares as at\n", opening=False, label="Name of promoter")
    ws.row_dimensions[ns.r - 1].height = 30
    for i, h in enumerate(prom):
        rr = cx.SHROW[h["name"]]
        rn = ns.line(f"{i + 1}. {h['name']}", lambda p, r: f"={q('Shareholder Register')}!{cx.SHC[p]}{rr}", fmt=ACC0, opening=False)
        rp = ns.line("   % of total shares", lambda p, r: f"=IFERROR({cx.NC[p]}{rn}/{cx.NC[p]}{re_},0)", fmt=PCT, opening=False, label_font=fI)
        ns.line("   % change during the year", lambda p, r: ("=0" if p == cx.PK[-1] else f"={cx.NC[p]}{rp}-{cx.NC[cx.PREV[p]]}{rp}"),
                fmt=PCT, opening=False, label_font=fI)
    ns.gap()
    ns.sub("(c) Details of shareholders holding more than 5% of shares")
    tot = {p: sum(float(h.get(p) or 0) for h in SH) for p in cx.PK}
    big = [h for h in SH if any(tot[p] and float(h.get(p) or 0) / tot[p] > 0.05 for p in cx.PK)]
    for i, h in enumerate(big):
        rr = cx.SHROW[h["name"]]
        rn = ns.line(f"{i + 1}. {h['name']} (No. of shares)", lambda p, r: f"={q('Shareholder Register')}!{cx.SHC[p]}{rr}", fmt=ACC0, opening=False)
        ns.line("   % holding", lambda p, r: f"=IFERROR({cx.NC[p]}{rn}/{cx.NC[p]}{re_},0)", fmt=PCT, opening=False, label_font=fI)
    ns.para(cx.text("sc_records", "As per records of the Company, including its register of shareholders / members. There are no shares reserved for issue "
                    "under options, no bonus shares, no shares allotted for consideration other than cash and no shares bought back during the five years "
                    f"immediately preceding {long_date(cx.DATE['Y1'])}."), height=39)
    ns.gap()
    # reserves and surplus
    ns.note("RS", "Reserves and surplus")
    ns.colhead()
    ns.sub("Surplus in the Statement of Profit and Loss")
    ro = ns.r
    other = [m for m in (I.get("reserve_movements") or [])]
    ns.line("Opening balance", lambda p, r: f"={cx.NC[cx.PREV[p]]}{{{{rs_end}}}}", opening=False)
    ns.line("Add: Restated profit for the year / period", lambda p, r: f"=PL_2!{cx.BC[p]}{{{{PL2_pat}}}}", opening=False)
    for mv in other:
        ns.line(mv["label"], lambda p, r, mv=mv: f"={float(mv.get(p) or 0)}/{S}" if mv.get(p) else 0, opening=False)
    rs = ns.r
    op_f = ("=" + cx.tb("OPEN", qs("Reserves and surplus")) + "+(" + cx.adj_bs(qs("Reserves and surplus"), "OPEN", "Cr")[1:-1]
            + "-(" + _adj_pl_total(cx, "OPEN") + "))")
    ns.line("Balance at the end of the year / period", lambda p, r: op_f if p == "OPEN" else f"=SUM({cx.NC[p]}{ro}:{cx.NC[p]}{r - 1})",
            font=fB, bd=B_TB2, ref="rs_end")
    # movements in the books' reserves other than the profit for the period must be reversed by restatement entries
    w(ws, f"C{ns.r}", "Check: movements in reserves in the books other than profit, not reversed by restatement entries (should be nil)", fChkI)
    RSq = qs("Reserves and surplus")
    for p in cx.PK:
        w(ws, f"{cx.NC[p]}{ns.r}", f"=ROUND({cx.tb(p, RSq)}-{cx.tb(cx.PREV[p], RSq)}-{cx.tb(p, qs('Net profit as per books'))}"
                                   f"+{cx.adj_pl(RSq, p, 'Cr')},2)", fChk, ACC)
    ns.r += 1
    w(ws, f"{CL(5 + cx.N)}{rs}", "Check vs Annexure V", fChkI)
    for i, p in enumerate(cx.PK):
        w(ws, f"{CL(6 + cx.N + i)}{rs}", f"=ROUND({cx.NC[p]}{rs}+{cx.NC[p]}{rt}-SOA_5!{cx.SC[p]}{{{{soa_sh_restated}}}},2)", fChk, ACC)
    cx.finish_print(ws, ns.r)
    return ns


def _adj_pl_total(cx, p):
    """all P&L restatement entries up to the date of p (Dr - Cr): already carried to reserves through restated profits"""
    d = cx.MDC[p]
    rng = cx._rng
    return (f"SUMIFS({rng('E')},{rng('D')},\"P&L\",{rng('A')},\"<=\"&{d})-SUMIFS({rng('F')},{rng('D')},\"P&L\",{rng('A')},\"<=\"&{d})")


def inr(x):
    """Indian digit grouping"""
    x = round(float(x))
    s = str(abs(int(x)))
    if len(s) > 3:
        head_, tail = s[:-3], s[-3:]
        parts = []
        while len(head_) > 2:
            parts.insert(0, head_[-2:])
            head_ = head_[:-2]
        if head_:
            parts.insert(0, head_)
        s = ",".join(parts) + "," + tail
    return ("-" if x < 0 else "") + s


# ====================================================================== Notes_Liabilities
def notes_liabilities(cx):
    I = cx.I
    ns = NoteSheet(cx, "Notes_Liabilities", True, prev_ref=f"{q('Notes_Equity & Reserves')}!$B${{{{note:RS}}}}")
    ws = ns.ws
    S = cx.SCALE
    # ---- 5 long-term borrowings
    ns.note("LTB", "Long-term borrowings")
    ns.colhead()
    first = ns.r
    sec = [m for m, _ in BM.note_lines("LTB_SEC", cx.custom) if cx.nonzero(m)]
    if sec:
        ns.sub("Secured")
        for m in sec:
            if m == "Term loans from banks and NBFCs (secured)":
                ns.master_line(m, minus=lambda p: f"{cx.NC[p]}{{{{cur_mat}}}}", ref="ltb_sec_tl")
            else:
                ns.master_line(m)
    dirs = sorted({m for (p, m) in cx.tbval if m.startswith(BM.DIRECTOR_PREFIX)}, key=lambda m: cx.dir_order(m))
    uns = [m for m in dirs if cx.nonzero(m)] + [m for m, _ in BM.note_lines("LTB_UNS", cx.custom) if cx.nonzero(m)]
    cx.DIR_LOAN_ROW = {}
    if uns:
        ns.sub("Unsecured - Loans from directors (interest free)" if dirs else "Unsecured")
        for m in uns:
            rr = ns.master_line(m)
            cx.DIR_LOAN_ROW[m] = rr
    rt = ns.total(first, ns.r - 1, ref="ltb_total")
    ns.para(cx.text("ltb_terms", "Loans from directors are interest free and repayable on demand at the option of the lender. As there is no stipulation for "
                    "repayment within twelve months, these have been classified as long-term borrowings. Term loans are secured by hypothecation of the "
                    "respective assets; the portion repayable within twelve months is shown under short-term borrowings (current maturities). "
                    "Refer Note {NOTE:BT} for the principal terms of borrowings."), height=None)
    # movement in director loans
    cx.DIR_MOV = {"acc": {}, "rep": {}}
    if dirs:
        ns.gap()
        ns.sub("Movement in loans from directors during the year")
        for kind, lbl, col in (("acc", "Loans accepted", "D"), ("rep", "Loans repaid", "C")):
            mv = [m for m in dirs if any(cx.tbmov.get((p, m, col), 0) > 0.0000001 for p in cx.PK)]
            if not mv:
                continue
            ns.colhead("pl", opening=False, label=lbl)
            ws[f"C{ns.r - 1}"].font = fBI
            for m in mv:
                rr = ns.line(m, lambda p, r, m=m, col=col: f"=SUMIF({q(cx.TBN[p])}!$G:$G,$C{r},{q(cx.TBN[p])}!${col}:${col})/{S}", opening=False)
                cx.DIR_MOV[kind][m] = rr
    ns.gap()
    # ---- 6 other long-term liabilities
    ol = [m for m, _ in BM.note_lines("OLTL", cx.custom) if cx.nonzero(m)]
    cx.has_oltl = bool(ol)
    if ol:
        ns.note("OLTL", "Other long-term liabilities")
        ns.colhead()
        f0 = ns.r
        for m in ol:
            ns.master_line(m)
        ns.total(f0, ns.r - 1, ref="oltl_total")
        ns.gap()
    # ---- 7 long-term provisions
    G = "Provision for gratuity"
    lp = [m for m, _ in BM.note_lines("LTP", cx.custom) if cx.nonzero(m)]
    cx.has_ltp = cx.nonzero(G) or bool(lp)
    if cx.has_ltp:
        ns.note("LTP", "Long-term provisions")
        ns.colhead()
        f0 = ns.r
        if cx.nonzero(G):
            ns.sub("Provision for employee benefits")
            f0 = ns.r
            ns.line(f'="Provision for gratuity (non-current portion) (refer Note "&{q("Gratuity and contingencies")}!$B${{{{note:EB}}}}&")"',
                    lambda p, r: "=" + cx.tb(p, qs(G)) + "+" + cx.adj_bs(qs(G), p, "Cr") + f"-{cx.NC[p]}{{{{grat_cur}}}}"
                    if G in cx.adj_masters else "=" + cx.tb(p, qs(G)) + f"-{cx.NC[p]}{{{{grat_cur}}}}", ref="ltp_grat")
        for m in lp:
            ns.master_line(m)
        ns.total(f0, ns.r - 1, ref="ltp_total")
        ns.gap()
    # ---- 8 short-term borrowings
    ns.note("STB", "Short-term borrowings")
    ns.colhead()
    f0 = ns.r
    CM = I.get("current_maturities") or {}
    TL = "Term loans from banks and NBFCs (secured)"

    def cm(p, r):
        v = CM.get(p)
        if isinstance(v, str) and v.lower() == "all":
            return "=" + cx.tb(p, qs(TL)) if p != "OPEN" else "=" + cx.tb("OPEN", qs(TL))
        return f"={float(v or 0):g}/{S}" if v not in (None, "") else 0
    rcm = ns.line(f'="Current maturities of long-term borrowings (secured) - refer Note "&$B${cx.refs["note:LTB"]}', cm, ref="cur_mat")
    for p in ns.keys():
        ws[f"{cx.NC[p]}{rcm}"].font = fBlue if p != "OPEN" else Font(name=TN, size=10, italic=True, color="0000FF")
    for m, _ in BM.note_lines("STB", cx.custom):
        ns.master_line(m)
    ns.total(f0, ns.r - 1, ref="stb_total")
    ns.para(cx.text("stb_security", f'="The security and other terms of the short-term borrowings are set out in Note "&{q("Borrowing Terms")}!$B$8&'
                                    '" (principal terms of borrowings)."'))
    ns.gap()
    # ---- 9 trade payables
    ns.note("TP", "Trade payables")
    ns.colhead()
    f0 = ns.r
    ns.master_line("Dues to micro and small enterprises", keep=True, ref="tp_msme")
    ns.master_line("Dues to other creditors", keep=True, ref="tp_oth")
    for m, _ in BM.note_lines("TP", cx.custom)[2:]:
        ns.master_line(m)
    rtp = ns.total(f0, ns.r - 1, ref="tp_total")
    ns.gap()
    tpn = cx.refs["note:TP"]
    w(ws, f"C{ns.r}", f'=$B${tpn}&"(a) Dues to Micro and Small Enterprises"', fB)
    ns.r += 1
    merge(ws, f"C{ns.r}:{cx.LASTN}{ns.r}", 'Disclosure as per Section 22 of the Micro, Small and Medium Enterprises Development Act, 2006 ("MSMED Act"):', fB,
          A_LW, 13.5)
    ns.r += 1
    r = ns.r
    w(ws, f"C{r}", "Particulars", fB, bd=B_ALL, al=A_CW)
    for p in cx.PK:
        w(ws, f"{cx.NC[p]}{r}", cx.dtxt(p), fB, bd=B_ALL, al=A_CW)
    ws.row_dimensions[r].height = 27.75
    ns.r += 1
    MI = "Interest accrued on dues to MSME"
    rows = [
        ("a) Dues remaining unpaid to any supplier as at the end of each accounting year:", None, 27.0),
        ("   - Principal amount due to micro and small enterprises", lambda p: f"={cx.NC[p]}{{{{tp_msme}}}}", 13.5),
        ("   - Interest due on the above", lambda p: f"={cx.NC[p]}{{{{ocl_msme_int}}}}" if cx.nonzero(MI) else 0, 13.5),
        ("b) The amount of interest paid by the buyer in terms of Section 16 of the MSMED Act, along with the amount of the payment made to the supplier "
         "beyond the appointed day during each accounting year", lambda p: 0, 39.75),
        ("c) The amount of interest due and payable for the period of delay in making payment (which have been paid but beyond the appointed day during "
         "the year) but without adding the interest specified under the MSMED Act",
         lambda p: f"='Notes_P&L'!{cx.NC[p]}{{{{pl_msme_int}}}}" if cx.nonzero("Interest on delayed payments to MSME", cx.PK) else 0, 39.75),
        ("d) The amount of interest accrued and remaining unpaid at the end of each accounting year",
         lambda p: f"={cx.NC[p]}{{{{ocl_msme_int}}}}" if cx.nonzero(MI) else 0, 27.0),
        ("e) The amount of further interest remaining due and payable even in the succeeding years, until such date when the interest dues above are actually "
         "paid to the small enterprise, for the purpose of disallowance of a deductible expenditure under Section 23 of the MSMED Act",
         lambda p: f"={cx.NC[p]}{{{{ocl_msme_int}}}}" if cx.nonzero(MI) else 0, 39.75),
    ]
    for i, (lbl, f, h) in enumerate(rows):
        r = ns.r
        bd = B_BLR if i == len(rows) - 1 else B_LR
        w(ws, f"C{r}", lbl, fN, bd=bd, al=A_LW)
        for p in cx.PK:
            w(ws, f"{cx.NC[p]}{r}", f(p) if f else None, fN, ACC, bd=bd)
        ws.row_dimensions[r].height = h
        ns.r += 1
    merge(ws, f"C{ns.r}:{cx.LASTN}{ns.r}", cx.text("msme_note", "Note: The above information regarding micro and small enterprises has been determined to "
                                                   "the extent such parties have been identified on the basis of information available with the Company. "
                                                   "Interest has been computed by the management at three times the bank rate notified by the Reserve Bank of India."),
          fI, A_J, 39.0)
    ns.r += 2
    # ageing schedules
    letters = "bcdefgh"
    for i, p in enumerate(cx.PK):
        r = ns.r
        w(ws, f"C{r}", f'=$B${tpn}&"({letters[i]}) "&"Trade payables ageing schedule as at "&TEXT({cx.MDC[p]},"mmmm dd, yyyy")', fB)
        r += 1
        w(ws, f"C{r}", "Particulars", fB, bd=B_ALL, al=A_CW)
        ws.merge_cells(f"C{r}:C{r + 1}")
        w(ws, f"D{r}", "Outstanding for following periods from due date of payment", fB, bd=B_ALL, al=A_CW)
        ws.merge_cells(f"D{r}:H{r}")
        for j, h in enumerate(["Less than 1 year", "1 - 2 years", "2 - 3 years", "More than 3 years", "Total"]):
            w(ws, f"{CL(4 + j)}{r + 1}", h, fB, bd=B_ALL if j < 4 else B_LR, al=A_CW)
        r += 2
        f0 = r
        base = cx.CRE[p]
        for k, lbl in enumerate(["(A) Total outstanding dues of micro and small enterprises",
                                 "(B) Total outstanding dues of creditors other than micro and small enterprises",
                                 "(C) Disputed dues of micro and small enterprises", "(D) Disputed dues of creditors other than micro and small enterprises"]):
            w(ws, f"C{r}", lbl, fN, bd=B_LR, al=A_LW)
            for j in range(4):
                w(ws, f"{CL(4 + j)}{r}", f"={q('Creditor Ageing')}!{CL(3 + j)}{base + k}/{S}", fN, ACC, bd=B_LR)
            w(ws, f"H{r}", f"=SUM(D{r}:G{r})", fN, ACC, bd=B_LR)
            ws.row_dimensions[r].height = 13.5 if k else 13.5
            r += 1
        w(ws, f"C{r}", "Total", fB, bd=B_ALL)
        for j in range(5):
            c = CL(4 + j)
            w(ws, f"{c}{r}", f"=SUM({c}{f0}:{c}{r - 1})", fB, ACC, bd=B_ALL)
        w(ws, f"I{r}", f"=ROUND(H{r}-{cx.NC[p]}{rtp},2)", fChk, ACC)
        ns.r = r + 2
    # ---- 10 other current liabilities
    ns.note("OCL", "Other current liabilities")
    ns.colhead()
    f0 = ns.r
    for m, _ in BM.note_lines("OCL", cx.custom):
        ns.master_line(m, ref="ocl_msme_int" if m == MI else None)
    if "ocl_msme_int" not in cx.refs:
        cx.setref("ocl_msme_int", f0)
    ns.total(f0, ns.r - 1, ref="ocl_total")
    ns.gap()
    # ---- 11 short-term provisions
    ns.note("STP", "Short-term provisions")
    ns.colhead()
    f0 = ns.r
    for m, _ in BM.note_lines("STP", cx.custom):
        ns.master_line(m, ref="stp_tax" if m.startswith("Provision for income tax") else None)
    if "stp_tax" not in cx.refs:
        cx.setref("stp_tax", f0)
        cx.no_stp_tax = True
    if cx.nonzero(G):
        ns.line(f'="Provision for gratuity (current portion) (refer Note "&{q("Gratuity and contingencies")}!$B${{{{note:EB}}}}&")"',
                lambda p, r: f"={q('Gratuity and contingencies')}!{cx.NC[p]}{{{{gr_cur}}}}", ref="grat_cur")
    else:
        cx.setref("grat_cur", f0)
        cx.no_grat = True
    ns.total(f0, ns.r - 1, ref="stp_total")
    cx.finish_print(ws, ns.r)
    return ns


# ====================================================================== Notes_Assets
def notes_assets(cx):
    I = cx.I
    ns = NoteSheet(cx, "Notes_Assets", True, prev_ref="Notes_PPE!$B$8")
    ws = ns.ws
    S = cx.SCALE
    T = cx.TAX
    # ---- 13 deferred tax
    DTA = "Deferred tax asset (net)"
    ns.note("DTA", "Deferred tax assets / (liabilities) (net)")
    ns.colhead()
    ro = ns.r
    ns.line("Opening balance of deferred tax asset / (liability) (net)", lambda p, r: f"={cx.NC[cx.PREV[p]]}{r + 2}", opening=False)
    ns.line("Add: Deferred tax credit / (expense) for the year / period", lambda p, r: f"=-'Notes_P&L'!{cx.NC[p]}{{{{pl_deftax}}}}", opening=False)

    def dclose(p, r):
        f = "=" + cx.master_val(p, DTA, qs(DTA), "BS", "Dr")
        if cx.nonzero("Deferred tax liability (net)"):
            f += "-" + cx.master_val(p, "Deferred tax liability (net)", qs("Deferred tax liability (net)"), "BS", "Cr")
        return f
    rc = ns.line("Closing balance of deferred tax asset / (liability) (net)", dclose, font=fB, bd=B_TB, ref="dta_close")
    for i, p in enumerate(cx.PK):
        w(ws, f"{CL(6 + cx.N + i)}{rc}", f"=ROUND({cx.NC[p]}{ro}+{cx.NC[p]}{ro + 1}-{cx.NC[p]}{rc},2)", fChk, ACC)
    ns.gap()
    ns.sub("Components of deferred tax asset / (liabilities)")
    rg = ns.r + 1
    ns.line(cx.text("dta_dep_label", "On account of difference between depreciation as per the Companies Act, 2013 and the Income-tax Act, 1961 and other "
                    "timing differences (Sections 43B(h) / 40(a)(ia))"), lambda p, r: f"={cx.NC[p]}{rc}-{cx.NC[p]}{r + 1}", label_font=Font(name=TN, size=10))
    ws[f"C{ns.r - 1}"].alignment = A_W
    ws.row_dimensions[ns.r - 1].height = 39
    if not getattr(cx, "no_grat", False):
        ns.line(cx.text("dta_grat_label", "On account of expenses disallowed under the Income-tax Act, 1961 (provision for gratuity - Section 43B / 40A(7))"),
                lambda p, r: f"=ROUND((Notes_Liabilities!{cx.NC[p]}{{{{ltp_grat}}}}+Notes_Liabilities!{cx.NC[p]}{{{{grat_cur}}}})*{T},2)")
        ws[f"C{ns.r - 1}"].alignment = A_W
        ws.row_dimensions[ns.r - 1].height = 27
    else:
        ns.line("On account of other timing differences", lambda p, r: 0)
    ns.total(rg - 1, ns.r - 1, "Net deferred tax asset / (liability)", bd=B_TB)
    ns.para(cx.text("dta_note", '="Deferred tax has been recognised on timing differences, being the difference between taxable income and accounting income that '
                    'originate in one period and are capable of reversal in one or more subsequent periods, at the rate of tax applicable under Section 115BAA of '
                    'the Income-tax Act, 1961 ("&TEXT(MasterData!$B$17,"0.000%")&"). The component on account of depreciation is derived as the balancing figure '
                    'of the deferred tax asset (net) as per the books and restatement adjustments."'), fI, last_col=cx.LASTN)
    ns.gap()
    # ---- 14 long-term loans and advances
    ll = [m for m, _ in BM.note_lines("LTLA", cx.custom) if cx.nonzero(m)]
    cx.has_ltla = bool(ll)
    if ll:
        ns.note("LTLA", "Long-term loans and advances")
        ns.colhead()
        ns.sub("Unsecured, considered good")
        f0 = ns.r
        for m in ll:
            ns.master_line(m)
        ns.total(f0, ns.r - 1, ref="ltla_total")
        ns.gap()
    # ---- 15 inventories
    ns.note("INV", "Inventories")
    ns.colhead()
    ns.sub("(Valued at lower of cost and net realisable value, as taken, valued and certified by the management)", fI)

    def stock(p, r):
        if p == "S":
            return f"={MD}!$B$19/{S}"
        if p == "Y1":
            return f"={MD}!$B$16/{S}"
        later = {"Y2": "Y1", "Y3": "Y2", "OPEN": "Y3"}[p]
        return "=" + cx.tb(later, qs("Opening stock"))
    rs_ = ns.line(cx.text("inv_label", "Stock-in-trade"), stock, ref="inv_line")
    ns.total(rs_, rs_, ref="inv_total")
    ns.para(cx.text("inv_note", f"Closing stock as at {long_date(cx.DATE['Y1'])} is as per the audited financial statements; closing stock of earlier years is the "
                    "opening stock carried in the next year's books."))
    ns.gap()
    # ---- 16 trade receivables
    ns.note("TR", "Trade receivables")
    ns.colhead()
    f0 = ns.r
    ns.master_line("Trade receivables", "Undisputed trade receivables - considered good", keep=True)
    for lbl in ("Undisputed trade receivables - considered doubtful", "Disputed trade receivables - considered good",
                "Disputed trade receivables - considered doubtful", "Less: Allowance for bad and doubtful debts"):
        ns.line(lbl, lambda p, r: 0)
    rtr = ns.total(f0, ns.r - 1, bd=B_TB, ref="tr_total")
    ns.gap()
    trn = cx.refs["note:TR"]
    letters = "abcdefg"
    for i, p in enumerate(cx.PK):
        r = ns.r
        w(ws, f"C{r}", f'=$B${trn}&"({letters[i]}) "&"Trade receivables ageing schedule as at "&TEXT({cx.MDC[p]},"mmmm dd, yyyy")', fB)
        r += 1
        w(ws, f"C{r}", "Particulars", fB, bd=B_ALL, al=A_CW)
        ws.merge_cells(f"C{r}:C{r + 1}")
        w(ws, f"D{r}", "Outstanding for following periods from due date of payment", fB, bd=B_ALL, al=A_CW)
        ws.merge_cells(f"D{r}:I{r}")
        for j, h in enumerate(["Less than 6 months", "6 months - 1 year", "1 - 2 years", "2 - 3 years", "More than 3 years", "Total"]):
            w(ws, f"{CL(4 + j)}{r + 1}", h, fB, bd=B_ALL if j < 5 else B_LR, al=A_CW)
        r += 2
        f1 = r
        base = cx.DEB[p]
        for k, lbl in enumerate(["Undisputed trade receivables - considered good", "Undisputed trade receivables - considered doubtful",
                                 "Disputed trade receivables - considered good", "Disputed trade receivables - considered doubtful"]):
            w(ws, f"C{r}", lbl, fN, bd=B_LR, al=A_LW)
            for j in range(5):
                w(ws, f"{CL(4 + j)}{r}", f"={q('Debtor Ageing')}!{CL(3 + j)}{base + k}/{S}", fN, ACC, bd=B_LR)
            w(ws, f"I{r}", f"=SUM(D{r}:H{r})", fN, ACC, bd=B_LR)
            r += 1
        w(ws, f"C{r}", "Less: Allowance for bad and doubtful debts", fN, bd=B_LR, al=A_LW)
        for j in range(5):
            w(ws, f"{CL(4 + j)}{r}", 0, fN, ACC, bd=B_LR)
        w(ws, f"I{r}", f"=SUM(D{r}:H{r})", fN, ACC, bd=B_LR)
        r += 1
        w(ws, f"C{r}", "Total", fB, bd=B_ALL)
        for j in range(6):
            c = CL(4 + j)
            w(ws, f"{c}{r}", f"=SUM({c}{f1}:{c}{r - 1})", fB, ACC, bd=B_ALL)
        w(ws, f"J{r}", f"=ROUND(I{r}-{cx.NC[p]}{rtr},2)", fChkI, ACC)
        ns.r = r + 2
    # ---- 17 cash and bank balances
    ns.note("CASH", "Cash and bank balances")
    ns.colhead()
    ns.sub("Cash and cash equivalents")
    f0 = ns.r
    cash_rows = []
    for m, _ in BM.note_lines("CASH", cx.custom):
        rr = ns.master_line(m, keep=m in ("Cash on hand", "Balances with banks in current accounts"))
        if rr:
            cash_rows.append(rr)
            cx.setref("cash:" + m, rr)
    rcce = ns.line("Total cash and cash equivalents", lambda p, r: "=" + "+".join(f"{cx.NC[p]}{x}" for x in cash_rows), font=fB, bd=B_T, ref="cce_total")
    cx.CASH_ROWS = cash_rows
    ns.sub("Other bank balances")
    fd_rows = []
    for m, _ in BM.note_lines("BANK", cx.custom):
        rr = ns.master_line(m, keep=m == "Fixed deposits with banks", ref="fd_line" if m == "Fixed deposits with banks" else None)
        if rr:
            fd_rows.append(rr)
    ns.line("Total", lambda p, r: f"={cx.NC[p]}{rcce}+" + "+".join(f"{cx.NC[p]}{x}" for x in fd_rows), font=fB, bd=B_TB2, ref="cash_total")
    ns.para(cx.text("cash_note", "Fixed deposits include deposits held as margin money / lien against bank guarantees and overdraft facilities."))
    ns.gap()
    # ---- 18 short-term loans and advances
    sl = [m for m, _ in BM.note_lines("STLA", cx.custom) if cx.nonzero(m)]
    ns.note("STLA", "Short-term loans and advances")
    ns.colhead()
    ns.sub("Unsecured, considered good")
    f0 = ns.r
    for m in sl:
        ns.master_line(m)
    if not sl:
        ns.line("Advances", lambda p, r: 0)
    ns.total(f0, ns.r - 1, ref="stla_total")
    ns.gap()
    # ---- 19 other current assets
    ns.note("OCA", "Other current assets")
    ns.colhead()
    f0 = ns.r
    for m, _ in BM.note_lines("OCA", cx.custom):
        key = {"Advance tax, TDS and TCS receivable": "oca_tax", "Interest accrued on deposits": "oca_int"}.get(m)
        rr = ns.master_line(m, ref=key)
        if key and not rr:
            cx.setref(key, None)
    for key in ("oca_tax", "oca_int"):
        if cx.refs.get(key) is None:
            cx.refs.pop(key, None)
    if ns.r == f0:
        ns.line("Other current assets", lambda p, r: 0)
    ns.total(f0, ns.r - 1, ref="oca_total")
    cx.finish_print(ws, ns.r)
    return ns


# ====================================================================== Notes_P&L
def notes_pl(cx):
    I = cx.I
    ns = NoteSheet(cx, "Notes_P&L", False, prev_ref=f"Notes_Assets!$B${{{{note:OCA}}}}")
    ws = ns.ws
    S = cx.SCALE

    def simple(key, title, lines, refkey, extra=None, keep=()):
        ns.note(key, title)
        ns.colhead("pl")
        f0 = ns.r
        if extra:
            extra()
        for m, _ in lines:
            rk = {"Interest on delayed payments to MSME": "pl_msme_int", "Gratuity": "pl_grat", "Directors' remuneration": "pl_dirrem",
                  "Contribution to provident and other funds": "pl_pf", "Payment to auditors": "pl_aud", "Current tax": "pl_curtax",
                  "Deferred tax": "pl_deftax", "Purchases of stock-in-trade": "pl_pur", "Miscellaneous expenses": "pl_misc"}.get(m)
            rr = ns.master_line(m, keep=m in keep)
            if rk:
                cx.setref(rk, rr if rr else f0)
                if not rr:
                    cx.missing = getattr(cx, "missing", set()) | {rk}
        if ns.r == f0:
            ns.line(lines[0][0] if lines else title, lambda p, r: 0)
        rt = ns.total(f0, ns.r - 1, ref=refkey)
        return f0, rt

    simple("REV", "Revenue from operations", BM.note_lines("REV", cx.custom), "pl_rev",
           extra=lambda: ns.sub("Sale of goods and services"))
    ns.gap()
    # ---- 21 other income (ledger-wise lines)
    ns.note("OI", "Other income")
    ns.colhead("pl")
    f0 = ns.r
    cx.OI_INT_ROWS, cx.OI_FD_ROW, cx.OI_PPE_ROWS = [], None, []
    for ln in cx.oi_lines:
        def f(p, r, ln=ln):
            parts = []
            for led in ln.get("ledgers") or []:
                s = q(cx.TBN[p])
                parts.append(f"SUMIFS({s}!$I:$I,{s}!$G:$G,{qs(ln['master'])},{s}!$A:$A,{qs(led)})")
            if not ln.get("ledgers"):
                parts.append(cx.tb(p, qs(ln["master"])))
            return "=" + "+".join(parts)
        rr = ns.line(ln["label"], f)
        if ln.get("interest"):
            cx.OI_INT_ROWS.append(rr)
        if ln.get("fd"):
            cx.OI_FD_ROW = rr
        if ln.get("ppe_sale"):
            cx.OI_PPE_ROWS.append(rr)
    # restatement entries on the Other income masters (the lines above are ledger-wise, from the TB)
    oi_adj = [m for m, _ in BM.note_lines("OI", cx.custom) if m in cx.adj_masters]
    if oi_adj:
        ns.line(cx.text("oi_adj_label", "Restatement adjustments"), lambda p, r: "=" + "+".join(cx.adj_pl(qs(m), p, "Cr") for m in oi_adj))
    if ns.r == f0:
        ns.line("Other income", lambda p, r: 0)
    rt = ns.total(f0, ns.r - 1, ref="pl_oi")
    for i, p in enumerate(cx.PK):        # check: ledger-wise lines agree with the Other income masters
        tot = "+".join(cx.tb(p, qs(m)) + ("+" + cx.adj_pl(qs(m), p, "Cr") if m in oi_adj else "") for m, _ in BM.note_lines("OI", cx.custom))
        w(ws, f"{CL(5 + cx.N + i + 1)}{rt}", f"=ROUND({cx.NC[p]}{rt}-({tot}),2)", fChk, ACC)
    ns.line("Restated profit before tax", lambda p, r: f"=PL_2!{cx.BC[p]}{{{{PL2_pbt}}}}", font=fI)
    ns.line("Other income as a percentage of restated profit before tax", lambda p, r: f"=IFERROR({cx.NC[p]}{rt}/{cx.NC[p]}{r - 1},0)", font=fI, fmt=PCT)
    ns.para(cx.text("oi_note", "Interest income on fixed deposits and income tax refund is recurring in nature and not related to the main business activity. "
                    "Other non-operating income (discount received, balances written back, exchange gain) is non-recurring in nature."))
    ns.gap()
    simple("PUR", "Purchases of stock-in-trade", BM.note_lines("PUR", cx.custom), "pl_purtot", keep=("Purchases of stock-in-trade",))
    ns.gap()
    # ---- 23 changes in inventories
    ns.note("CHINV", "Changes in inventories of stock-in-trade")
    ns.colhead("pl")
    ra = ns.line("Inventories at the beginning of the year (a)", lambda p, r: f"=Notes_Assets!{cx.NC[cx.PREV[p]]}{{{{inv_line}}}}")
    ns.line("Inventories at the end of the year (b)", lambda p, r: f"=Notes_Assets!{cx.NC[p]}{{{{inv_line}}}}")
    ns.line("(Increase) / decrease in inventories (a - b)", lambda p, r: f"={cx.NC[p]}{ra}-{cx.NC[p]}{ra + 1}", font=fB, bd=B_TB2, ref="pl_chinv")
    ns.gap()
    simple("EMP", "Employee benefits expense", BM.note_lines("EMP", cx.custom), "pl_emp", keep=("Salaries, wages and incentives",))
    ns.gap()
    simple("FIN", "Finance costs", BM.note_lines("FIN", cx.custom), "pl_fin")
    ns.gap()
    simple("DEP", "Depreciation and amortisation expense", BM.note_lines("DEP", cx.custom), "pl_dep", keep=("Depreciation and amortisation",))
    ns.gap()
    simple("OE", "Other expenses", BM.note_lines("OE", cx.custom), "pl_oe")
    if "pl_aud" not in getattr(cx, "missing", set()):
        ns.gap()
        ns.sub("Payment to auditors")
        ns.line("As statutory auditor (including tax audit)", lambda p, r: f"='Notes_P&L'!{cx.NC[p]}{{{{pl_aud}}}}", ref="pl_aud_sub")
    ns.gap()
    simple("TAX", "Tax expense", BM.note_lines("TAX", cx.custom), "pl_tax", keep=("Current tax", "Deferred tax"))
    ns.gap()
    ns.line("Net profit after tax as per books of account (memorandum)", lambda p, r: "=" + cx.tb(p, qs("Net profit as per books")), font=fI, ref="pl_bookpat")
    ns.gap()
    # ---- 29 earnings per share
    ns.note("EPS", "Earnings per share")
    ns.colhead("pl")
    ns.line("Restated profit for the year attributable to equity shareholders", lambda p, r: f"=PL_2!{cx.BC[p]}{{{{PL2_pat}}}}")
    ns.line("Weighted average number of equity shares (adjusted for bonus / split)", lambda p, r: f"='Accounting ratios'!{cx.NC[p]}{{{{ar_wans_adj}}}}",
            fmt=NUM0)
    ns.line("Basic and diluted earnings per share (Rs.)", lambda p, r: f"='Accounting ratios'!{cx.NC[p]}{{{{ar_eps}}}}", font=fB, fmt=EPS, ref="pl_eps")
    ns.line("Face value per equity share (Rs.)", lambda p, r: f"={MD}!$B$7", fmt=EPS)
    cx.finish_print(ws, ns.r)
    return ns
