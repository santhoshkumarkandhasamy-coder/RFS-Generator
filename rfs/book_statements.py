"""Annexures I - V: BS_1, PL_2, CFS_3, SAP_4, SOA_5."""
from openpyxl.styles import Font, Alignment
from .book_core import *          # noqa: F401,F403
from . import book_masters as BM
from .book_notes import head

FV = f"{MD}!$B$7"


def signature(cx, ws, r, cL="B", cM="E", cR="G"):
    """auditor (left) and four signatories of the Board (two rows of two)"""
    w(ws, f"{cL}{r}", "As per our report of even date attached")
    r += 2
    w(ws, f"{cL}{r}", f'="For "&{MD}!$B$20', fB)
    w(ws, f"{cM}{r}", "For and on behalf of the Board of Directors of", fB)
    r += 1
    w(ws, f"{cL}{r}", "Chartered Accountants")
    w(ws, f"{cM}{r}", f"={MD}!$B$3", fB, ACC)
    r += 1
    w(ws, f"{cL}{r}", f'="Firm Registration No.: "&{MD}!$B$21')
    w(ws, f"{cM}{r}", f'="CIN: "&{MD}!$B$2')
    r += 1
    w(ws, f"{cL}{r}", f'="Peer Review Certificate No.: "&{MD}!$B$22')
    r += 3
    for c in (cL, cM, cR):
        w(ws, f"{c}{r}", "Sd/-")
    r += 1
    w(ws, f"{cL}{r}", f'=IF({MD}!$B$23="","[Partner]",{MD}!$B$23)', fB, ACC)
    w(ws, f"{cM}{r}", f'=IF({MD}!$F$4="","[Director 1]",PROPER({MD}!$F$4))', fB, ACC)
    w(ws, f"{cR}{r}", f'=IF({MD}!$F$5="","[Director 2]",PROPER({MD}!$F$5))', fB, ACC)
    r += 1
    w(ws, f"{cL}{r}", "Partner")
    w(ws, f"{cM}{r}", f'=""&{MD}!$G$4', fN, ACC)
    w(ws, f"{cR}{r}", f'=""&{MD}!$G$5', fN, ACC)
    r += 1
    w(ws, f"{cL}{r}", f'="Membership No.: "&{MD}!$B$24')
    w(ws, f"{cM}{r}", f'="DIN: "&{MD}!$E$4')
    w(ws, f"{cR}{r}", f'="DIN: "&{MD}!$E$5')
    r += 1
    w(ws, f"{cL}{r}", f'="UDIN: "&{MD}!$B$26')
    r += 2
    w(ws, f"{cM}{r}", "Sd/-")
    w(ws, f"{cR}{r}", "Sd/-")
    r += 1
    w(ws, f"{cM}{r}", f'=IF({MD}!$F$6="","[CFO]",PROPER({MD}!$F$6))', fB, ACC)
    w(ws, f"{cR}{r}", f'=IF({MD}!$F$7="","[Company Secretary]",PROPER({MD}!$F$7))', fB, ACC)
    r += 1
    w(ws, f"{cM}{r}", f'=""&{MD}!$G$6', fN, ACC)
    w(ws, f"{cR}{r}", f'=""&{MD}!$G$7', fN, ACC)
    r += 1
    w(ws, f"{cM}{r}", f'="DIN: "&{MD}!$E$6')
    r += 2
    for c in (cL, cM, cR):
        w(ws, f"{c}{r}", f'="Place: "&{MD}!$B$27')
    r += 1
    for c in (cL, cM, cR):
        w(ws, f"{c}{r}", f'="Date: "&{MD}!$B$25')
    return r


def _row(cx, ws, r, label, f, keys, cols, font=fN, bd=None, note=None, lab_font=None, lab_bd=False):
    w(ws, f"B{r}", label, lab_font or font, bd=bd if lab_bd else None)
    if note:
        w(ws, f"C{r}", note, fN, al=A_C)
    for k in keys:
        v = f(k)
        if v is None:
            continue
        fo = font
        if k == "OPEN":
            fo = Font(name=TN, size=10, italic=True, color="808080")
        w(ws, f"{cols[k]}{r}", v, fo, ACC, bd=bd)


# ====================================================================== BS_1
def bs1(cx):
    BC = cx.BC
    last = BC["OPEN"]
    ws = cx.sheet("BS_1", {"A": 1.0, "B": 52.0, "C": 11.0, **{BC[p]: 13.0 for p in cx.PK}, cx.BGAP: 3.0, last: 15.0}, print_cols=last)
    head(cx, ws, "ANNEXURE I - RESTATED SUMMARY STATEMENT OF ASSETS AND LIABILITIES", last)
    keys = cx.PK + ["OPEN"]
    w(ws, "B8", "Particulars", fB, bd=B_TB)
    w(ws, "C8", "Annexure VI Notes", fB, bd=B_TB, al=A_CW)
    for k in keys:
        w(ws, f"{BC[k]}8", cx.dtxt(k), fB if k != "OPEN" else fGreyI, bd=B_TB, al=A_CW)
    ws.row_dimensions[8].height = 30
    NL, NA, NE = "Notes_Liabilities", "Notes_Assets", q("Notes_Equity & Reserves")
    nc = lambda k: cx.NC[k]
    r = 9
    rows = {}

    def put(name, label, f, note=None, font=fN, bd=None):
        nonlocal r
        _row(cx, ws, r, label, f, keys, BC, font, bd, note)
        rows[name] = r
        r += 1

    def hdr(t, gap_before=0):
        nonlocal r
        r += gap_before
        w(ws, f"B{r}", t, fB)
        r += 1

    hdr("EQUITY AND LIABILITIES")
    hdr("Shareholders' funds", 1)
    put("sc", "Share capital", lambda k: f"={NE}!{nc(k)}{{{{sc_total}}}}", f"={NE}!$B${{{{note:SC}}}}")
    put("rs", "Reserves and surplus", lambda k: f"={NE}!{nc(k)}{{{{rs_end}}}}", f"={NE}!$B${{{{note:RS}}}}")
    put("shf", "Total Shareholders' funds", lambda k: f"={BC[k]}{rows['sc']}+{BC[k]}{rows['rs']}", font=fB, bd=B_T)
    hdr("Non-current liabilities", 1)
    ncl = []
    put("ltb", "Long-term borrowings", lambda k: f"={NL}!{nc(k)}{{{{ltb_total}}}}", f"={NL}!$B${{{{note:LTB}}}}")
    ncl.append(rows["ltb"])
    if getattr(cx, "has_oltl", False):
        put("oltl", "Other long-term liabilities", lambda k: f"={NL}!{nc(k)}{{{{oltl_total}}}}", f"={NL}!$B${{{{note:OLTL}}}}")
        ncl.append(rows["oltl"])
    if getattr(cx, "has_ltp", False):
        put("ltp", "Long-term provisions", lambda k: f"={NL}!{nc(k)}{{{{ltp_total}}}}", f"={NL}!$B${{{{note:LTP}}}}")
        ncl.append(rows["ltp"])
    put("ncl", "Total Non-current liabilities", lambda k: f"=SUM({BC[k]}{ncl[0]}:{BC[k]}{ncl[-1]})", font=fB, bd=B_T)
    hdr("Current liabilities", 1)
    put("stb", "Short-term borrowings", lambda k: f"={NL}!{nc(k)}{{{{stb_total}}}}", f"={NL}!$B${{{{note:STB}}}}")
    w(ws, f"B{r}", "Trade payables")
    w(ws, f"C{r}", f"={NL}!$B${{{{note:TP}}}}", al=A_C)
    r += 1
    put("tpm", "- Total outstanding dues of micro and small enterprises", lambda k: f"={NL}!{nc(k)}{{{{tp_msme}}}}")
    put("tpo", "- Total outstanding dues of creditors other than micro and small enterprises", lambda k: f"={NL}!{nc(k)}{{{{tp_total}}}}-{NL}!{nc(k)}{{{{tp_msme}}}}"
        if len(BM.note_lines("TP", cx.custom)) > 2 else f"={NL}!{nc(k)}{{{{tp_oth}}}}")
    put("ocl", "Other current liabilities", lambda k: f"={NL}!{nc(k)}{{{{ocl_total}}}}", f"={NL}!$B${{{{note:OCL}}}}")
    put("stp", "Short-term provisions", lambda k: f"={NL}!{nc(k)}{{{{stp_total}}}}", f"={NL}!$B${{{{note:STP}}}}")
    cl_ = [rows[x] for x in ("stb", "tpm", "tpo", "ocl", "stp")]
    put("cl", "Total Current liabilities", lambda k: "=SUM(" + ",".join(f"{BC[k]}{x}" for x in cl_) + ")", font=fB, bd=B_T)
    r += 1
    put("tel", "Total Equity and Liabilities", lambda k: f"=SUM({BC[k]}{rows['shf']},{BC[k]}{rows['ncl']},{BC[k]}{rows['cl']})", font=fB, bd=B_TB2)
    hdr("ASSETS", 1)
    hdr("Non-current assets", 1)
    w(ws, f"B{r}", "Property, Plant and Equipment and Intangible assets")
    w(ws, f"C{r}", "=Notes_PPE!$B$8", al=A_C)
    r += 1
    nca = []

    def ppe(k, grp):
        if getattr(cx, "FAN", None):
            kk = k if k != "OPEN" else cx.EARLY
            col = "M" if k != "OPEN" else "L"
            rr = cx.PPE_ROWS.get((kk, grp))
            return f"=Notes_PPE!{col}{rr}" if rr else 0
        rr = cx.PPE_ROWS.get(("FB", grp))
        return f"=Notes_PPE!{cx.FALLBACK_COLS[k]}{rr}"
    put("ppe", "   Property, Plant and Equipment", lambda k: ppe(k, "PPE"))
    nca.append(rows["ppe"])
    if any(g == "INTANG" for _, g in _classes(cx)) or not getattr(cx, "FAN", None):
        put("int", "   Intangible assets", lambda k: ppe(k, "INTANG"))
        nca.append(rows["int"])

    def cw(k, m):
        kk = k if k != "OPEN" else cx.EARLY
        rr = cx.CWIP_ROWS[(kk, m)]
        return f"=Notes_PPE!{'G' if k != 'OPEN' else 'D'}{rr}"
    for m, lbl in (("Capital work-in-progress", "   Capital work-in-progress"), ("Intangible assets under development", "   Intangible assets under development")):
        if cx.nonzero(m) or m == "Capital work-in-progress":
            put(m, lbl, lambda k, m=m: cw(k, m))
            nca.append(rows[m])
    put("dta", "Deferred tax assets (Net)", lambda k: f"={NA}!{nc(k)}{{{{dta_close}}}}", f"={NA}!$B${{{{note:DTA}}}}")
    nca.append(rows["dta"])
    if getattr(cx, "has_ltla", False):
        put("ltla", "Long-term loans and advances", lambda k: f"={NA}!{nc(k)}{{{{ltla_total}}}}", f"={NA}!$B${{{{note:LTLA}}}}")
        nca.append(rows["ltla"])
    put("nca", "Total Non-current assets", lambda k: "=SUM(" + ",".join(f"{BC[k]}{x}" for x in nca) + ")", font=fB, bd=B_T)
    hdr("Current assets", 1)
    put("inv", "Inventories", lambda k: f"={NA}!{nc(k)}{{{{inv_total}}}}", f"={NA}!$B${{{{note:INV}}}}")
    put("tr", "Trade receivables", lambda k: f"={NA}!{nc(k)}{{{{tr_total}}}}", f"={NA}!$B${{{{note:TR}}}}")
    put("cash", "Cash and bank balances", lambda k: f"={NA}!{nc(k)}{{{{cash_total}}}}", f"={NA}!$B${{{{note:CASH}}}}")
    put("stla", "Short-term loans and advances", lambda k: f"={NA}!{nc(k)}{{{{stla_total}}}}", f"={NA}!$B${{{{note:STLA}}}}")
    put("oca", "Other current assets", lambda k: f"={NA}!{nc(k)}{{{{oca_total}}}}", f"={NA}!$B${{{{note:OCA}}}}")
    ca = [rows[x] for x in ("inv", "tr", "cash", "stla", "oca")]
    put("ca", "Total Current assets", lambda k: "=SUM(" + ",".join(f"{BC[k]}{x}" for x in ca) + ")", font=fB, bd=B_T)
    r += 1
    put("ta", "Total Assets", lambda k: f"={BC[k]}{rows['nca']}+{BC[k]}{rows['ca']}", font=fB, bd=B_TB2)
    for k in keys:
        w(ws, f"{BC[k]}1", f"=ROUND({BC[k]}{rows['tel']}-{BC[k]}{rows['ta']},2)", fChk, ACC)
    r += 1
    t = cx.text("bs_readwith", "The above statement should be read with Annexure IV - Background, basis of preparation and significant accounting policies, "
                               "Annexure V - Statement of adjustments to audited financial statements and Annexure VI - Notes to the restated summary statements.")
    w(ws, f"B{r}", t, fN, al=A_J)
    ws.merge_cells(f"B{r}:{cx.LASTP}{r + 1}")
    cx.BS_READWITH = r
    r += 3
    r = signature(cx, ws, r, "B", BC[cx.PK[1]], BC[cx.PK[-1]])
    cx.BS = rows
    for k, v in rows.items():
        cx.setref("bs_" + k, v)
    cx.finish_print(ws, r)


def _classes(cx):
    from .book_fa import fa_classes
    return fa_classes(cx) if getattr(cx, "FAN", None) else []


# ====================================================================== PL_2
def pl2(cx):
    BC = cx.BC
    ws = cx.sheet("PL_2", {"A": 1.0, "B": 60.0, "C": 11.0, **{BC[p]: 13.0 for p in cx.PK}}, print_cols=cx.LASTP)
    head(cx, ws, "ANNEXURE II - RESTATED SUMMARY STATEMENT OF PROFIT AND LOSS", cx.LASTP)
    keys = cx.PK
    w(ws, "B7", "Particulars", fB, bd=B_TB)
    w(ws, "C7", "Annexure VI Notes", fB, bd=B_TB, al=A_CW)
    for k in keys:
        w(ws, f"{BC[k]}7", cx.dtxt(k, "pl"), fB, bd=B_TB, al=A_CW)
    ws.row_dimensions[7].height = 39.75
    NP = "'Notes_P&L'"
    nc = lambda k: cx.NC[k]
    rows = {}
    r = 8

    def put(name, label, f, note=None, font=fN, bd=None, lab_font=None):
        nonlocal r
        _row(cx, ws, r, label, f, keys, BC, font, bd, note, lab_font, lab_bd=True)
        rows[name] = r
        r += 1
    w(ws, f"B{r}", "Income", fB)
    r += 1
    put("rev", "I. Revenue from operations", lambda k: f"={NP}!{nc(k)}{{{{pl_rev}}}}", f"={NP}!$B${{{{note:REV}}}}")
    put("oi", "II. Other income", lambda k: f"={NP}!{nc(k)}{{{{pl_oi}}}}", f"={NP}!$B${{{{note:OI}}}}")
    put("ti", "III. Total Income (I + II)", lambda k: f"=SUM({BC[k]}{rows['rev']}:{BC[k]}{rows['oi']})", font=fB, bd=B_T)
    w(ws, f"B{r}", "IV. Expenses", fB)
    r += 1
    e0 = r
    put("pur", "Purchases of stock-in-trade", lambda k: f"={NP}!{nc(k)}{{{{pl_purtot}}}}", f"={NP}!$B${{{{note:PUR}}}}")
    put("chinv", "Changes in inventories of stock-in-trade", lambda k: f"={NP}!{nc(k)}{{{{pl_chinv}}}}", f"={NP}!$B${{{{note:CHINV}}}}")
    put("emp", "Employee benefits expense", lambda k: f"={NP}!{nc(k)}{{{{pl_emp}}}}", f"={NP}!$B${{{{note:EMP}}}}")
    put("fin", "Finance costs", lambda k: f"={NP}!{nc(k)}{{{{pl_fin}}}}", f"={NP}!$B${{{{note:FIN}}}}")
    put("dep", "Depreciation and amortisation expense", lambda k: f"={NP}!{nc(k)}{{{{pl_dep}}}}", f"={NP}!$B${{{{note:DEP}}}}")
    put("oe", "Other expenses", lambda k: f"={NP}!{nc(k)}{{{{pl_oe}}}}", f"={NP}!$B${{{{note:OE}}}}")
    put("te", "Total Expenses", lambda k: f"=SUM({BC[k]}{e0}:{BC[k]}{r - 1})", font=fB, bd=B_T)
    put("pbt", "V. Profit before tax (III - IV)", lambda k: f"={BC[k]}{rows['ti']}-{BC[k]}{rows['te']}", font=fB, bd=B_T)
    ws[f"B{rows['pbt']}"].border = B_T
    w(ws, f"B{r}", "VI. Tax expense:", fB)
    r += 1
    put("ct", "Current tax", lambda k: f"={NP}!{nc(k)}{{{{pl_curtax}}}}", f"={NP}!$B${{{{note:TAX}}}}")
    put("dt", "Deferred tax", lambda k: f"={NP}!{nc(k)}{{{{pl_deftax}}}}")
    put("tt", "Total tax expense", lambda k: f"=SUM({BC[k]}{rows['ct']}:{BC[k]}{rows['dt']})", bd=B_T)
    put("pat", "VII. Profit for the year / period (V - VI)", lambda k: f"={BC[k]}{rows['pbt']}-{BC[k]}{rows['tt']}", font=fB, bd=B_TB2)
    r += 1
    w(ws, f"B{r}", f'="VIII. Earnings per equity share (Face value of Rs. "&{FV}&" each)"', fB)
    r += 1
    _row(cx, ws, r, "Basic and Diluted (in Rs.)", lambda k: f"={NP}!{nc(k)}{{{{pl_eps}}}}", keys, BC, fN, None, f"={NP}!$B${{{{note:EPS}}}}")
    for k in keys:
        ws[f"{BC[k]}{r}"].number_format = EPS
    rows["eps"] = r
    r += 1
    if cx.stub:
        w(ws, f"B{r}", f"Earnings per share for the period ended {long_date(cx.stub)} are not annualised.", fI)
        r += 1
    w(ws, f"B{r}", f"=BS_1!B{cx.BS_READWITH}", fN, al=A_J)
    ws.merge_cells(f"B{r}:{cx.LASTP}{r + 1}")
    r += 3
    # row 1 check: restated PAT = book PAT + P&L restatement entries
    for k in keys:
        s, e = cx.MDC[cx.PREV[k]], cx.MDC[k]
        rng = cx._rng
        crit = f'{rng("D")},"P&L",{rng("A")},">"&{s},{rng("A")},"<="&{e}'
        w(ws, f"{BC[k]}1", f"=ROUND({BC[k]}{rows['pat']}-{NP}!{nc(k)}{{{{pl_bookpat}}}}+(SUMIFS({rng('E')},{crit})-SUMIFS({rng('F')},{crit})),2)", fChk, ACC)
    r = signature(cx, ws, r, "B", BC[cx.PK[1]], BC[cx.PK[-1]])
    cx.PL = rows
    for k, v in rows.items():
        cx.setref("PL2_" + k, v)
    cx.finish_print(ws, r)


# ====================================================================== CFS_3
def cfs3(cx):
    CC = cx.CC
    keys = cx.PK
    last = CC[keys[-1]]
    ws = cx.sheet("CFS_3", {"A": 1.0, "B": 62.0, **{CC[p]: 13.0 for p in keys}}, print_cols=last)
    head(cx, ws, "ANNEXURE III - RESTATED SUMMARY STATEMENT OF CASH FLOWS", last)
    w(ws, "B8", "Particulars", fB, bd=B_TB)
    for k in keys:
        w(ws, f"{CC[k]}8", cx.dtxt(k, "pl"), fB, bd=B_TB, al=A_CW)
    ws.row_dimensions[8].height = 39.75
    B = cx.BS
    BC = cx.BC
    NA, NL, NP = "Notes_Assets", "Notes_Liabilities", "'Notes_P&L'"
    nc = lambda k: cx.NC[k]
    pv = lambda k: cx.PREV[k]
    bs = lambda k, name: f"BS_1!{BC[k]}{B[name]}" if name in B else "0"
    rows = {}
    r = 9

    def put(name, label, f, font=fN, bd=None):
        nonlocal r
        _row(cx, ws, r, label, f, keys, CC, font, bd, lab_bd=True)
        rows[name] = r
        r += 1

    def hdr(t):
        nonlocal r
        w(ws, f"B{r}", t, fB)
        r += 1

    int_rows = cx.OI_INT_ROWS
    intinc = lambda k: "(" + "+".join(f"{NP}!{nc(k)}{x}" for x in int_rows) + ")" if int_rows else "0"
    oca_x = [x for x in ("oca_tax", "oca_int") if x in cx.refs]
    ppe_sale = lambda k: ("-(" + "+".join(f"{NP}!{nc(k)}{x}" for x in cx.OI_PPE_ROWS) + ")") if cx.OI_PPE_ROWS else "0"
    loss_led = (cx.I.get("cfs") or {}).get("ppe_loss_ledgers") or []
    fa_c = (lambda k: f"-'FA Restatement'!{cx.FBC[k]}{cx.FBR['c']}") if getattr(cx, "FBR", None) else (lambda k: "")
    ppe_loss = lambda k: (("+(" + "+".join(cx.tbl(k, l) for l in loss_led) + ")") if loss_led else "") + fa_c(k)   # restated loss = books - (c)

    hdr("A) Cash flows from operating activities")
    put("pbt", "Net profit before tax", lambda k: f"=PL_2!{BC[k]}{cx.PL['pbt']}")
    w(ws, f"B{r}", "Adjustments for:")
    r += 1
    a0 = rows["pbt"]
    put("dep", "Depreciation and amortisation expense", lambda k: f"=PL_2!{BC[k]}{cx.PL['dep']}")
    put("fin", "Finance costs", lambda k: f"=PL_2!{BC[k]}{cx.PL['fin']}")
    put("int", "Interest income", lambda k: f"=-{intinc(k)}")
    put("ppe", "(Profit) / loss on sale or write off of property, plant and equipment, and adjustments made directly in reserves",
        lambda k: f"={ppe_sale(k)}{ppe_loss(k)}")
    put("op", "Operating profit before working capital changes", lambda k: f"=SUM({CC[k]}{a0}:{CC[k]}{r - 1})", font=fB, bd=B_T)
    w(ws, f"B{r}", "Adjustments for changes in working capital:")
    r += 1
    w0 = r
    put("winv", "(Increase) / decrease in inventories", lambda k: f"={bs(pv(k), 'inv')}-{bs(k, 'inv')}")
    put("wtr", "(Increase) / decrease in trade receivables", lambda k: f"={bs(pv(k), 'tr')}-{bs(k, 'tr')}")
    put("wst", "(Increase) / decrease in short-term loans and advances", lambda k: f"={bs(pv(k), 'stla')}-{bs(k, 'stla')}")
    if "ltla" in B:
        put("wlt", "(Increase) / decrease in long-term loans and advances", lambda k: f"={bs(pv(k), 'ltla')}-{bs(k, 'ltla')}")
    ocx = lambda k: "".join(f"-{NA}!{nc(k)}{{{{{x}}}}}" for x in oca_x)
    put("woca", "(Increase) / decrease in other current assets", lambda k: f"=({bs(pv(k), 'oca')}{ocx(pv(k))})-({bs(k, 'oca')}{ocx(k)})")
    put("wtp", "Increase / (decrease) in trade payables", lambda k: f"=({bs(k, 'tpm')}+{bs(k, 'tpo')})-({bs(pv(k), 'tpm')}+{bs(pv(k), 'tpo')})")
    put("wocl", "Increase / (decrease) in other current liabilities", lambda k: f"={bs(k, 'ocl')}-{bs(pv(k), 'ocl')}")
    gcur = (lambda k: f"+{NL}!{nc(k)}{{{{grat_cur}}}}") if not getattr(cx, "no_grat", False) else (lambda k: "")
    prov = lambda k: "+".join(x for x in (bs(k, "ltp") if "ltp" in B else "", bs(k, "oltl") if "oltl" in B else "") if x) or "0"
    put("wprov", "Increase / (decrease) in provisions and other long-term liabilities",
        lambda k: f"=({prov(k)}{gcur(k)})-({prov(pv(k))}{gcur(pv(k))})")
    put("cgo", "Cash generated from operations", lambda k: f"=SUM({CC[k]}{rows['op']}:{CC[k]}{r - 1})", font=fB, bd=B_T)
    stp = (lambda k: f"+({NL}!{nc(k)}{{{{stp_tax}}}}-{NL}!{nc(pv(k))}{{{{stp_tax}}}})") if not getattr(cx, "no_stp_tax", False) else (lambda k: "")
    advt = (lambda k: f"-({NA}!{nc(k)}{{{{oca_tax}}}}-{NA}!{nc(pv(k))}{{{{oca_tax}}}})") if "oca_tax" in cx.refs else (lambda k: "")
    put("tax", "Income taxes paid (net of refunds)", lambda k: f"=-PL_2!{BC[k]}{cx.PL['tt']}{stp(k)}{advt(k)}-({bs(k, 'dta')}-{bs(pv(k), 'dta')})")
    put("A", "Net cash flow from / (used in) operating activities (A)", lambda k: f"={CC[k]}{rows['cgo']}+{CC[k]}{rows['tax']}", font=fB, bd=B_TB2)
    r += 1
    hdr("B) Cash flows from investing activities")
    b0 = r
    fa = [x for x in ("ppe", "int", "Capital work-in-progress", "Intangible assets under development") if x in B]
    fasum = lambda k: "(" + "+".join(bs(k, x) for x in fa) + ")"
    put("capex", "Purchase of property, plant and equipment and intangible assets (net of sale proceeds)",
        lambda k: f"=-({fasum(k)}-{fasum(pv(k))}+PL_2!{BC[k]}{cx.PL['dep']})-{CC[k]}{rows['ppe']}")
    if "fd_line" in cx.refs:
        put("fd", "(Investment in) / redemption of fixed deposits with banks", lambda k: f"={NA}!{nc(pv(k))}{{{{fd_line}}}}-{NA}!{nc(k)}{{{{fd_line}}}}")
    oint = (lambda k: f"-({NA}!{nc(k)}{{{{oca_int}}}}-{NA}!{nc(pv(k))}{{{{oca_int}}}})") if "oca_int" in cx.refs else (lambda k: "")
    put("intr", "Interest received", lambda k: f"={intinc(k)}{oint(k)}")
    put("B", "Net cash flow from / (used in) investing activities (B)", lambda k: f"=SUM({CC[k]}{b0}:{CC[k]}{r - 1})", font=fB, bd=B_TB2)
    r += 1
    hdr("C) Cash flows from financing activities")
    c0 = r
    shc = any(abs(cx.tbval.get((k, "Equity share capital"), 0) - cx.tbval.get((cx.PREV[k], "Equity share capital"), 0)) > 0.000001 for k in keys)
    if shc:
        put("sh", "Proceeds from issue of equity shares", lambda k: f"={bs(k, 'sc')}-{bs(pv(k), 'sc')}")
    put("ltb", "Proceeds from / (repayment of) long-term borrowings (net)", lambda k: f"={bs(k, 'ltb')}-{bs(pv(k), 'ltb')}")
    put("stb", "Proceeds from / (repayment of) short-term borrowings (net)", lambda k: f"={bs(k, 'stb')}-{bs(pv(k), 'stb')}")
    put("finp", "Finance costs paid", lambda k: f"=-PL_2!{BC[k]}{cx.PL['fin']}")
    put("C", "Net cash flow from / (used in) financing activities (C)", lambda k: f"=SUM({CC[k]}{c0}:{CC[k]}{r - 1})", font=fB, bd=B_TB2)
    r += 1
    put("net", "Net increase / (decrease) in cash and cash equivalents (A+B+C)", lambda k: f"={CC[k]}{rows['A']}+{CC[k]}{rows['B']}+{CC[k]}{rows['C']}", font=fB)
    put("open", "Cash and cash equivalents at the beginning of the year", lambda k: f"={NA}!{nc(pv(k))}{{{{cce_total}}}}")
    put("close", "Cash and cash equivalents at the end of the year", lambda k: f"={CC[k]}{rows['net']}+{CC[k]}{rows['open']}", font=fB, bd=B_TB2)
    r += 1
    hdr("Components of cash and cash equivalents:")
    k0 = r
    for m in ("Cash on hand", "Balances with banks in current accounts"):
        if "cash:" + m in cx.refs:
            put("c:" + m, m, lambda k, m=m: f"={NA}!{nc(k)}{{{{cash:{m}}}}}")
    put("comp", "Total cash and cash equivalents", lambda k: f"=SUM({CC[k]}{k0}:{CC[k]}{r - 1})", font=fB, bd=B_TB2)
    for k in keys:
        w(ws, f"{CC[k]}1", f"=ROUND({CC[k]}{rows['comp']}-{CC[k]}{rows['close']},2)", fChk, ACC)
    r += 1
    t = cx.text("cfs_note", "Note: The Cash Flow Statement has been prepared under the indirect method as set out in Accounting Standard 3 'Cash Flow Statements'. "
                            "Bank overdrafts are shown under short-term borrowings and their movement is included in financing activities.")
    w(ws, f"B{r}", t, fN, al=A_J)
    ws.merge_cells(f"B{r}:{last}{r + 1}")
    r += 3
    r = signature(cx, ws, r, "B", CC[keys[1]], CC[keys[-1]])
    cx.CFS = rows
    cx.finish_print(ws, r)


# ====================================================================== SAP_4
POLICIES = [
    ("pol_basis", "a) Basis of preparation\nThe Restated Summary Statement of Assets and Liabilities of the Company as at {DATES} and the related Restated Summary "
     "Statements of Profit and Loss and Cash Flows for the {PERIODS} then ended, together with the annexures thereto (collectively, the \"Restated Summary "
     "Statements\"), have been prepared specifically for inclusion in the Draft Red Herring Prospectus / Red Herring Prospectus / Prospectus to be filed by "
     "the Company with the {EXCHANGE} in connection with its proposed Initial Public Offer. The Restated Summary Statements have been prepared to comply with "
     "Section 26 of Part I of Chapter III of the Companies Act, 2013, the Securities and Exchange Board of India (Issue of Capital and Disclosure Requirements) "
     "Regulations, 2018, as amended, and the Guidance Note on Reports in Company Prospectuses (Revised 2019) issued by the Institute of Chartered Accountants "
     "of India.", 126.0),
    ("pol_compiled", "The Restated Summary Statements have been compiled from the audited financial statements of the Company for the {PERIODS_LONG}, "
     "prepared under the historical cost convention on accrual basis in accordance with the Accounting Standards specified under Section 133 of the Companies "
     "Act, 2013 read with the Companies (Accounting Standards) Rules, 2021, and are presented in the format prescribed in Division I of Schedule III to the Act.",
     55.5),
    ("pol_cycle", "Operating Cycle:\nAll assets and liabilities have been classified as current or non-current as per the Company's normal operating cycle and "
     "other criteria set out in Schedule III. Based on the nature of products and the time between acquisition of assets and their realisation in cash, the "
     "Company has ascertained its operating cycle as twelve months.", 55.5),
    ("pol_estimates", "b) Use of estimates\nThe preparation of the financial statements in conformity with Indian GAAP requires the management to make estimates "
     "and assumptions that affect the reported amounts of assets and liabilities, disclosure of contingent liabilities and the reported income and expenses "
     "during the year. Differences between actual results and estimates are recognised in the period in which the results are known / materialise.", 69.75),
    ("pol_revenue", "c) Revenue recognition\nRevenue from sale of goods is recognised when the significant risks and rewards of ownership are transferred to "
     "the buyer, which generally coincides with delivery. Revenue from services is recognised as the services are rendered, on a proportionate basis over "
     "the period of the contract. Services rendered but not billed at the year end are shown as unbilled revenue. Revenue is stated net of GST, trade "
     "discounts and returns. Interest income is recognised on a time proportion basis.", 84.0),
    ("pol_cash", "d) Cash and bank balances\nCash and cash equivalents comprise cash in hand and balances in current accounts with banks. Fixed deposits with "
     "banks, including those held as margin money, are disclosed as other bank balances.", 42.0),
    ("pol_cfs", "e) Cash flow statement\nCash flows are reported using the indirect method, whereby profit before tax is adjusted for the effects of "
     "transactions of a non-cash nature and any deferrals or accruals of past or future cash receipts or payments.", 42.0),
    ("pol_ppe", "f) Property, Plant and Equipment and Intangible assets\nProperty, Plant and Equipment are stated at cost less accumulated depreciation. Cost "
     "comprises the purchase price and any attributable cost of bringing the asset to its working condition for its intended use. Intangible assets are "
     "stated at cost less accumulated amortisation. Expenditure on development of software meeting the recognition criteria of AS 26 is carried as "
     "intangible assets under development until the asset is ready for use.", 69.75),
    ("pol_dep", "g) Depreciation\nDepreciation on Property, Plant and Equipment is provided on the written down value method based on the useful lives "
     "prescribed in Schedule II to the Companies Act, 2013, retaining a residual value of 5% of the original cost, as per the Fixed Asset Register maintained "
     "by the Company. Land is not depreciated. Computer software is amortised over its estimated useful life of three years.", 55.5),
    ("pol_inv", "h) Inventories\nStock-in-trade is valued at lower of cost and net realisable value. Cost includes purchase price and other costs incurred in "
     "bringing the inventories to their present location and condition.", 42.0),
    ("pol_emp", "i) Employee benefits\nShort term employee benefits are recognised as an expense in the year in which the related service is rendered. "
     "Contributions to provident fund and ESI are charged to the Statement of Profit and Loss. Liability for gratuity, a defined benefit obligation, is "
     "provided on the basis of valuation as at the balance sheet date.", 55.5),
    ("pol_fx", "j) Foreign currency transactions\nTransactions in foreign currency are recorded at the exchange rate prevailing on the date of the transaction. "
     "Monetary items outstanding at the year end are restated at the closing rate and the resultant differences are recognised in the Statement of Profit "
     "and Loss.", 55.5),
    ("pol_tax", "k) Taxes on income\nCurrent tax is determined as the amount of tax payable in respect of taxable income for the year under the Income-tax "
     "Act, 1961; the Company has opted for the concessional regime under Section 115BAA. Deferred tax is recognised on timing differences, subject to "
     "consideration of prudence, using tax rates enacted or substantively enacted at the balance sheet date.", 55.5),
    ("pol_lease", "l) Leases\nLease rentals in respect of assets taken on operating lease are charged to the Statement of Profit and Loss over the lease term.",
     42.0),
    ("pol_borrow", "m) Borrowing costs\nBorrowing costs attributable to the acquisition of qualifying assets are capitalised as part of the cost of such "
     "assets. All other borrowing costs are charged to the Statement of Profit and Loss.", 42.0),
    ("pol_prov", "n) Provisions and contingencies\nA provision is recognised when the Company has a present obligation as a result of a past event, it is "
     "probable that an outflow of resources will be required to settle the obligation and a reliable estimate can be made. Contingent liabilities are "
     "disclosed in the notes. Contingent assets are not recognised.", 55.5),
    ("pol_eps", "o) Earnings per share\nBasic earnings per share is computed by dividing the restated profit after tax by the weighted average number of equity "
     "shares outstanding during the year, adjusted for bonus issue / share split, if any, including those occurring after the balance sheet date. There are "
     "no dilutive potential equity shares.", 55.5),
    ("pol_segment", "p) Segment reporting\nThe Company is engaged in a single business segment and operates primarily in India. Accordingly, no separate "
     "segment information is disclosed.", 42.0),
    ("pol_impair", "q) Impairment of assets\nThe carrying amounts of assets are reviewed at each balance sheet date for any indication of impairment based on "
     "internal / external factors. An impairment loss is recognised wherever the carrying amount of an asset exceeds its recoverable amount, being the higher "
     "of the net selling price and value in use.", 54.0),
]

USEFUL_LIVES = [("Buildings", 30, 0.095), ("Plant and machinery", 15, 0.1810), ("Furniture and fittings", 10, 0.2589), ("Electrical installations", 10, 0.2589),
                ("Office equipment", 5, 0.4507), ("Office equipment - air conditioners", 10, 0.2589), ("Office equipment - mobile phones", 3, 0.6316),
                ("Computers", 3, 0.6316), ("Vehicles - motor cars", 8, 0.3123), ("Vehicles - two wheelers", 10, 0.2589), ("Computer software", 3, 0.3333),
                ("Land", "Not depreciated", "-")]


def policy_text(cx, key, t):
    dates = cx.dates_text()
    if cx.stub:
        per_long = f"period ended {long_date(cx.stub)} and the years ended " + cx.dates_text(["Y1", "Y2", "Y3"])
        periods = "period / years"
    else:
        per_long = "years ended " + cx.dates_text(["Y1", "Y2", "Y3"])
        periods = "years"
    t = cx.text(key, t)
    return (t.replace("{DATES}", dates).replace("{PERIODS_LONG}", per_long).replace("{PERIODS}", periods)
            .replace("{EXCHANGE}", (cx.I.get("company") or {}).get("exchange") or "SME Platform of the Stock Exchange"))


def sap4(cx):
    ws = cx.sheet("SAP_4", {"A": 1.0, "B": 5.0, "C": 110.0, "D": 16.0, "E": 13.0}, print_cols="H")
    head(cx, ws, "ANNEXURE IV - BACKGROUND, BASIS OF PREPARATION AND SIGNIFICANT ACCOUNTING POLICIES", "H", units=False)
    C = cx.I.get("company") or {}
    w(ws, "B8", 1, fB, al=A_L)
    w(ws, "C8", "BACKGROUND", fB)
    inc = C.get("incorporation_date")
    bg = (f'M/s. {C.get("name_proper") or (C.get("name") or "").title()} ("the Company") was incorporated on '
          f'{long_date(ddate(inc)) if inc else "[date of incorporation]"} as a {C.get("incorporated_as") or "private limited company"} under the Companies Act, '
          f'2013 with the Registrar of Companies, {C.get("roc") or "[ROC]"}. The registered office of the Company is at {C.get("address") or "[address]"}.')
    t = cx.text("sap_background", bg)
    w(ws, "C10", t, fN, al=A_J)
    ws.row_dimensions[10].height = est_height(t, 110)
    t = cx.text("sap_objects", "Main objects of the Company:\n" + (C.get("objects") or "[Main objects / business of the Company]"))
    w(ws, "C11", t, fN, al=A_J)
    ws.row_dimensions[11].height = est_height(t, 110)
    w(ws, "B12", "=B8+1", fB, al=A_L)
    w(ws, "C12", "SIGNIFICANT ACCOUNTING POLICIES", fB)
    r = 13
    pols = list(POLICIES)
    for key, t, h in pols:
        if key == "pol_impair":
            pass
        tt = policy_text(cx, key, t)
        w(ws, f"C{r}", tt, fN, al=A_J)
        ws.row_dimensions[r].height = max(h, est_height(tt, 110))
        r += 1
    for k, v in sorted((cx.I.get("texts") or {}).items()):
        if k.startswith("pol_extra"):
            w(ws, f"C{r}", v, fN, al=A_J)
            ws.row_dimensions[r].height = est_height(v, 110)
            r += 1
    r += 1
    w(ws, f"C{r}", "r) Useful lives and rates of depreciation (written down value method) as per the Fixed Asset Register", fB)
    r += 1
    for col, h in zip("CDE", ["Class of asset", "Useful life (years)", "Rate (WDV)"]):
        w(ws, f"{col}{r}", h, fB, bd=B_ALL, al=A_CW)
    r += 1
    for cl_, life, rate in (cx.I.get("useful_lives") or USEFUL_LIVES):
        w(ws, f"C{r}", cl_, fN, bd=B_ALL)
        w(ws, f"D{r}", life, fBlue, bd=B_ALL, al=A_C)
        w(ws, f"E{r}", rate, fBlue, PCT, bd=B_ALL, al=A_C)
        r += 1
    cx.finish_print(ws, r)


# ====================================================================== SOA_5
def soa5(cx):
    SC = cx.SC
    keys = cx.PK
    lastp = SC[keys[-1]]
    lastn = SC["OPEN"]
    ws = cx.sheet("SOA_5", {"A": 1.5, "B": 5.0, "C": 52.0, "D": 7.0, **{SC[k]: 16.0 for k in keys + ["OPEN"]}}, print_cols=lastn)
    head(cx, ws, "ANNEXURE V - STATEMENT OF ADJUSTMENTS TO AUDITED FINANCIAL STATEMENTS", lastn)
    r = 8
    w(ws, f"B{r}", 1, fB, al=A_L)
    w(ws, f"C{r}", "Reconciliation of restated profits", fB)
    r += 2
    w(ws, f"C{r}", "Particulars", fB, bd=B_TB, al=A_CW)
    w(ws, f"D{r}", "Note", fB, bd=B_TB, al=A_C)
    for k in keys:
        w(ws, f"{SC[k]}{r}", cx.dtxt(k, "pl"), fB, bd=B_TB, al=A_CW)
    ws.row_dimensions[r].height = 39.75
    r += 1
    rb = r
    w(ws, f"C{r}", "Net profit / (loss) after tax as per audited financial statements *", fB)
    for k in keys:
        w(ws, f"{SC[k]}{r}", f"='Notes_P&L'!{cx.NC[k]}{{{{pl_bookpat}}}}", fB, ACC)
    r += 1
    w(ws, f"C{r}", "Adjustments for:", fB)
    r += 1
    groups = cx.soa_pl_groups()
    letters = "ABCDEFGHIJKLMNOP"
    a0 = r
    for i, (label, masters, _) in enumerate(groups):
        w(ws, f"C{r}", label)
        w(ws, f"D{r}", letters[i], fN, al=A_C)
        for k in keys:
            parts = []
            for m in masters:                       # effect on profit: Cr - Dr
                parts.append(cx.adj_pl(qs(m), k, "Cr"))
            w(ws, f"{SC[k]}{r}", "=ROUND(" + "+".join(parts) + ",6)" if parts else 0, fN, ACC)
        r += 1
    if not groups:
        w(ws, f"C{r}", "No adjustments")
        r += 1
    w(ws, f"C{r}", "Net adjustment in the Statement of Profit and Loss", fB, bd=B_TB)
    for k in keys:
        w(ws, f"{SC[k]}{r}", f"=SUM({SC[k]}{a0}:{SC[k]}{r - 1})", fB, ACC, bd=B_TB)
    rn = r
    r += 1
    w(ws, f"C{r}", "Net profit / (loss) after tax as per Restated Summary Statements", fB, bd=B_TB)
    for i, k in enumerate(keys):
        w(ws, f"{SC[k]}{r}", f"={SC[k]}{rb}+{SC[k]}{rn}", fB, ACC, bd=B_TB)
        w(ws, f"{CL(CI(lastn) + i)}{r}", f"=ROUND({SC[k]}{r}-PL_2!{cx.BC[k]}{cx.PL['pat']},2)", Font(name=TN, size=8, italic=True), ACC)
    r += 1
    star = (f"* For the period ended {long_date(cx.stub)}: net profit as per the books of account (unaudited; no provision for tax made in the books for the "
            "interim period).") if cx.stub else "* As per the audited financial statements."
    t = cx.text("soa_pl_star", star)
    merge(ws, f"C{r}:{lastp}{r}", t, fI, A_J, est_height(t, 52 + 7 + 16 * cx.N))
    r += 2
    w(ws, f"C{r}", "Notes for reconciliation of profit:", fB)
    r += 1
    for i, (label, masters, key) in enumerate(groups):
        w(ws, f"C{r}", f"{letters[i]}) {cx.text('soa_pl_head_' + key, label)}:", fB)
        r += 1
        t = cx.text("soa_pl_" + key, SOA_PL_DEFAULT.get(key, "[Explanation of the adjustment]"))
        merge(ws, f"C{r}:{lastp}{r}", t, fN, A_J, est_height(t, 52 + 7 + 16 * cx.N))
        r += 2
    # ---- shareholders' funds
    w(ws, f"B{r}", 2, fB, al=A_L)
    w(ws, f"C{r}", "Reconciliation of restated shareholders' funds", fB)
    r += 2
    w(ws, f"C{r}", "Particulars", fB, bd=B_TB, al=A_CW)
    w(ws, f"D{r}", "Note", fB, bd=B_TB, al=A_C)
    for k in keys + ["OPEN"]:
        w(ws, f"{SC[k]}{r}", cx.dtxt(k), fB, bd=B_TB, al=A_CW)
    ws.row_dimensions[r].height = 30
    r += 1
    ra = r
    w(ws, f"C{r}", "Shareholders' funds as per audited financial statements (a) *", fB)
    for k in keys + ["OPEN"]:
        w(ws, f"{SC[k]}{r}", "=" + cx.tb(k, qs("Equity share capital")) + "+" + cx.tb(k, qs("Reserves and surplus")), fB, ACC)
    r += 1
    w(ws, f"C{r}", "Adjustments:", fB)
    r += 1
    bgroups = cx.soa_bs_groups()
    b0 = r
    for i, (label, masters, _) in enumerate(bgroups):
        w(ws, f"C{r}", label)
        w(ws, f"D{r}", letters[i], fN, al=A_C)
        for k in keys + ["OPEN"]:
            parts = []
            for m in masters:                       # effect on shareholders' funds: Dr - Cr
                parts.append(cx.adj_bs(qs(m), k, "Dr"))
            w(ws, f"{SC[k]}{r}", "=ROUND(" + "+".join(parts) + ",6)" if parts else 0, fN, ACC)
        r += 1
    if not bgroups:
        w(ws, f"C{r}", "No adjustments")
        r += 1
    w(ws, f"C{r}", "Total adjustments (b)", fB, bd=B_TB)
    for k in keys + ["OPEN"]:
        w(ws, f"{SC[k]}{r}", f"=SUM({SC[k]}{b0}:{SC[k]}{r - 1})", fB, ACC, bd=B_TB)
    rt = r
    r += 1
    w(ws, f"C{r}", "Shareholders' funds as per Restated Summary Statements (a + b)", fB, bd=B_TB)
    chk0 = CI(lastn) + 1
    for i, k in enumerate(keys + ["OPEN"]):
        w(ws, f"{SC[k]}{r}", f"={SC[k]}{ra}+{SC[k]}{rt}", fB, ACC, bd=B_TB)
        w(ws, f"{CL(chk0 + i)}{r}", f"=ROUND({SC[k]}{r}-BS_1!{cx.BC[k]}{cx.BS['shf']},2)", Font(name=TN, size=8, italic=True), ACC)
    cx.setref("soa_sh_restated", r)
    r += 1
    star = (f"* For {long_date(cx.stub)}: as per the books of account (unaudited). " if cx.stub else "* ") + \
        f'The column "As at {long_date(cx.opening)}" represents the adjustments to the opening reserves as at {long_date(cx.opening + __import__("datetime").timedelta(days=1))}.'
    t = cx.text("soa_bs_star", star)
    merge(ws, f"C{r}:{lastn}{r}", t, fI, A_J, est_height(t, 52 + 7 + 16 * (cx.N + 1)))
    r += 2
    w(ws, f"C{r}", "Notes for reconciliation of shareholders' funds:", fB)
    r += 1
    for i, (label, masters, key) in enumerate(bgroups):
        w(ws, f"C{r}", f"{letters[i]}) {cx.text('soa_bs_head_' + key, label)}:", fB)
        r += 1
        t = cx.text("soa_bs_" + key, SOA_BS_DEFAULT.get(key, "[Explanation of the adjustment]"))
        merge(ws, f"C{r}:{lastn}{r}", t, fN, A_J, est_height(t, 52 + 7 + 16 * (cx.N + 1)))
        r += 2
    w(ws, f"B{r}", 3, fB, al=A_L)
    w(ws, f"C{r}", "Material regrouping", fB)
    r += 1
    t = cx.text("soa_regroup", "Appropriate regroupings have been made in the Restated Summary Statements, wherever required, by reclassification of the "
                               "corresponding items of income, expenses, assets, liabilities and cash flows, in order to bring them in line with the accounting "
                               "policies and classification adopted in the latest audited financial statements and the requirements of the SEBI ICDR Regulations.")
    merge(ws, f"C{r}:{lastn}{r}", t, fN, A_J, est_height(t, 52 + 7 + 16 * (cx.N + 1)))
    r += 2
    w(ws, f"B{r}", 4, fB, al=A_L)
    w(ws, f"C{r}", "Non-adjusting items", fB)
    r += 1
    yrs = cx.dates_text(["Y1", "Y2", "Y3"])
    t = cx.text("soa_nonadj", f"There are no audit qualifications in the auditor's reports for the years ended {yrs} which require any adjustment in the "
                              "Restated Summary Statements.")
    merge(ws, f"C{r}:{lastn}{r}", t, fN, A_J, est_height(t, 52 + 7 + 16 * (cx.N + 1)))
    for row in ws.iter_rows(min_col=3, max_col=5):          # line labels wrap as in the final format (not headings / sub-totals)
        c = row[0]
        if (c.value is not None and c.alignment.horizontal != "justify" and row[2].value is not None
                and not str(c.value).startswith(("Net adjustment", "Total adjustments"))):
            c.alignment = Alignment(horizontal=c.alignment.horizontal, vertical=c.alignment.vertical, wrap_text=True)
    cx.finish_print(ws, r + 1)


SOA_PL_DEFAULT = {
    "dep": "During the preparation of the Restated Financial Statements, it was observed that depreciation in the audited financial statements was not "
           "computed in accordance with the Fixed Asset Register and the useful lives prescribed in Schedule II to the Companies Act, 2013. Depreciation "
           "(including the loss on deletion of assets) has been recomputed and the difference adjusted in the respective years / period.",
    "grat": "During the preparation of the Restated Financial Statements, it was observed that the gratuity liability was not provided in the respective years "
            "on the basis of actuarial valuation. Gratuity expense has been recognised in the respective years on the basis of actuarial valuation, and the "
            "portion relating to earlier periods has been adjusted in the opening reserves.",
    "msme": "Interest payable under Section 16 of the MSMED Act, 2006 on dues to micro and small enterprises outstanding beyond the appointed day was not "
            "provided in the books. The same has been provided in the respective years / period.",
    "curtax": "Current tax has been restated to the tax as per the returns of income filed / the restated taxable income of the respective years / period.",
    "deftax": "Deferred tax has been recomputed on the restatement adjustments, being timing differences, and reversed in the years in which the audited "
              "financial statements carry the corresponding balances. No deferred tax is recognised on interest on delayed payments to MSME, being a permanent "
              "disallowance under Section 23 of the MSMED Act.",
}
SOA_BS_DEFAULT = {
    "ppe": "Represents the cumulative effect of restating depreciation as per the Fixed Asset Register, including the alignment of the net block as at the "
           "opening balance sheet date with the Fixed Asset Register (adjusted in the opening reserves) (refer Note A of profit reconciliation).",
    "grat": "Represents the gratuity liability recognised on the basis of actuarial valuation, including the portion relating to periods up to the opening "
            "balance sheet date adjusted in the opening reserves (refer Note B of profit reconciliation).",
    "msme": "Represents the cumulative interest under Section 16 of the MSMED Act, 2006 not provided in the books (refer Note C of profit reconciliation).",
    "taxprov": "Represents the difference in the provision for current tax on account of restatement of current tax (refer Note D of profit reconciliation).",
    "dta": "Represents deferred tax on the above timing differences, including the deferred tax on the opening adjustments (refer Note E of profit "
           "reconciliation).",
}
