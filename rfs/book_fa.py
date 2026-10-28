"""Fixed assets: FA Register (standard template or the client's own register mapped), FA Restatement
(Part A class-wise summary, Part B reconciliation with books, Part C note 12 working), Notes_PPE and
the automatic restatement entries for Property, Plant and Equipment.

Both register modes end in the same normalised block of columns per asset (class, cost, purchase date,
deletion date and, for every period, opening WDV / additions / depreciation / WDV of deletions / closing WDV),
which Part C sums by class."""
import copy
import openpyxl
from openpyxl.styles import Font
from .book_core import *          # noqa: F401,F403
from . import book_masters as BM
from .book_notes import head

FRG, FRS = "FA Register", "FA Restatement"
PPE_M = "Property, Plant and Equipment"


def _n(v):
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def fa_classes(cx):
    """note-12 classes used, in note order: (class, 'PPE' / 'INTANG')"""
    used = set(cx.FA_CLASS_USED)
    out = [(m, "PPE") for m, _ in BM.note_lines("PPE", cx.custom) if m in used]
    out += [(m, "INTANG") for m, _ in BM.note_lines("INTANG", cx.custom) if m in used]
    return out


# ---------------------------------------------------------------------------------------------- register
def fa_register(cx):
    """writes 'FA Register'; sets cx.FAN = normalised column letters, cx.FA_ROWS = (first, last)"""
    FA = cx.I.get("fa") or {}
    cx.FA_MODE = FA.get("mode") or ("standard" if FA.get("assets") else None)
    cx.FA_CLASS_USED = []
    if not cx.FA_MODE:
        return False
    chron = list(reversed(cx.PK))
    if cx.FA_MODE == "mapped":
        src = openpyxl.load_workbook(FA["source"])[FA.get("sheet") or 0] if isinstance(FA.get("sheet"), str) else openpyxl.load_workbook(FA["source"]).worksheets[0]
        ws = cx.wb.create_sheet(FRG)
        for row in src.iter_rows():
            for c in row:
                if c.value is None and not c.has_style:
                    continue
                n = ws.cell(row=c.row, column=c.column, value=c.value)
                if c.has_style:
                    n.font = copy.copy(c.font)
                    n.number_format = c.number_format
                    n.alignment = copy.copy(c.alignment)
                    n.border = copy.copy(c.border)
                    n.fill = copy.copy(c.fill)
        for m in src.merged_cells.ranges:
            ws.merge_cells(str(m))
        for k, d in src.column_dimensions.items():
            ws.column_dimensions[k].width = d.width
        M = FA["map"]
        first, last = int(M["first_row"]), int(M["last_row"])
        hdr = int(M.get("header_row", first - 1))
        start = (src.max_column or 1) + 2
        col = start
        N = {}

        def nc(key, title):
            nonlocal col
            N[key] = CL(col)
            w(ws, f"{CL(col)}{hdr}", title, fB, bd=B_TB, al=A_CW)
            ws.column_dimensions[CL(col)].width = 14
            col += 1
        nc("cls", "Class in Note 12 (normalised)")
        nc("cost", "Cost")
        nc("date", "Date of purchase")
        nc("del", "Date of deletion")
        nc("open", "WDV as at " + long_date(cx.opening))
        for p in chron:
            pl = "period" if p == "S" else "FY " + fy_label(cx.DATE[p])
            for k, t in (("o", "Opening WDV"), ("a", "Additions"), ("d", "Depreciation (restated)"), ("x", "WDV of deletions"), ("c", "Closing WDV")):
                nc(f"{k}{p}", f"{t}\n{pl}")
        ws.row_dimensions[hdr].height = 45
        cmap = FA.get("class_map") or {}
        cx.FA_CLASSMAP = cmap
        for r in range(first, last + 1):
            ccol, cost, dcol = M["class"], M["cost"], M["date"]
            w(ws, f"{N['cls']}{r}", f"=IFERROR(VLOOKUP({ccol}{r},{q(FRS)}!$A$5:$B${4 + max(1, len(cmap))},2,FALSE()),\"\")", fN)
            w(ws, f"{N['cost']}{r}", f"=N({cost}{r})", fN, ACC)
            w(ws, f"{N['date']}{r}", f"={dcol}{r}", fN, DMY)
            w(ws, f"{N['del']}{r}", f'=IF({M["del"]}{r}="","",{M["del"]}{r})' if M.get("del") else "", fN, DMY)
            ow = M["periods"]["OPEN"]["wdv"]
            w(ws, f"{N['open']}{r}", f'=IF(AND({N["date"]}{r}<>"",{N["date"]}{r}>{cx.MDC["OPEN"]}),0,N({ow}{r}))', fN, ACC)
            prev = f"{N['open']}{r}"
            for p in chron:
                pm = M["periods"][p]
                S_, E_ = cx.MDC[cx.PREV[p]], cx.MDC[p]
                d, wv = pm["dep"], pm["wdv"]
                dr = pm.get("dep_reg") or d
                live = f'OR({N["date"]}{r}="",{N["date"]}{r}<={E_})'
                w(ws, f"{N['o' + p]}{r}", f"={prev}", fN, ACC)
                w(ws, f"{N['a' + p]}{r}", f'=IF(AND({N["date"]}{r}>{S_},{N["date"]}{r}<={E_}),{N["cost"]}{r},0)', fN, ACC)
                w(ws, f"{N['d' + p]}{r}", f"=IF({live},N({d}{r}),0)", fN, ACC)
                w(ws, f"{N['x' + p]}{r}", f"=IF({live},{N['o' + p]}{r}+{N['a' + p]}{r}-N({dr}{r})-N({wv}{r}),0)", fN, ACC)
                w(ws, f"{N['c' + p]}{r}", f"=IF({live},N({wv}{r})+N({dr}{r})-N({d}{r}),0)", fN, ACC)
                prev = f"{N['c' + p]}{r}"
        cx.FA_CLASS_USED = sorted(set(cmap.values()))
        cx.FAN, cx.FA_ROWS = N, (first, last)
        cx.FA_CLASSIFICATIONS = list(cmap.items())
        w(ws, f"{CL(start)}{max(1, hdr - 1)}", "Normalised columns used by 'FA Restatement' (restated WDV; additions / deletions by date)", fGreyI)
        return True

    # ---- standard template: the register is generated with Schedule II WDV formulas
    assets = FA.get("assets") or []
    ws = cx.sheet(FRG, {"A": 2, "B": 6, "C": 8, "D": 34, "E": 30, "F": 24, "G": 10, "H": 14, "I": 12, "J": 10, "K": 9, "L": 12, "M": 12, "N": 12})
    w(ws, "B1", "Fixed Asset Register (Companies Act, 2013 - Schedule II, WDV method, residual value) - amount in Rs.", fB)
    w(ws, "B2", "Depreciation pro rata for the days of use in the year of acquisition / deletion; written down to the residual value when the useful "
                "life is completed. Inputs in blue.", fI)
    heads = ["Sl.no", "FY", "Description", "FA Classification", "Class in Note 12", "Rate (WDV)", "Cost", "Date of purchase", "Useful life (years)",
             "Residual value %", "Date of deletion", "Sale value", "WDV as at " + long_date(cx.opening)]
    for i, h in enumerate(heads):
        w(ws, f"{CL(2 + i)}5", h, fB, bd=B_TB, al=A_CW)
    N = {"cls": "F", "cost": "H", "date": "I", "del": "L", "open": "N"}
    col = 15
    for p in chron:
        pl = "period" if p == "S" else "FY " + fy_label(cx.DATE[p])
        for k, t in (("o", "Opening WDV"), ("a", "Additions"), ("d", "Depreciation"), ("x", "WDV of deletions"), ("c", "Closing WDV")):
            N[f"{k}{p}"] = CL(col)
            w(ws, f"{CL(col)}5", f"{t}\n{pl}", fB, bd=B_TB, al=A_CW)
            ws.column_dimensions[CL(col)].width = 13
            col += 1
    ws.row_dimensions[5].height = 45
    O = cx.MDC["OPEN"]
    r = 6
    cls_used = set()
    for k, a in enumerate(assets):
        cls = a.get("class") or ""
        cls_used.add(cls)
        d = ddate(a.get("date")) if a.get("date") else None
        w(ws, f"B{r}", k + 1)
        w(ws, f"C{r}", a.get("fy") or (fy_label(d)[2:] if d else ""), fBlue)
        w(ws, f"D{r}", a.get("name", ""), fBlue)
        w(ws, f"E{r}", a.get("classification") or cls, fBlue)
        w(ws, f"F{r}", cls, fBlue)
        life = _n(a.get("life"))
        res = _n(a.get("residual", 5)) if a.get("residual") not in (None, "") else 5.0
        if a.get("rate") not in (None, ""):
            w(ws, f"G{r}", _n(a["rate"]), fBlue, "0.00%")
        elif life:
            w(ws, f"G{r}", f"=IF(J{r}>0,1-(K{r}/100)^(1/J{r}),0)", fN, "0.00%")
        else:
            w(ws, f"G{r}", 0, fBlue, "0.00%")
        w(ws, f"H{r}", _n(a.get("cost")), fBlue, ACC)
        w(ws, f"I{r}", d, fBlue, DMY)
        w(ws, f"J{r}", life or None, fBlue)
        w(ws, f"K{r}", res, fBlue)
        if a.get("del_date"):
            w(ws, f"L{r}", ddate(a["del_date"]), fBlue, DMY)
        if a.get("sale") not in (None, ""):
            w(ws, f"M{r}", _n(a["sale"]), fBlue, ACC)
        if a.get("wdv_open") not in (None, ""):
            w(ws, f"N{r}", _n(a["wdv_open"]), fBlue, ACC)      # WDV as per the existing register at the opening date
        else:
            # computed: full Schedule II WDV from the date of purchase to the opening date
            yrs = f"MAX(0,({O}-I{r}+1)/365)"
            w(ws, f"N{r}", f'=IF(OR(I{r}="",I{r}>{O},AND(L{r}<>"",L{r}<={O})),0,IF(AND(J{r}>0,{yrs}>=J{r}),H{r}*K{r}/100,'
                           f'MAX(H{r}*K{r}/100,H{r}*(1-G{r})^{yrs})))', fN, ACC)
        prev = f"N{r}"
        for p in chron:
            S_, E_ = cx.MDC[cx.PREV[p]], cx.MDC[p]
            o, ad, dp, x, c = (N[f"{k_}{p}"] for k_ in "oadxc")
            dl = f'AND(L{r}<>"",L{r}>{S_},L{r}<={E_})'
            upto = f"IF({dl},L{r},{E_})"
            resid = f"H{r}*K{r}/100"
            done = f"AND(J{r}>0,({upto}-I{r}+1)/365>=J{r})"
            w(ws, f"{o}{r}", f"={prev}", fN, ACC)
            w(ws, f"{ad}{r}", f"=IF(AND(I{r}>{S_},I{r}<={E_}),H{r},0)", fN, ACC)
            base = f"({o}{r}+{ad}{r})"
            days = f"IF({ad}{r}>0,{upto}-I{r}+1,{upto}-{S_})"
            w(ws, f"{dp}{r}", f"=IF({base}<=0,0,IF({done},MAX(0,{base}-{resid}),MIN(MAX(0,{base}-{resid}),{base}*G{r}*{days}/365)))", fN, ACC)
            w(ws, f"{x}{r}", f"=IF({dl},{o}{r}+{ad}{r}-{dp}{r},0)", fN, ACC)
            w(ws, f"{c}{r}", f"={o}{r}+{ad}{r}-{dp}{r}-{x}{r}", fN, ACC)
            prev = f"{c}{r}"
        r += 1
    last = max(6, r - 1)
    w(ws, f"D{r}", "Total", fB, bd=B_TB2)
    for cc in ["H", "N"] + [N[f"{k_}{p}"] for p in chron for k_ in "oadxc"]:
        w(ws, f"{cc}{r}", f"=SUM({cc}6:{cc}{last})", fB, ACC, bd=B_TB2)
    ws.freeze_panes = "E6"
    cx.FAN, cx.FA_ROWS = N, (6, last)
    cx.FA_CLASS_USED = sorted(c for c in cls_used if c)
    cx.FA_CLASSIFICATIONS = sorted({(a.get("classification") or a.get("class") or "", a.get("class") or "") for a in assets})
    return True


# ---------------------------------------------------------------------------------------------- FA Restatement
def fa_restatement(cx):
    """Part A (by FA classification), Part B (reconciliation with books, Rs. lakhs), Part C (by note-12 class, Rs.)"""
    N = cx.FAN
    f1, l1 = cx.FA_ROWS
    chron = list(reversed(cx.PK))
    ws = cx.sheet(FRS, {"A": 62, "B": 22, **{CL(3 + i): 14 for i in range(40)}})
    rg = lambda c: f"{q(FRG)}!${c}${f1}:${c}${l1}"
    S = cx.SCALE
    w(ws, "A1", "Restatement of Property, Plant and Equipment and Intangible assets as per Fixed Asset Register", fB)
    w(ws, "A2", "Part A - Class-wise summary of the Fixed Asset Register (amount in Rs.)", fB)
    # ---- Part A (rows 5.. : FA classification -> class in note 12; the register looks this table up)
    heads = ["FA classification", "Class in Note 12", "Cost", "WDV as at\n" + cx.opening.strftime("%d.%m.%Y")]
    for p in chron:
        tag = ("period " if p == "S" else "FY ") + (cx.DATE[p].strftime("%d.%m.%Y") if p == "S" else fy_label(cx.DATE[p]))
        heads += [f"Additions\n{tag}", f"Depreciation\n{tag}", f"Deletions\n{tag}", f"WDV as at\n{cx.DATE[p]:%d.%m.%Y}"]
    for i, h in enumerate(heads):
        w(ws, f"{CL(1 + i)}4", h, fB, bd=B_TB, al=A_CW)
    ws.row_dimensions[4].height = 30
    if cx.FA_MODE == "mapped":
        clsrc = f"{q(FRG)}!${cx.I['fa']['map']['class']}${f1}:${cx.I['fa']['map']['class']}${l1}"
    else:
        clsrc = rg("E")
    r = 5
    pairs = cx.FA_CLASSIFICATIONS or [(c, c) for c in cx.FA_CLASS_USED]
    for fc, nc in pairs:
        w(ws, f"A{r}", fc)
        w(ws, f"B{r}", nc, fBlue)
        crit = f"{clsrc},$A{r}"
        w(ws, f"C{r}", f"=SUMIFS({rg(N['cost'])},{crit})", fmt=ACC)
        w(ws, f"D{r}", f"=SUMIFS({rg(N['open'])},{crit})", fmt=ACC)
        c = 5
        for p in chron:
            for k in "adxc":
                w(ws, f"{CL(c)}{r}", f"=SUMIFS({rg(N[k + p])},{crit})", fmt=ACC)
                c += 1
        r += 1
    last_a = r - 1
    w(ws, f"A{r}", "Total", fB, bd=B_TB)
    for c in range(3, 5 + 4 * len(chron)):
        w(ws, f"{CL(c)}{r}", f"=SUM({CL(c)}5:{CL(c)}{max(5, last_a)})", fB, ACC, bd=B_TB)
    tot_a = r
    r += 1
    w(ws, f"A{r}", "Check: total cost agrees with the Fixed Asset Register (unclassified assets show here)", fGrey)
    w(ws, f"C{r}", f"=ROUND(C{tot_a}-SUM({rg(N['cost'])}),0)", fChk, ACC)
    r += 1
    w(ws, f"A{r}", cx.text("fa_parta_note", "Deletions are the WDV of assets sold / written off as per the Fixed Asset Register (opening WDV + additions - "
                                            "depreciation - closing WDV)."), fI)
    r += 2
    # ---- Part C first (Part B refers to it): by note-12 class, amount in Rs.
    classes = fa_classes(cx)
    rC = r + 30 + 0
    # Part B position
    rB = r
    # layout of Part C columns
    pc = {}
    c = 3
    pc["g:OPEN"] = CL(c); c += 1
    for p in chron:
        for k in ("add", "del", "g"):
            pc[f"{k}:{p}"] = CL(c); c += 1
    pc["w:OPEN"] = CL(c); c += 1
    for p in chron:
        pc[f"w:{p}"] = CL(c); c += 1
    for p in chron:
        pc[f"dep:{p}"] = CL(c); c += 1
    for p in chron:
        pc[f"x:{p}"] = CL(c); c += 1
    cx.PC = pc
    # Part B rows (Rs. lakhs)
    w(ws, f"A{rB}", "Part B - Reconciliation with books of account and restatement adjustments (Rs. in lakhs)", fB)
    rB += 1
    keys = ["OPEN"] + chron
    BCOL = {k: CL(3 + i) for i, k in enumerate(keys)}
    cx.FBC = BCOL
    w(ws, f"A{rB}", "Particulars", fB, bd=B_TB)
    for k in keys:
        lab = ("As at / period ended\n" if k == "S" else "As at / FY ended\n") + cx.DATE[k].strftime("%d.%m.%Y")
        w(ws, f"{BCOL[k]}{rB}", lab, fB, bd=B_TB, al=A_CW)
    ws.row_dimensions[rB].height = 30
    rB += 1
    R = {}
    cls_lbls = [cl for cl, _ in classes]
    for nm, lbl in (("nb_reg", "Net block (PPE and intangible assets) as per Fixed Asset Register (restated)"),
                    ("nb_book", "Net block (PPE and intangible assets) as per books / audited financial statements"),
                    ("a", "Difference - to be restated (a)"), (None, None),
                    ("dep_reg", "Depreciation and amortisation as per Fixed Asset Register (restated)"),
                    ("dep_book", "Depreciation and amortisation as per books / audited financial statements"),
                    ("b", "Excess / (short) depreciation charged in books (b)"), (None, None),
                    ("d", "Net change in Property, Plant and Equipment credited / (debited) directly to Reserves and surplus in the books (d) - input"),
                    ("c", "Difference on additions / deletions taken to the Statement of Profit and Loss (c) = (a) - (a) previous - (b) + (d)"), (None, None),
                    ("rate", "Tax rate"), ("dt", "Deferred tax on restatement adjustment - charge / (credit)"), ("chk", None)):
        if nm is None:
            rB += 1
            continue
        R[nm] = rB
        if lbl:
            w(ws, f"A{rB}", lbl, fB if nm in ("a", "b", "c") else fN)
        rB += 1
    cx.FBR = R
    cls_list = ",".join(qs(x) for x in cls_lbls) if cls_lbls else qs("Land")
    D_IN = (cx.I.get("fa") or {}).get("direct_reserve") or {}
    for k in keys:
        col = BCOL[k]
        w(ws, f"{col}{R['nb_reg']}", f"=SUM({pc['w:' + k]}{rC + 3}:{pc['w:' + k]}{rC + 2 + max(1, len(classes))})/{S}", fmt=ACC)
        if k == "OPEN":
            tbp = q(cx.TBN[cx.EARLY])
            w(ws, f"{col}{R['nb_book']}", f"=SUMPRODUCT(SUMIF({tbp}!$J:$J,{{{cls_list}}},{tbp}!$L:$L))", fmt=ACC)
        else:
            tbp = q(cx.TBN[k])
            w(ws, f"{col}{R['nb_book']}", f"=SUMPRODUCT(SUMIF({tbp}!$G:$G,{{{cls_list}}},{tbp}!$I:$I))", fmt=ACC)
        w(ws, f"{col}{R['a']}", f"={col}{R['nb_reg']}-{col}{R['nb_book']}", fB, ACC, bd=B_T)
        w(ws, f"{col}{R['rate']}", f"={cx.TAX}", fmt=PCT3)
        if k == "OPEN":
            w(ws, f"{col}{R['dt']}", f"=ROUND({col}{R['a']}*{col}{R['rate']},2)", fmt=ACC)
            continue
        pv = BCOL[cx.PREV[k]]
        w(ws, f"{col}{R['dep_reg']}", f"=SUM({pc['dep:' + k]}{rC + 3}:{pc['dep:' + k]}{rC + 2 + max(1, len(classes))})/{S}", fmt=ACC)
        w(ws, f"{col}{R['dep_book']}", "=" + cx.tb(k, qs("Depreciation and amortisation")), fmt=ACC)
        w(ws, f"{col}{R['b']}", f"={col}{R['dep_book']}-{col}{R['dep_reg']}", fB, ACC, bd=B_T)
        dv = D_IN.get(k)
        w(ws, f"{col}{R['d']}", f"={_n(dv):g}/{S}" if dv not in (None, "") else 0, fBlue, ACC, fill=Y_IN)
        w(ws, f"{col}{R['c']}", f"={col}{R['a']}-{pv}{R['a']}-{col}{R['b']}+{col}{R['d']}", fB, ACC, bd=B_T)
        # deferred tax: on the P&L part; the part reversed through reserves takes back the deferred tax recognised on it earlier
        cum = "+".join(f"{BCOL[x]}{R['dt']}" for x in keys[:keys.index(k)])
        w(ws, f"{col}{R['dt']}", f"=ROUND(({col}{R['b']}+{col}{R['c']})*{col}{R['rate']},2)+IF(ABS({pv}{R['a']})>0.000001,"
                                 f"-ROUND(({cum})*{col}{R['d']}/{pv}{R['a']},2),ROUND({col}{R['d']}*{col}{R['rate']},2))", fmt=ACC)
    w(ws, f"A{R['chk']}", cx.text("fa_partb_note", "Restatement entries are recorded in sheet \"Adj Entry RFS\" (rows marked \"FA Register\"). (c) should "
                                                   "normally be the difference between the loss / profit on deletion as per books and as per the register."), fI)
    r = max(rB + 1, rC)
    rC = r
    # ---- Part C
    w(ws, f"A{rC}", "Part C - Gross block, WDV and depreciation by class as per the Fixed Asset Register (amount in Rs.) - used in Note 12", fB)
    hdr = rC + 1
    w(ws, f"A{hdr}", "Class in Note 12", fB, bd=B_TB, al=A_CW)
    lab = lambda k: cx.DATE[k].strftime("%d.%m.%Y")
    tg = lambda p: ("period " if p == "S" else "FY ") + (lab(p) if p == "S" else fy_label(cx.DATE[p]))
    titles = {"g:OPEN": f"Gross block\n{lab('OPEN')}", "w:OPEN": f"WDV\n{lab('OPEN')}"}
    for p in chron:
        titles[f"add:{p}"] = f"Additions\n{tg(p)}"
        titles[f"del:{p}"] = f"Deletions\n{tg(p)}"
        titles[f"g:{p}"] = f"Gross block\n{lab(p)}"
        titles[f"w:{p}"] = f"WDV\n{lab(p)}"
        titles[f"dep:{p}"] = f"Depreciation\n{tg(p)}"
        titles[f"x:{p}"] = f"WDV of deletions\n{tg(p)}"
    for k, cl in pc.items():
        w(ws, f"{cl}{hdr}", titles[k], fB, bd=B_TB, al=A_CW)
    ws.row_dimensions[hdr].height = 30
    r = hdr + 1
    r0 = r
    cx.PCROW = {}
    cr = rg(N["cls"])
    for cl_, _ in classes:
        w(ws, f"A{r}", cl_)
        crit = f"{cr},$A{r}"
        gross = lambda d: (f'=SUMIFS({rg(N["cost"])},{crit},{rg(N["date"])},"<="&{d})-SUMIFS({rg(N["cost"])},{crit},{rg(N["del"])},"<="&{d})')
        w(ws, f"{pc['g:OPEN']}{r}", gross(cx.MDC["OPEN"]), fmt=ACC)
        w(ws, f"{pc['w:OPEN']}{r}", f"=SUMIFS({rg(N['open'])},{crit})", fmt=ACC)
        for p in chron:
            S_, E_ = cx.MDC[cx.PREV[p]], cx.MDC[p]
            w(ws, f"{pc['add:' + p]}{r}", f'=SUMIFS({rg(N["cost"])},{crit},{rg(N["date"])},">"&{S_},{rg(N["date"])},"<="&{E_})', fmt=ACC)
            w(ws, f"{pc['del:' + p]}{r}", f'=SUMIFS({rg(N["cost"])},{crit},{rg(N["del"])},">"&{S_},{rg(N["del"])},"<="&{E_})', fmt=ACC)
            w(ws, f"{pc['g:' + p]}{r}", gross(E_), fmt=ACC)
            w(ws, f"{pc['w:' + p]}{r}", f"=SUMIFS({rg(N['c' + p])},{crit})", fmt=ACC)
            w(ws, f"{pc['dep:' + p]}{r}", f"=SUMIFS({rg(N['d' + p])},{crit})", fmt=ACC)
            w(ws, f"{pc['x:' + p]}{r}", f"=SUMIFS({rg(N['x' + p])},{crit})", fmt=ACC)
        cx.PCROW[cl_] = r
        r += 1
    if not classes:
        r += 1
    w(ws, f"A{r}", "Total", fB)
    for k, cl in pc.items():
        w(ws, f"{cl}{r}", f"=SUM({cl}{r0}:{cl}{max(r0, r - 1)})", fB, ACC, bd=B_TB)
    r += 1
    w(ws, f"A{r}", "Check: gross block movement (opening + additions - deletions - closing) should be nil", fGrey)
    for p in chron:
        g0 = pc["g:" + cx.PREV[p]] if cx.PREV[p] != "OPEN" else pc["g:OPEN"]
        w(ws, f"{pc['g:' + p]}{r}", f"=ROUND({g0}{r - 1}+{pc['add:' + p]}{r - 1}-{pc['del:' + p]}{r - 1}-{pc['g:' + p]}{r - 1},0)", fChk, ACC)
    # Part B now knows Part C rows: fix the references written with the provisional rC
    real0, realn = r0, max(r0, r0 + len(classes) - 1)
    for k in keys:
        col = BCOL[k]
        ws[f"{col}{R['nb_reg']}"].value = f"=SUM({pc['w:' + k]}{real0}:{pc['w:' + k]}{realn})/{S}"
        if k != "OPEN":
            ws[f"{col}{R['dep_reg']}"].value = f"=SUM({pc['dep:' + k]}{real0}:{pc['dep:' + k]}{realn})/{S}"
    cx.finish_print(ws, r)


def ppe_entries(cx):
    """automatic restatement entries (single master 'Property, Plant and Equipment' for the PPE notes' check rows)"""
    R, B = cx.FBR, cx.FBC
    rs = f"{q(FRS)}!"
    chron = list(reversed(cx.PK))
    E = []
    a0 = f"{rs}${B['OPEN']}${R['a']}"
    dt0 = f"{rs}${B['OPEN']}${R['dt']}"

    def L(master, amt, part, to=False):
        return {"master": master, "particulars": part, "dr": f"=MAX({amt},0)", "cr": f"=MAX(-({amt}),0)", "to": to}
    E.append({"date": cx.opening, "lines": [L(PPE_M, a0, PPE_M), L("Reserves and surplus", f"-{a0}", "Reserves and surplus", True)],
              "narration": f"Being net block as at {long_date(cx.opening)} aligned with the Fixed Asset Register - FA Register",
              "working": f"FA Register: WDV difference as at {cx.opening:%d.%m.%Y} ('FA Restatement' Part B (a))"})
    E.append({"date": cx.opening, "lines": [L("Reserves and surplus", dt0, "Reserves and surplus"), L("Deferred tax asset (net)", f"-{dt0}", "Deferred tax asset (net)", True)],
              "narration": "Being deferred tax on the above opening adjustment - FA Register", "working": "FA Register: deferred tax on the above"})
    for p in chron:
        c = B[p]
        b, cc, d, dt = (f"{rs}${c}${R[k]}" for k in ("b", "c", "d", "dt"))
        pw = cx.per_word(p)
        E.append({"date": cx.DATE[p], "lines": [L(PPE_M, b, PPE_M), L("Depreciation and amortisation", f"-{b}", f"Depreciation and amortisation (P&L - {_tag(cx, p)})", True)],
                  "narration": f"Being depreciation for the {pw} ended {long_date(cx.DATE[p])} restated as per the Fixed Asset Register - FA Register",
                  "working": "FA Register: excess / (short) depreciation charged in books"})
        E.append({"date": cx.DATE[p], "lines": [L(PPE_M, cc, PPE_M), L("Miscellaneous expenses", f"-{cc}", f"Miscellaneous expenses (loss on deletion - {_tag(cx, p)})", True)],
                  "narration": f"Being difference on additions / deletions of assets for the {pw} ended {long_date(cx.DATE[p])} as per the Fixed Asset Register - FA Register",
                  "working": "FA Register: difference on deletions (c)"})
        E.append({"date": cx.DATE[p], "lines": [L("Reserves and surplus", d, "Reserves and surplus"), L(PPE_M, f"-{d}", PPE_M, True)],
                  "narration": f"Being reversal of the alignment of Property, Plant and Equipment credited directly to reserves in the books for the {pw} ended "
                               f"{long_date(cx.DATE[p])} - FA Register", "working": "FA Register: (d)"})
        E.append({"date": cx.DATE[p], "lines": [L("Deferred tax", dt, f"Deferred tax (P&L - {_tag(cx, p)})"), L("Deferred tax asset (net)", f"-{dt}", "Deferred tax asset (net)", True)],
                  "narration": "Being deferred tax on the above - FA Register", "working": "FA Register: deferred tax on the above"})
    return E


def _tag(cx, p):
    return f"period ended {cx.DATE[p]:%d.%m.%Y}" if p == "S" else "FY " + fy_label(cx.DATE[p])


# ---------------------------------------------------------------------------------------------- Notes_PPE
def notes_ppe(cx):
    widths = {"A": 1.0, "B": 5.0, "C": 38.0, **{CL(4 + i): 13.5 for i in range(10)}, "N": 2.0, "O": 11.0}
    ws = cx.sheet("Notes_PPE", widths, print_cols="M")
    head(cx, ws, "ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS", "M")
    S = cx.SCALE
    w(ws, "B8", f"=Notes_Liabilities!$B${{{{note:STP}}}}+1", fB, al=A_L)
    w(ws, "C8", "Property, Plant and Equipment and Intangible assets", fB)
    cx.noteno_counter += 1
    cx.noteno["PPE"] = cx.noteno_counter
    cx.setref("note:PPE", 8)
    r = 10
    cx.PPE_ROWS = {}
    if not getattr(cx, "FAN", None):
        return _ppe_fallback(cx, ws, r)
    pc = cx.PC
    FR = q(FRS)
    classes = fa_classes(cx)
    for p in cx.PK:
        pw = "period" if p == "S" else "year"
        prev = cx.PREV[p]
        w(ws, f"C{r}", f'="As at "&TEXT({cx.MDC[p]},"mmmm dd, yyyy")', fB)
        r += 1
        w(ws, f"C{r}", "Particulars", fB, bd=B_ALL, al=A_CW)
        ws.merge_cells(f"C{r}:C{r + 1}")
        for c0, c1, t in (("D", "G", "Gross Block"), ("H", "K", "Accumulated Depreciation and Amortisation"), ("L", "M", "Net Block")):
            w(ws, f"{c0}{r}", t, fB, bd=B_ALL, al=A_CW)
            ws.merge_cells(f"{c0}{r}:{c1}{r}")
        r += 1
        sd = f'="As at "&TEXT({cx.MDC[prev]}+1,"mmmm d, yyyy")'
        ed = f'="As at "&TEXT({cx.MDC[p]},"mmmm dd, yyyy")'
        pe = f'="As at "&TEXT({cx.MDC[prev]},"mmmm dd, yyyy")'
        for col, t in zip("DEFGHIJKLM", [sd, f"Additions during the {pw}", f"Deletions during the {pw}", ed, sd, f"Depreciation / amortisation for the {pw}",
                                         "Depreciation on deletions", ed, pe, ed]):
            w(ws, f"{col}{r}", t, fB, bd=B_ALL, al=A_CW)
        ws.row_dimensions[r].height = 42
        r += 1
        blocks = []
        for grp, title, tot in (("PPE", "(a) Property, Plant and Equipment", "Total (a)"), ("INTANG", "(b) Intangible assets", "Total (b)")):
            cls = [c for c, g in classes if g == grp]
            if not cls:
                continue
            w(ws, f"C{r}", title, fBI)
            r += 1
            f0 = r
            for cl_ in cls:
                pr = cx.PCROW[cl_]
                g0 = pc["g:" + prev]
                w0 = pc["w:" + prev]
                w(ws, f"C{r}", cl_, fN, bd=B_LR)
                vals = {
                    "D": f"={FR}!{g0}{pr}/{S}", "E": f"={FR}!{pc['add:' + p]}{pr}/{S}", "F": f"={FR}!{pc['del:' + p]}{pr}/{S}", "G": f"=D{r}+E{r}-F{r}",
                    "H": f"=({FR}!{g0}{pr}-{FR}!{w0}{pr})/{S}", "I": f"={FR}!{pc['dep:' + p]}{pr}/{S}",
                    "J": f"=({FR}!{pc['del:' + p]}{pr}-{FR}!{pc['x:' + p]}{pr})/{S}", "K": f"=H{r}+I{r}-J{r}", "L": f"=D{r}-H{r}", "M": f"=G{r}-K{r}"}
                fGreen = Font(name=TN, size=10, color="008000")      # values linked from 'FA Restatement'
                for col, v in vals.items():
                    w(ws, f"{col}{r}", v, fGreen if col in "DEFHIJ" else fN, ACC, bd=B_LR)
                w(ws, f"O{r}", f"=ROUND(M{r}-{FR}!{pc['w:' + p]}{pr}/{S},2)", fChk, ACC)
                r += 1
            w(ws, f"C{r}", tot, fB, bd=B_ALL)
            for col in "DEFGHIJKLM":
                w(ws, f"{col}{r}", f"=SUM({col}{f0}:{col}{r - 1})", fB, ACC, bd=B_ALL)
            blocks.append(r)
            cx.PPE_ROWS[(p, grp)] = r
            r += 1
        if len(blocks) == 2:
            w(ws, f"C{r}", "Total (a) + (b)", fB, bd=B_ALL)
            for col in "DEFGHIJKLM":
                w(ws, f"{col}{r}", f"={col}{blocks[0]}+{col}{blocks[1]}", fB, ACC, bd=B_ALL)
            r += 1
        r += 2
    r = _cwip(cx, ws, r)
    # notes
    r += 1
    w(ws, f"C{r}", "Notes:", fB)
    r += 1
    dflt = [
        f"There is no capital work-in-progress as at {cx.dates_text()}; no project is overdue or has exceeded its cost compared to its original plan."
        if not cx.nonzero("Capital work-in-progress") else "",
        "The Company maintains its fixed assets on written down value basis. Property, Plant and Equipment and intangible assets have been restated as per "
        "the Fixed Asset Register (depreciation on WDV method at rates based on useful lives prescribed in Schedule II to the Companies Act, 2013, with "
        "residual value of 5%) - refer Annexure V.",
        "Gross block and accumulated depreciation are as per the Fixed Asset Register.",
        "The Company has not revalued its Property, Plant and Equipment or intangible assets during the years presented.",
    ]
    for i, t in enumerate(cx.text_list("ppe_notes", dflt, "Note 12 (PPE) - notes (one per line)")):
        if not t:
            continue
        merge(ws, f"C{r}:M{r}", t, fN, A_J, est_height(t, 160))
        r += 1
    # checks: net block = books + restatement entries
    r += 1
    w(ws, f"C{r}", "Checks (should be nil)", fGreyB)
    r += 1
    for p in cx.PK + ["OPEN"]:
        lbl = (f"{'Period' if p == 'S' else 'Year'} ended {cx.DATE[p]:%b-%y}: net block = books + restatement entries" if p != "OPEN"
               else f"{long_date(cx.opening)}: opening net block = books + restatement entries")
        w(ws, f"C{r}", lbl, fGrey)
        nb = _nb_ref(cx, p)
        if nb:
            book = f"{FR}!{cx.FBC[p]}{cx.FBR['nb_book']}"
            w(ws, f"D{r}", f"=ROUND({nb}-({book}+{cx.adj_bs(qs(PPE_M), p, 'Dr')[1:-1]}),2)", fGrey, ACC)
        r += 1
    cx.finish_print(ws, r)


def _nb_ref(cx, p):
    """net block (PPE + intangibles) cells in Notes_PPE for period p (opening: column L of the earliest block)"""
    k = p if p != "OPEN" else cx.EARLY
    col = "M" if p != "OPEN" else "L"
    rows = [cx.PPE_ROWS.get((k, g)) for g in ("PPE", "INTANG")]
    rows = [x for x in rows if x]
    return "(" + "+".join(f"{col}{x}" for x in rows) + ")" if rows else None


def _cwip(cx, ws, r):
    S = cx.SCALE
    cw = [m for m, _ in BM.note_lines("CWIP", cx.custom)]
    w(ws, f"C{r}", "(c) Capital work-in-progress and (d) Intangible assets under development", fB)
    r += 1
    cx.CWIP_ROWS = {}
    for p in cx.PK:
        pw = "period" if p == "S" else "year"
        prev = cx.PREV[p]
        w(ws, f"C{r}", f'="As at "&TEXT({cx.MDC[p]},"mmmm dd, yyyy")', fB)
        r += 1
        for col, t in zip("CDEFG", ["Particulars", f'="As at "&TEXT({cx.MDC[prev]}+1,"mmmm d, yyyy")', f"Additions during the {pw}",
                                     f"Capitalised / deletions during the {pw}", f'="As at "&TEXT({cx.MDC[p]},"mmmm dd, yyyy")']):
            w(ws, f"{col}{r}", t, fB, bd=B_ALL, al=A_CW)
        ws.row_dimensions[r].height = 30
        r += 1
        f0 = r
        for m in cw:
            w(ws, f"C{r}", m, fN, bd=B_LR)
            fGreen = Font(name=TN, size=10, color="008000")
            w(ws, f"D{r}", "=" + cx.tb(prev, f"$C{r}"), fGreen, ACC, bd=B_LR)
            w(ws, f"E{r}", "=" + cx.tb(p, f"$C{r}", col="C") + f"/{S}", fGreen, ACC, bd=B_LR)
            w(ws, f"F{r}", f"=D{r}+E{r}-G{r}", fN, ACC, bd=B_LR)
            w(ws, f"G{r}", "=" + cx.tb(p, f"$C{r}"), fGreen, ACC, bd=B_LR)
            cx.CWIP_ROWS[(p, m)] = r
            r += 1
        w(ws, f"C{r}", "Total", fB, bd=B_ALL)
        for col in "DEFG":
            w(ws, f"{col}{r}", f"=SUM({col}{f0}:{col}{r - 1})", fB, ACC, bd=B_ALL)
        r += 2
    # ageing of intangible assets under development / CWIP (cost incurred in each period)
    for m, ttl in (("Intangible assets under development", "Ageing of intangible assets under development (projects in progress)"),
                   ("Capital work-in-progress", "Ageing of capital work-in-progress (projects in progress)")):
        if not cx.nonzero(m, cx.PK):              # ageing only when there is a balance at a period end
            continue
        w(ws, f"C{r}", ttl, fB)
        r += 1
        for i, p in enumerate(cx.PK):
            w(ws, f"C{r}", f'="As at "&TEXT({cx.MDC[p]},"mmmm dd, yyyy")', fB, bd=B_ALL)
            for col, t in zip("DEFGH", ["Less than 1 year", "1-2 years", "2-3 years", "More than 3 years", "Total"]):
                w(ws, f"{col}{r}", t, fB, bd=B_ALL, al=A_CW)
            r += 1
            w(ws, f"C{r}", cx.text("cwip_project", "Projects in progress"), fN, bd=B_ALL)
            later = cx.PK[i:]
            adds = [f"E{cx.CWIP_ROWS[(k, m)]}" for k in later]
            close = f"G{cx.CWIP_ROWS[(p, m)]}"
            fs = []
            for j in range(3):
                if j < len(adds):
                    if j == 0:
                        fs.append(f"={adds[0]}" if p != "S" or True else "")
                    else:
                        fs.append(f"=MIN({adds[j]},{close}-" + "-".join(f"{CL(4 + t)}{r}" for t in range(j)) + ")")
                else:
                    fs.append(0)
            for j, v in enumerate(fs):
                w(ws, f"{CL(4 + j)}{r}", v, fN, ACC, bd=B_ALL)
            w(ws, f"G{r}", f"={close}-D{r}-E{r}-F{r}", fN, ACC, bd=B_ALL)
            w(ws, f"H{r}", f"=SUM(D{r}:G{r})", fN, ACC, bd=B_ALL)
            r += 2
    t = cx.text("cwip_note", "Ageing is based on the period in which the cost was incurred (stub period / financial year). No project is overdue or has "
                             "exceeded its cost compared to its original plan.")
    merge(ws, f"C{r}:M{r}", t, fI, A_J, est_height(t, 160))
    return r + 1


def _ppe_fallback(cx, ws, r):
    """no fixed asset register: net block by class as per the trial balances"""
    S = cx.SCALE
    w(ws, f"C{r}", "Net block as per books (Fixed Asset Register not furnished)", fBI)
    r += 1
    w(ws, f"C{r}", "Particulars", fB, bd=B_TB)
    keys = cx.PK + ["OPEN"]
    for i, k in enumerate(keys):
        w(ws, f"{CL(4 + i)}{r}", cx.dtxt(k), fB, bd=B_TB, al=A_CW)
    r += 1
    for grp in ("PPE", "INTANG"):
        f0 = r
        for m, _ in BM.note_lines(grp, cx.custom):
            if not cx.nonzero(m):
                continue
            w(ws, f"C{r}", m)
            for i, k in enumerate(keys):
                w(ws, f"{CL(4 + i)}{r}", "=" + cx.tb(k, f"$C{r}"), fN, ACC)
            r += 1
        w(ws, f"C{r}", "Total (a)" if grp == "PPE" else "Total (b)", fB, bd=B_TB)
        for i, k in enumerate(keys):
            c = CL(4 + i)
            w(ws, f"{c}{r}", f"=SUM({c}{f0}:{c}{max(f0, r - 1)})" if r > f0 else 0, fB, ACC, bd=B_TB)
        cx.PPE_ROWS[("FB", grp)] = r
        r += 2
    cx.FALLBACK_COLS = {k: CL(4 + i) for i, k in enumerate(keys)}
    r = _cwip(cx, ws, r + 1)
    cx.finish_print(ws, r)
