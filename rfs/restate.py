"""Readers for restatement inputs: fixed asset register, MSME interest computations, actuarial (gratuity) reports,
income-tax returns and sanction letters. Best effort - every value is shown in the app for review before it is used."""
import datetime as dt
import os
import re

from .classify import _fixed_asset
from .docreader import extract_tables
from .mapping import LBL
from .regulatory import _parse_date

FA_CLASSES = ["Land", "Buildings", "Furniture and fittings", "Vehicles", "Office equipment", "Electrical installations", "Computers",
              "Computer software"]
# Schedule II useful lives (years) used when the register gives neither a rate nor a life
DEFAULT_LIFE = {"Land": 0, "Buildings": 60, "Furniture and fittings": 10, "Vehicles": 8, "Office equipment": 5, "Electrical installations": 10,
                "Computers": 3, "Computer software": 6}


def _num(v):
    s = str(v or "").replace(",", "").replace("₹", "").replace("Rs.", "").strip()
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    try:
        x = float(s)
    except ValueError:
        return None
    return -x if neg else x


def fa_class(text):
    """Register classification / asset name -> PPE note line."""
    t = re.sub(r"\s+", " ", (text or "").lower())
    t = " " + t.replace("airconditioner", "air conditioner") + " "     # padded so whole-word keywords (' car', ' ups') match
    code = _fixed_asset(t)
    return LBL[code][0] if code in LBL else "Office equipment"


def _find(head, pats, avoid=()):
    for p in pats:
        for j, c in enumerate(head):
            if p in c and not any(a in c for a in avoid):
                return j
    return None


def read_fa_register(path, work_dir=None):
    """Assets from a fixed asset register (Excel / CSV / Word / PDF table): one dict per asset."""
    best = []
    for t in extract_tables(path, work_dir):
        for hi in range(min(25, len(t))):
            head = [re.sub(r"\s+", " ", c.lower()) for c in t[hi]]
            jd = _find(head, ["date of purchase", "put to use", "date of acquisition", "acquisition date", "purchase date", "capitalisation date",
                              "capitalization date", "invoice date", "date"], avoid=("sale", "deletion", "disposal", "sold"))
            jc = _find(head, ["original cost", "gross block", "cost", "amount", "value"],
                       avoid=("residual", "wdv", "written down", "accumulated", "depreciation", "sale", "salvage", "disposal"))
            if jd is None or jc is None:
                continue
            jn1 = _find(head, ["new description"])
            jn2 = _find(head, ["description", "asset name", "name of asset", "name of the asset", "particulars", "asset", "item"], avoid=("new",))
            jk = _find(head, ["fa classification", "classification", "class", "block", "category", "asset group", "group", "type"])
            jr = _find(head, ["depreciation rate", "rate of depreciation", "rate"], avoid=("interest",))
            jl = _find(head, ["useful life"], avoid=("remaining", "balance"))
            jx = _find(head, ["date of sale", "date of deletion", "deletion date", "disposal date", "date of disposal", "sold on", "discarded on",
                              "period of deletion"])
            js = _find(head, ["sale value", "sale proceeds", "sale amount", "disposal value", "sold for", "realisation"])
            rows = []
            for r in t[hi + 1:]:
                g = lambda j: (r[j] if j is not None and j < len(r) else "") or ""
                d = _parse_date(g(jd))
                c = _num(g(jc))
                if not d or not c or c <= 0:
                    continue
                name = (g(jn1) or g(jn2)).strip()
                src = (g(jk) or name).strip()
                rate = _num(g(jr))
                if rate is not None and rate > 1:
                    rate = rate / 100
                life = _num(g(jl))
                xd = _parse_date(g(jx))
                rows.append({"name": name or src, "source_class": src, "class": fa_class(src or name), "date": d.isoformat(), "cost": round(c, 2),
                             "rate": round(rate, 4) if rate else None, "life": life if life else None,
                             "del_date": xd.isoformat() if xd else None, "sale": _num(g(js)) or 0.0})
            if len(rows) > len(best):
                best = rows
            if rows:
                break
    return best


def _msme_xlsx(path):
    """Excel computation: the sheet with an 'Interest due' column; interest = sum over the bill rows (rows with a party name),
    period end = the latest date above the heading row (else the year in the file name)"""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = sorted(wb.worksheets, key=lambda w: 0 if "interest" in w.title.lower() else 1 if "msme" in w.title.lower() else 2)
    for ws in sheets:
        rows = [list(r) for r in ws.iter_rows(max_col=min(ws.max_column, 40), values_only=True)]
        for hi, r in enumerate(rows[:20]):
            head = [re.sub(r"\s+", " ", str(c or "")).strip().lower() for c in r]
            j = next((i for i, h in enumerate(head) if re.search(r"interest (due|payable|amount|u/s ?16)", h)), None)
            if j is None:
                j = next((i for i, h in enumerate(head) if h == "interest"), None)
            if j is None:
                continue
            pj = next((i for i, h in enumerate(head) if re.search(r"party|supplier|vendor|creditor|name", h)), None)
            end = None
            for rr in rows[:hi]:
                for c in rr:
                    d = c.date() if isinstance(c, dt.datetime) else c if isinstance(c, dt.date) else _parse_date(c) if isinstance(c, str) else None
                    if d and (end is None or d > end):
                        end = d
            tot, n = 0.0, 0
            for rr in rows[hi + 1:]:
                if pj is not None and not str((rr[pj] if pj < len(rr) else None) or "").strip():
                    continue                       # totals / notes below the bills
                if pj is not None and re.search(r"^(grand )?total", str(rr[pj]).strip(), re.I):
                    continue
                v = rr[j] if j < len(rr) else None
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    tot += v; n += 1
            if n:
                return round(tot, 2), end
    return None, None


def read_msme_interest(path, work_dir=None):
    """(interest for the period in Rs, period end date or None) from an MSME interest computation workbook."""
    if str(path).lower().endswith((".xlsx", ".xlsm")):
        try:
            total, end = _msme_xlsx(path)
        except Exception:
            total, end = None, None
        if total is not None:
            if end is None:
                m = re.search(r"(?<!\d)(20)?(\d{2})\s*[-_ ]\s*(\d{2})(?!\d)", os.path.basename(path))
                if m:
                    end = dt.date(int(m.group(3)) + 2000, 3, 31)
            return total, end
    total, end = None, None
    for t in extract_tables(path, work_dir):
        for r in t[:8]:                       # the computation date is usually in the heading
            for c in r:
                d = _parse_date(c)
                if d and (end is None or d > end):
                    end = d
        for hi in range(min(15, len(t))):
            head = [c.lower() for c in t[hi]]
            j = _find(head, ["interest due", "interest payable", "interest amount", "interest u/s 16", "interest"], avoid=("annual", "rate", "%"))
            if j is None:
                continue
            vals = [_num(r[j]) for r in t[hi + 1:] if j < len(r)]
            vals = [v for v in vals if v is not None]
            if not vals:
                continue
            if len(vals) > 1 and abs(vals[-1] - sum(vals[:-1])) < 1:
                total = vals[-1]                  # the last figure is the column total
            else:
                total = sum(vals)
            break
        if total is not None:
            break
    if end is None:
        m = re.search(r"(?<!\d)(20)?(\d{2})\s*[-_ ]\s*(\d{2})(?!\d)", os.path.basename(path))
        if m:
            y = int(m.group(3)) + 2000
            end = dt.date(y, 3, 31)
    return (round(total, 2) if total is not None else None), end


_D = r"(\d{1,2}(?:st|nd|rd|th)?[\s./-]+(?:\w+|\d{1,2})[\s.,/-]+\d{4})"


def read_gratuity(text):
    """{valuation date: defined benefit obligation (Rs)} from an actuarial valuation report."""
    t = re.sub(r"[ \t]+", " ", text or "")
    vd = None
    m = re.search(r"(?:valuation date|as (?:at|on)|accounting (?:period|year) ending)\s*[:\-]?\s*" + _D, t, re.I)
    if m:
        vd = _parse_date(m.group(1)) or _parse_any(m.group(1))
    amt = None
    for pat in [r"present value of (?:the )?(?:defined benefit )?obligation (?:at|as at) (?:the )?end of (?:the )?(?:year|period)[^\d\n]{0,40}([\d,]+(?:\.\d+)?)",
                r"closing (?:defined benefit )?obligation[^\d\n]{0,40}([\d,]+(?:\.\d+)?)",
                r"(?:net )?liability (?:recognised|to be recognised) in (?:the )?balance sheet[^\d\n]{0,40}([\d,]+(?:\.\d+)?)",
                r"defined benefit obligation[^\d\n]{0,40}([\d,]+(?:\.\d+)?)"]:
        m = re.search(pat, t, re.I)
        if m and (_num(m.group(1)) or 0) >= 100:
            amt = _num(m.group(1))
            break
    return {vd.isoformat(): amt} if vd and amt else {}


def _parse_any(s):
    s = re.sub(r"(st|nd|rd|th)\b", "", s or "")
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%d %B %Y", "%d %b %Y", "%d-%b-%Y"):
        try:
            return dt.datetime.strptime(s.strip(" ,."), f).date()
        except ValueError:
            pass
    return None


def read_tax(text):
    """(tax per return in Rs, financial year end) from an ITR acknowledgement / computation of total income."""
    t = re.sub(r"[ \t]+", " ", text or "")
    end = None
    m = re.search(r"assessment year\s*[:\-]?\s*(20\d{2})\s*[-–]\s*(\d{2,4})", t, re.I)
    if m:
        end = dt.date(int(m.group(1)), 3, 31)          # AY 2024-25 -> year ended March 31, 2024
    amt = None
    for pat in [r"tax payable on total income", r"total tax(?:,| and) interest(?: and fee)? payable", r"net tax (?:liability|payable)",
                r"total tax liability", r"tax payable"]:
        m = re.search(pat + r"[^\d\n]{0,40}([\d,]+(?:\.\d+)?)", t, re.I)
        if m:
            amt = _num(m.group(1))
            break
    return amt, end


_LENDERS = ["State Bank of India", "HDFC Bank", "ICICI Bank", "Axis Bank", "Kotak Mahindra Bank", "Yes Bank", "IndusInd Bank", "IDFC First Bank",
            "Bank of Baroda", "Canara Bank", "Union Bank of India", "Indian Bank", "Punjab National Bank", "Bank of India", "Central Bank of India",
            "Indian Overseas Bank", "Federal Bank", "South Indian Bank", "Karur Vysya Bank", "City Union Bank", "Tamilnad Mercantile Bank", "RBL Bank",
            "DCB Bank", "Bandhan Bank", "AU Small Finance Bank", "Equitas Small Finance Bank", "Ujjivan Small Finance Bank", "Unity Small Finance Bank",
            "Bajaj Finance", "Sundaram Finance", "Tata Capital", "Cholamandalam Investment and Finance", "Mahindra & Mahindra Financial Services",
            "Hero FinCorp", "Aditya Birla Finance", "SIDBI", "Small Industries Development Bank of India", "Ugro Capital", "Shriram Finance"]
_FACILITIES = [("channel finan", "Channel finance"), ("dealer finan", "Dealer finance"), ("bill discount", "Bill discounting"),
               ("invoice discount", "Invoice discounting"), ("working capital demand loan", "Working capital demand loan"), ("wcdl", "Working capital demand loan"),
               ("cash credit", "Cash credit"), ("overdraft", "Overdraft"), ("od limit", "Overdraft"), ("term loan", "Term loan"),
               ("vehicle loan", "Vehicle loan"), ("car loan", "Vehicle loan"), ("equipment loan", "Equipment loan"), ("bank guarantee", "Bank guarantee"),
               ("letter of credit", "Letter of credit"), ("business loan", "Business loan")]


def _amount_rs(num, unit):
    v = _num(num)
    if v is None:
        return None
    u = (unit or "").lower()
    if u.startswith("cr"):
        v *= 1e7
    elif u.startswith(("lak", "lac", "l")):
        v *= 1e5
    return v


def read_sanction(text, name=""):
    """Principal terms from a sanction letter: lender, facility, sanctioned amount (Rs), rate, tenure, security, guarantee."""
    t = re.sub(r"[ \t]+", " ", text or "")
    low = (t + " " + name).lower()
    lender = next((l for l in _LENDERS if l.lower() in low), "")
    if not lender:
        m = re.search(r"([A-Z][A-Za-z&. ]{2,40}(?:Bank|Finance|Financial Services|Capital|Fincorp|FinCorp))(?: Limited| Ltd\.?)?", t)
        lender = m.group(1).strip() if m else ""
    facility = next((f for k, f in _FACILITIES if k in low), "")
    amts = []
    for m in re.finditer(r"(?:Rs\.?|₹|INR)\s*([\d,]+(?:\.\d+)?)\s*(crores?|crs?\.?|cr\b|lakhs?|lacs?)?", t + " " + name, re.I):
        a = _amount_rs(m.group(1), m.group(2))
        if a:
            ctx = (t + " " + name)[max(0, m.start() - 80): m.start()].lower()
            amts.append((("limit" in ctx or "sanction" in ctx or "facility" in ctx or "amount" in ctx), a))
    m = re.search(r"([\d.]+)\s*(crores?|crs?\.?|lakhs?|lacs?)\b", name, re.I)
    if m:
        amts.append((True, _amount_rs(m.group(1), m.group(2))))
    sanctioned = max((a for k, a in amts if k), default=max((a for _, a in amts), default=None))
    rate = ""
    m = re.search(r"((?:repo|mclr|eblr|rllr|t-?bill|base rate|plr)[^\n]{0,60}?\d+(?:\.\d+)?\s*%[^\n]{0,40})", t, re.I) or \
        re.search(r"(?:rate of interest|interest rate|roi)\s*[:\-]?\s*([^\n]{3,80}?%[^\n]{0,30})", t, re.I) or \
        re.search(r"(\d{1,2}(?:\.\d{1,2})?\s*%\s*(?:p\.?\s*a\.?|per annum))", t, re.I)
    if m:
        rate = re.sub(r"\s+", " ", m.group(1)).strip(" .,:;")
    tenure = ""
    m = re.search(r"(?:tenure|tenor|period of (?:the )?(?:loan|facility)|validity|repayment)\s*[:\-]?\s*([^\n]{3,90})", t, re.I) or \
        re.search(r"(\d{1,3}\s*(?:months|equated monthly instal\w+|emis?))", t, re.I)
    if m:
        tenure = re.sub(r"\s+", " ", m.group(1)).strip(" .,:;")[:90]
    sec = []
    for pat in [r"primary security\s*[:\-]?\s*([^\n]{8,220})", r"collateral(?: security)?\s*[:\-]?\s*([^\n]{8,220})", r"security\s*[:\-]\s*([^\n]{8,220})"]:
        m = re.search(pat, t, re.I)
        if m:
            sec.append(re.sub(r"\s+", " ", m.group(1)).strip(" .,:;"))
    m = re.search(r"((?:personal|corporate) guarantee[^\n]{0,180})", t, re.I) or re.search(r"guarantee\s*[:\-]\s*([^\n]{5,180})", t, re.I)
    guarantee = re.sub(r"\s+", " ", m.group(1)).strip(" .,:;") if m else ""
    return {"lender": lender, "facility": facility, "sanctioned": round(sanctioned, 2) if sanctioned else None, "rate": rate, "tenure": tenure,
            "security": "; ".join(dict.fromkeys(sec))[:300], "guarantee": guarantee[:200], "ledger": ""}
