"""App side of the final-format builder (book.py): input screens for the notes / annexures, editable texts,
the client fixed asset register mapping, and make_input() which turns the app's project into book.build's input."""
import datetime as dt
import openpyxl
import pandas as pd
import streamlit as st
from openpyxl.utils import get_column_letter as CL

from . import mapping, book_masters as BM

SLOT2P = {"FYS": "S", "FY26": "Y1", "FY25": "Y2", "FY24": "Y3"}
DK2P = {"dS": "S", "d26": "Y1", "d25": "Y2", "d24": "Y3", "d23": "OPEN"}


_SEQ = [0]


def _k(key):
    """widget key unique within a run (several tables share a store key name, e.g. 'other')"""
    _SEQ[0] += 1
    return f"nt_{key}_{_SEQ[0]}_{st.session_state.get('_fill_n', 0)}"


def _f(v):
    try:
        return None if v in (None, "") or (isinstance(v, float) and pd.isna(v)) else float(v)
    except (TypeError, ValueError):
        return None


def per_table(store, key, label, pkeys, plab, num=True, help_=None):
    """one editable column keyed by period keys (Rs); stored in store[key] = {p: value}"""
    cur = store.setdefault(key, {})
    vals = [cur.get(k) for k in pkeys]
    if not num:                                # text column: editor needs strings
        vals = [None if v is None or (isinstance(v, float) and pd.isna(v)) else (f"{v:g}" if isinstance(v, float) else str(v)) for v in vals]
    df = pd.DataFrame({"Period / date": [plab(k) for k in pkeys], label: pd.Series(vals, dtype=object if not num else "float64")})
    cfg = {label: st.column_config.NumberColumn(label, format="%.2f", help=help_) if num else st.column_config.TextColumn(label, help=help_)}
    e = st.data_editor(df, hide_index=True, width="stretch", key=_k(key), disabled=["Period / date"], column_config=cfg)
    store[key] = {k: (None if (v is None or (isinstance(v, float) and pd.isna(v)) or v == "") else (float(v) if num else v))
                  for k, v in zip(pkeys, e[label])}
    return store[key]


def grid(store, key, columns, title=None, help_=None, cfg=None, rows=None):
    """editable table stored as a list of dicts"""
    if title:
        st.caption(title)
    data = store.get(key) or rows or []
    df = pd.DataFrame(data, columns=list(columns))
    cfg = cfg or {}
    for c in df.columns:                      # column types the editor accepts for empty columns
        kind = getattr(cfg.get(c), "get", lambda *_: None)("type_config") or {}
        t = (kind.get("type") if isinstance(kind, dict) else None) or ""
        if c == "date" or t == "date":
            df[c] = pd.to_datetime(df[c], errors="coerce").dt.date
        elif t == "checkbox" or c in ("guarantee", "interest", "fd", "ppe_sale", "promoter"):
            df[c] = df[c].fillna(False).astype(bool)
        elif t == "number" or c in ("amount", "limit_rs", "life", "rate", "S", "Y1", "Y2", "Y3", "OPEN") and t not in ("text",):
            df[c] = pd.to_numeric(df[c], errors="coerce")
        else:
            df[c] = df[c].astype(object).where(df[c].notna(), None)
    e = st.data_editor(df, num_rows="dynamic", width="stretch", key=_k(key), column_config=cfg)
    out = [{k: (None if (isinstance(v, float) and pd.isna(v)) or v is pd.NaT else (v.isoformat() if hasattr(v, "isoformat") else v))
            for k, v in r.items()} for r in e.to_dict("records")]
    store[key] = [r for r in out if any(v not in (None, "", False) for v in r.values())]
    return store[key]


# ======================================================================================== notes tab
def notes_tab(ss, P, SLOTS, plabel, years, ready, tb_ledgers):
    """inputs for the notes and annexures (Annexure VI)"""
    _SEQ[0] = 0
    N = ss.project.setdefault("notes", {})
    pk = [SLOT2P[s] for s, _ in SLOTS]
    plab = {SLOT2P[s]: plabel(s) for s, _ in SLOTS}
    open_d = years["d24"].replace(year=years["d24"].year - 1)
    plab["OPEN"] = f"As at {open_d:%d-%b-%Y} (opening)"
    lab = lambda k: plab[k]
    fy = [p for p in pk if p != "S"]
    st.write("Inputs for the notes and annexures. Amounts are in Rs (converted to lakhs in the workbook) unless stated. Leave a table empty "
             "where a disclosure is nil.")

    with st.expander("Groupings - custom masters (added to the note you choose)"):
        notes_opt = list(BM.NOTE_NAMES)
        cm = grid(N, "custom_masters", ["name", "note", "nature"], "Each custom master appears as an extra line in the note chosen; use it in tab 5.",
                  cfg={"name": st.column_config.TextColumn("Master (note line text)", width="large"),
                       "note": st.column_config.SelectboxColumn("Note", options=notes_opt, width="medium"),
                       "nature": st.column_config.SelectboxColumn("Natural balance", options=["Dr", "Cr"])})
        mapping.CUSTOM[:] = [c for c in cm if c.get("name") and c.get("note")]

    with st.expander("Share capital - authorised capital history, promoters (shareholders in tab 6)"):
        grid(N, "auth_history", ["date", "amount", "remarks"], "Authorised capital: one row on incorporation and one per increase (Rs).",
             cfg={"date": st.column_config.DateColumn("Date of resolution", format="DD-MM-YYYY"),
                  "amount": st.column_config.NumberColumn("Increase (Rs)", format="%.0f"), "remarks": st.column_config.TextColumn("Remarks", width="large")},
             rows=[{"date": None, "amount": P("authorised_capital"), "remarks": "On incorporation"}] if P("authorised_capital") else None)
        grid(N, "reserve_movements", ["label"] + pk, "Other movements in Reserves and surplus (e.g. bonus issue, dividend) - Rs, negative to reduce:",
             cfg={"label": st.column_config.TextColumn("Particulars", width="large")})

    with st.expander("Borrowings - current maturities, terms (Note 5 / 8 and the terms annexure)"):
        st.caption("Current maturities of long-term borrowings at each date (Rs) - shown under short-term borrowings. Type 'all' when the whole "
                   "term loan is repayable within twelve months.")
        per_table(N, "current_maturities", "Current maturities (Rs or 'all')", pk + ["OPEN"], lab, num=False)
        grid(N, "borrowings", ["lender", "nature", "date", "limit_rs", "roi", "tenor", "security", "ledgers", "guarantee"],
             "Facilities (one row each). Ledgers: TB ledger names whose balances are the outstanding (comma separated).",
             cfg={"lender": "Lender", "nature": st.column_config.TextColumn("Nature of facility", width="medium"),
                  "date": st.column_config.DateColumn("Date of sanction", format="DD-MM-YYYY"),
                  "limit_rs": st.column_config.NumberColumn("Sanctioned limit (Rs)", format="%.0f"), "roi": "Rate of interest",
                  "tenor": "Tenor / repayment", "security": st.column_config.TextColumn("Security / guarantee", width="large"),
                  "ledgers": st.column_config.TextColumn("TB ledgers", width="large", help="Available: " + ", ".join(tb_ledgers.get("borrow", [])[:60])),
                  "guarantee": st.column_config.CheckboxColumn("Directors' personal guarantee")})
        grid(N, "principal_repaid", ["label"] + pk, "Principal repaid on long-term borrowings (Rs) - for the debt service coverage ratio (one row):")

    with st.expander("Other income - ledger-wise lines (Note on other income)"):
        if not N.get("oi_lines") and tb_ledgers.get("oi"):
            N["oi_lines"] = [{"label": l, "master": m, "ledgers": l, "interest": m == "Interest income", "fd": m == "Interest income" and "fd" in l.lower(),
                              "ppe_sale": False} for m, l in tb_ledgers["oi"]]
        grid(N, "oi_lines", ["label", "master", "ledgers", "interest", "fd", "ppe_sale"],
             "One line per item of other income; Ledgers = TB ledger names (comma separated). Tick interest (cash flow / tax), FD interest "
             "(return on investment) and profit on sale of PPE (cash flow).",
             cfg={"label": st.column_config.TextColumn("Line in the note", width="large"),
                  "master": st.column_config.SelectboxColumn("Master", options=["Interest income", "Other non-operating income"]),
                  "ledgers": st.column_config.TextColumn("TB ledgers", width="large"), "interest": st.column_config.CheckboxColumn("Interest"),
                  "fd": st.column_config.CheckboxColumn("FD interest"), "ppe_sale": st.column_config.CheckboxColumn("Profit on sale of PPE")})
        N["ppe_loss_ledgers"] = st.text_input("Ledgers of loss on sale / write off of PPE in the books (comma separated) - cash flow",
                                              N.get("ppe_loss_ledgers", ""), key="nt_ppeloss")

    with st.expander("Related parties (Note on related party transactions)"):
        ppl = [d for d in P("directors", []) if d.get("name")]
        st.caption("Nature of relationship of each person in the Company tab (same order):")
        kn = N.setdefault("kmp_nature", {})
        df = pd.DataFrame({"Name": [d["name"] for d in ppl], "Nature of relationship": [kn.get(d["name"], "") for d in ppl]})
        e = st.data_editor(df, hide_index=True, width="stretch", key="nt_kmp", disabled=["Name"])
        N["kmp_nature"] = dict(zip(e["Name"], e["Nature of relationship"]))
        grid(N, "remuneration", ["name"] + pk, "Directors' remuneration by director (Rs). Leave a period blank to show the total as 'not allocated'.",
             cfg={"name": st.column_config.TextColumn("Director", width="medium")})
        grid(N, "rpt_entities", ["name", "relation", "sales_S", "sales_Y1", "sales_Y2", "sales_Y3", "tr_ledger", "tp_ledger"],
             "Entities in which directors are interested: sales (Rs) and the TB ledgers of the receivable / payable balance.",
             cfg={"name": st.column_config.TextColumn("Entity", width="medium"), "relation": "Relationship",
                  "tr_ledger": "Receivable ledger", "tp_ledger": "Payable ledger"})
        grid(N, "loan_overrides", ["kind", "director"] + pk, "Director loans as disclosed in the audited statements, where they differ from the TB "
             "movements (Rs) - kind: accepted / repaid.",
             cfg={"kind": st.column_config.SelectboxColumn("Kind", options=["accepted", "repaid"])})

    with st.expander("Contingent liabilities and commitments"):
        for k, lbl in (("claims", "Claims not acknowledged as debt (Rs)"), ("guarantees", "Bank guarantees issued on behalf of the Company (Rs)"),
                       ("other", "Other money contingently liable (Rs)"), ("capital_commitments", "Capital commitments (Rs)"),
                       ("other_commitments", "Other commitments (Rs)")):
            per_table(N.setdefault("contingent", {}), k, lbl, pk, lab)

    with st.expander("Employee benefits - gratuity disclosure (actuarial valuation)"):
        G = N.setdefault("gratuity", {})
        G["actuary"] = st.text_input("Actuary", G.get("actuary", ""), key="nt_actuary")
        dflt_ass = [{"label": "Discount rate (per annum)"}, {"label": "Expected rate of salary increase (per annum)"}, {"label": "Attrition rate (per annum)"},
                    {"label": "Retirement age"}, {"label": "Mortality table"}, {"label": "Number of employees covered"}]
        grid(G, "assumptions", ["label"] + fy, "a) Actuarial assumptions (rates as decimals, e.g. 0.0735; text allowed):",
             cfg={"label": st.column_config.TextColumn("Assumption", width="large"), **{p: st.column_config.TextColumn(plab[p]) for p in fy}},
             rows=dflt_ass)
        st.caption("b) Change in the obligation (Rs). Closing: leave blank to add up; if given, the actuarial (gain) / loss is the balancing figure.")
        D = G.setdefault("dbo", {})
        rows = [("open", "Opening obligation"), ("service", "Current service cost"), ("interest", "Interest cost"), ("paid", "Benefits paid"),
                ("actuarial", "Actuarial (gain) / loss"), ("closing", "Closing obligation (optional)"), ("direct_paid", "Gratuity paid directly (charged to P&L)"),
                ("experience", "Experience adjustment - loss / (gain)")]
        df = pd.DataFrame([[lbl] + [(D.get(k) or [None] * 3)[i] for i in range(len(fy))] for k, lbl in rows], columns=["Item"] + [plab[p] for p in fy])
        e = st.data_editor(df, hide_index=True, width="stretch", key="nt_dbo", disabled=["Item"])
        for (k, _), (_, r) in zip(rows, e.iterrows()):
            D[k] = [_f(r[plab[p]]) for p in fy]
        st.caption("Current portion of the obligation at each date (Rs) - short-term provision; the rest is non-current:")
        per_table(G, "current", "Current portion (Rs)", pk + ["OPEN"], lab)
        st.caption("Sensitivity of the obligation (Rs) at the latest year end:")
        sv = G.setdefault("sensitivity", [[None, None]] * 3)
        df = pd.DataFrame(sv, columns=["Increase by 100 bp", "Decrease by 100 bp"], index=["Discount rate", "Salary escalation", "Attrition rate"])
        e = st.data_editor(df, width="stretch", key="nt_sens")
        G["sensitivity"] = [[_f(a), _f(b)] for a, b in e.values.tolist()]

    with st.expander("Tax computation (statement of tax shelter) and returns"):
        T = N.setdefault("tax", {})
        per_table(T, "it_dep", "Depreciation as per the Income-tax Act (Rs)", pk, lab)
        grid(T, "other_rows", ["label"] + pk, "Other timing differences (Rs; disallowance +, allowance -), e.g. 43B(h), 40(a)(ia):",
             cfg={"label": st.column_config.TextColumn("Particulars", width="large")},
             rows=[{"label": "Disallowance u/s 40(a)(ia) of the Income-tax Act, 1961"},
                   {"label": "Disallowance u/s 43B(h) - dues to micro and small enterprises"},
                   {"label": "Allowance u/s 43B(h) - dues to micro and small enterprises (disallowed in the earlier year)"}])
        grid(T, "permanent", ["label"] + pk, "Other permanent differences (Rs):", cfg={"label": st.column_config.TextColumn("Particulars", width="large")})
        st.caption("Return of income details:")
        itr = T.setdefault("itr", {})
        items = ["Assessment year", "Acknowledgement number", "Date of filing", "Total income (Rs)", "Tax payable (Rs)", "Taxes paid (Rs)", "Refund claimed (Rs)"]
        df = pd.DataFrame([[it] + [(itr.get(p) or [None] * 7)[i] for p in pk] for i, it in enumerate(items)], columns=["Item"] + [plab[p] for p in pk])
        e = st.data_editor(df.astype(object), hide_index=True, width="stretch", key="nt_itr", disabled=["Item"])
        T["itr"] = {p: [None if (isinstance(v, float) and pd.isna(v)) else v for v in e[plab[p]].tolist()] for p in pk}

    with st.expander("Additional disclosures - auditor fees, foreign currency, loans to related parties"):
        A = N.setdefault("auditor_fees", {})
        per_table(A, "tax_audit", "Tax audit fee included in payment to auditors (Rs)", pk, lab)
        per_table(A, "other", "Other services (Rs)", pk, lab)
        FX = N.setdefault("foreign", {})
        st.caption("Foreign currency (Rs, or text such as 'Not available'):")
        for k, l in (("imports", "CIF value of imports"), ("exports", "FOB value of exports"), ("expenditure", "Expenditure in foreign currency")):
            per_table(FX, k, l, pk, lab, num=False)
        FX["expenditure_label"] = st.text_input("Label for the expenditure line", FX.get("expenditure_label") or "C. Expenditure in foreign currency",
                                                key="nt_fxlab")

    with st.expander("Useful lives table (Annexure IV)"):
        from .book_statements import USEFUL_LIVES
        grid(N, "useful_lives", ["class", "life", "rate"], cfg={"class": st.column_config.TextColumn("Class", width="large")},
             rows=[{"class": c, "life": l, "rate": r} for c, l, r in USEFUL_LIVES])


# ======================================================================================== texts tab
def texts_tab(ss):
    T = ss.project.setdefault("texts", {})
    used = ss.get("text_defaults") or {}
    if not used:
        st.info("Generate the RFS once (tab Generate) - every paragraph then appears here with its default wording, ready to edit.")
        return
    st.write("Every paragraph of the RFS with its default wording (from the inputs). Edit any of them and generate again; clear a box to go back "
             "to the default. Use {NOTE:BT} style tokens for note numbers: BT borrowing terms, EB employee benefits, RPT, CL contingent, PPE.")
    q = st.text_input("Filter", key="tx_filter")
    for k, d in used.items():
        lbl = d.get("label") or k
        if q and q.lower() not in (lbl + k + str(d.get("default"))).lower():
            continue
        dflt = str(d.get("default") or "")
        if dflt.startswith("="):
            st.caption(f"{lbl}: built from a formula (note numbers / face value) - type a fixed text to replace it.")
        val = st.text_area(lbl, T.get(k) or ("" if dflt.startswith("=") else dflt), key="tx_" + k, height=min(300, 60 + len(dflt) // 3))
        if val and val != dflt:
            T[k] = val
        else:
            T.pop(k, None)


# ======================================================================================== client register mapping
def fa_mapped_ui(ss, RST, SLOTS, plabel, path):
    """map the columns of the client's own fixed asset register"""
    M = RST.setdefault("fa_map", {})
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        st.error(f"Could not open the register: {e}")
        return
    sheet = st.selectbox("Sheet", wb.sheetnames, index=wb.sheetnames.index(M["sheet"]) if M.get("sheet") in wb.sheetnames else 0, key="fam_sheet")
    ws = wb[sheet]
    c1, c2, c3 = st.columns(3)
    M["sheet"] = sheet
    M["header_row"] = c1.number_input("Header row", 1, 50, int(M.get("header_row") or 5), key="fam_h")
    M["first_row"] = c2.number_input("First asset row", 1, 5000, int(M.get("first_row") or M["header_row"] + 1), key="fam_f")
    M["last_row"] = c3.number_input("Last asset row", 1, 20000, int(M.get("last_row") or ws.max_row), key="fam_l")
    heads = {}
    for row in ws.iter_rows(min_row=M["header_row"], max_row=M["header_row"]):
        for c in row:
            if c.value not in (None, ""):
                heads[c.column_letter] = f"{c.column_letter} - {str(c.value)[:40]}"
    opts = [""] + list(heads)
    fmt = lambda c: heads.get(c, c or "-")
    sel = lambda lbl, key: st.selectbox(lbl, opts, index=opts.index(M.get(key)) if M.get(key) in opts else 0, format_func=fmt, key="fam_" + key)
    a, b, c, d = st.columns(4)
    with a: M["class"] = sel("FA classification column", "class")
    with b: M["cost"] = sel("Cost column", "cost")
    with c: M["date"] = sel("Date of purchase column", "date")
    with d: M["del"] = sel("Date of deletion column (optional)", "del")
    st.caption("For each date: WDV column; for each period: depreciation column (restated). For a stub period whose register depreciation is "
               "for the full period, give the register depreciation column too.")
    per = M.setdefault("periods", {})
    keys = ["OPEN"] + [SLOT2P[s] for s, _ in reversed(SLOTS)]
    for k in keys:
        cols = st.columns(4)
        lbl = "Opening" if k == "OPEN" else next(plabel(s) for s, _ in SLOTS if SLOT2P[s] == k)
        cols[0].write(lbl)
        pm = per.setdefault(k, {})
        pm["wdv"] = cols[1].selectbox("WDV at the end", opts, index=opts.index(pm.get("wdv")) if pm.get("wdv") in opts else 0, format_func=fmt, key=f"fam_w{k}")
        if k != "OPEN":
            pm["dep"] = cols[2].selectbox("Depreciation", opts, index=opts.index(pm.get("dep")) if pm.get("dep") in opts else 0, format_func=fmt, key=f"fam_d{k}")
            pm["dep_reg"] = cols[3].selectbox("Register depreciation (if different)", opts,
                                              index=opts.index(pm.get("dep_reg")) if pm.get("dep_reg") in opts else 0, format_func=fmt, key=f"fam_r{k}")
    if M.get("class"):
        vals = sorted({str(r[0]) for r in ws.iter_rows(min_row=M["first_row"], max_row=M["last_row"], min_col=openpyxl.utils.column_index_from_string(M["class"]),
                                                      max_col=openpyxl.utils.column_index_from_string(M["class"]), values_only=True) if r[0] not in (None, "")})
        cm = RST.setdefault("fa_classmap", {})
        classes = [m for m, _ in BM.note_lines("PPE", mapping.CUSTOM) + BM.note_lines("INTANG", mapping.CUSTOM)]
        df = pd.DataFrame({"Register classification": vals, "Class in Note 12": [cm.get(v) or _guess_class(v, classes) for v in vals]})
        e = st.data_editor(df, hide_index=True, width="stretch", key="fam_cm", disabled=["Register classification"],
                           column_config={"Class in Note 12": st.column_config.SelectboxColumn(options=classes, required=True)})
        RST["fa_classmap"] = dict(zip(e["Register classification"], e["Class in Note 12"]))


def _guess_class(v, classes):
    t = v.lower()
    for k, c in (("land", "Land"), ("build", "Buildings"), ("plant", "Plant and machinery"), ("machin", "Plant and machinery"), ("furni", "Furniture and fittings"),
                 ("car", "Vehicles"), ("vehic", "Vehicles"), ("bike", "Vehicles"), ("software", "Computer software"), ("comput", "Computers"),
                 ("laptop", "Computers"), ("electric", "Electrical installations"), ("air", "Office equipment"), ("mobile", "Office equipment"),
                 ("office", "Office equipment")):
        if k in t and c in classes:
            return c
    return None


# ======================================================================================== project -> builder input
def make_input(ss, P, SLOTS, years, stub_end, tb_title):
    """the app's project and trial balances -> input dict of book.build"""
    pr = ss.project
    N = pr.get("notes", {})
    RST = pr.get("restatement", {})
    pk = [SLOT2P[s] for s, _ in SLOTS]
    aud = pr.get("auditor", {})

    def d_(v):
        if v in (None, "") or (isinstance(v, float) and pd.isna(v)):
            return None
        return v.isoformat() if hasattr(v, "isoformat") else str(v)[:10]
    I = {"company": {"cin": pr.get("cin", ""), "cin_note": pr.get("cin_note", ""), "name": pr.get("name", ""), "former": pr.get("former", ""),
                     "address": pr.get("address", ""), "fv": pr.get("face_value", 10), "bonus": pr.get("bonus_factor", 1), "scale": 100000,
                     "tax_rate": pr.get("tax_rate", 0.25168), "place": aud.get("place", ""), "incorporation_date": d_(pr.get("incorporation_date")),
                     "roc": pr.get("roc", ""), "objects": pr.get("objects", ""), "exchange": pr.get("exchange", "")},
         "dates": {"stub": stub_end.isoformat() if stub_end else None, "fy": [years["d26"].isoformat(), years["d25"].isoformat(), years["d24"].isoformat()]},
         "people": [d for d in pr.get("directors", []) if d.get("name")],
         "auth_capital": [{"date": d_(a.get("date")), "amount": a.get("amount"), "remarks": a.get("remarks", "")} for a in N.get("auth_history", [])]
         or ([{"date": d_(pr.get("incorporation_date")), "amount": pr.get("authorised_capital"), "remarks": "On incorporation"}] if pr.get("authorised_capital") else []),
         "auditor": {k: aud.get(k, "") for k in ("firm", "frn", "peer", "partner", "mno", "sign_date", "udin")},
         "stat_auditor": pr.get("stat_auditor", {}),
         "pending": [x for x in str(pr.get("pending", "")).splitlines() if x.strip()],
         "closing_stock": {"Y1": pr.get("closing_stock"), "S": pr.get("stub_closing_stock")},
         "custom_masters": list(mapping.CUSTOM), "texts": dict(pr.get("texts", {}))}
    if pr.get("background") and "sap_background" not in I["texts"]:
        I["texts"]["sap_background"] = pr["background"]
    # trial balances
    I["tbs"] = {}
    from .book_core import fy_label as _fyl
    for s, _ in SLOTS:
        rows = []
        dkey = {"FYS": stub_end, "FY26": years["d26"], "FY25": years["d25"], "FY24": years["d24"]}[s]
        pre = "@TB " + _fyl(dkey) + "@"                  # mapping imported from an RFS workbook for this year's TB
        ovs = dict(ss.overrides)
        ovs.update({"@" + k[len(pre):]: v for k, v in ss.overrides.items() if k.startswith(pre)})
        for x in ss.tb[s]:
            r = {"name": x["name"], "op": x["v"][0], "dr": x["v"][1], "cr": x["v"][2], "cl": x["v"][3]}
            if x.get("type") == "X":
                r["dup"] = True
            elif not x.get("inc"):
                r["group"] = True
            else:
                m = mapping.master_for(x, 3, ovs)
                r["master"] = m[0] if m else None
                if SLOT2P[s] == "Y3":
                    ovo = dict(ovs)
                    ovo.update({"@" + k[6:]: v for k, v in ss.overrides.items() if k.startswith("@OPEN@")})
                    mo = mapping.master_for(x, 0, ovo)
                    r["op_master"] = mo[0] if mo else r["master"]
            rows.append(r)
        I["tbs"][SLOT2P[s]] = {"title": tb_title(s), "rows": rows}
    # adjustments (tab 8) and automatic entries (tab 7)
    I["adjustments"] = [{"date": d_(e["date"]), "narration": e.get("narration", ""),
                         "lines": [{"master": l["master"], "particulars": l.get("particulars") or l["master"], "dr": l["dr"], "cr": l["cr"]} for l in e["legs"]]}
                        for e in getattr(ss, "entries", [])]
    rst = {}
    dbo = {DK2P[k]: v for k, v in (RST.get("dbo") or {}).items() if v not in (None, "") and k in DK2P}
    if dbo:
        rst["gratuity"] = {"dbo": dbo, "direct_reserve": {DK2P[k]: v for k, v in (RST.get("gdir") or {}).items() if v}}
    if any((RST.get("msme") or {}).values()):
        rst["msme"] = {DK2P[k]: v for k, v in RST["msme"].items() if v not in (None, "")}
    tx = {DK2P[k]: v for k, v in (RST.get("tax") or {}).items() if v not in (None, "")}
    if stub_end and RST.get("tax_stub_computed"):
        tx["S"] = "computed"
    if tx:
        rst["tax"] = tx
    I["restatement"] = rst
    has_register = (RST.get("fa_mode") == "mapped" and RST.get("fa_map_path")) or (RST.get("fa_mode") != "mapped" and RST.get("fa_assets"))
    if has_register:          # the register generates the PPE entries: drop typed / imported ones marked 'FA Register'
        I["adjustments"] = [e for e in I["adjustments"] if "fa register" not in (e.get("narration") or "").lower()]
    # fixed assets
    fad = {DK2P[k]: v for k, v in (RST.get("fa_direct") or {}).items() if v}
    if RST.get("fa_mode") == "mapped" and RST.get("fa_map_path"):
        M = dict(RST.get("fa_map") or {})
        M["periods"] = {k: {kk: vv for kk, vv in v.items() if vv} for k, v in (M.get("periods") or {}).items() if k == "OPEN" or k in pk}
        I["fa"] = {"mode": "mapped", "source": RST["fa_map_path"], "sheet": M.get("sheet"), "map": M, "class_map": RST.get("fa_classmap") or {},
                   "direct_reserve": fad}
    elif RST.get("fa_assets"):
        I["fa"] = {"mode": "standard", "direct_reserve": fad,
                   "assets": [{"name": a.get("name"), "classification": a.get("source_class") or a.get("class"), "class": a.get("class"),
                               "cost": a.get("cost"), "date": a.get("date"), "rate": a.get("rate"), "life": a.get("life"),
                               "del_date": a.get("del_date"), "sale": a.get("sale"), "wdv_open": a.get("wdv_open")} for a in RST["fa_assets"] if a.get("date")]}
    # notes inputs
    I["shareholders"] = [{"name": h.get("name"), "promoter": bool(h.get("promoter")), **{SLOT2P[s]: h.get({"FYS": "sS", "FY26": "s26", "FY25": "s25", "FY24": "s24"}[s])
                                                                                       for s, _ in SLOTS}} for h in pr.get("shareholders", []) if h.get("name")]
    ag = lambda key: {DK2P[k]: v for k, v in (pr.get(key) or {}).items() if k in DK2P and any(any(x for x in r) for r in v)}
    I["debtor_ageing"], I["creditor_ageing"] = ag("ageing_dr"), ag("ageing_cr")
    cm = N.get("current_maturities") or {}
    I["current_maturities"] = {k: ("all" if str(v).strip().lower() == "all" else _f(v)) for k, v in cm.items() if v not in (None, "")}
    I["borrowings"] = [{"lender": b.get("lender"), "nature": b.get("nature"), "date": d_(b.get("date")),
                        "limit": (_f(b.get("limit_rs")) or 0) / 100000 if _f(b.get("limit_rs")) else None, "roi": b.get("roi"), "tenor": b.get("tenor"),
                        "security": b.get("security"), "ledgers": [l.strip() for l in str(b.get("ledgers") or "").split(",") if l.strip()],
                        "guarantee": bool(b.get("guarantee"))} for b in N.get("borrowings", [])]
    if N.get("oi_lines"):
        I["other_income_lines"] = [{**l, "ledgers": [x.strip() for x in str(l.get("ledgers") or "").split(",") if x.strip()]} for l in N["oi_lines"] if l.get("label")]
    I["cfs"] = {"ppe_loss_ledgers": [x.strip() for x in str(N.get("ppe_loss_ledgers", "")).split(",") if x.strip()]}
    pr_ = (N.get("principal_repaid") or [{}])[0] if N.get("principal_repaid") else {}
    I["analytical"] = {"principal_repaid": {p: pr_.get(p) for p in pk}}
    I["reserve_movements"] = [m for m in N.get("reserve_movements", []) if m.get("label")]
    kn = N.get("kmp_nature", {})
    ov = {"acc": {}, "rep": {}}
    for o in N.get("loan_overrides", []):
        if o.get("director"):
            ov["acc" if o.get("kind") == "accepted" else "rep"][o["director"]] = {p: o.get(p) for p in pk}
    I["rpt"] = {"kmp_nature": [kn.get(d.get("name"), "") for d in I["people"]],
                "remuneration": [r for r in N.get("remuneration", []) if r.get("name")],
                "entities": [{"name": e["name"], "relation": e.get("relation"), "sales": {p: e.get("sales_" + p) for p in pk},
                              "tr_ledger": e.get("tr_ledger"), "tp_ledger": e.get("tp_ledger")} for e in N.get("rpt_entities", []) if e.get("name")],
                "loan_overrides": ov}
    I["contingent"] = N.get("contingent", {})
    G = N.get("gratuity", {})
    if G:
        D = G.get("dbo", {})
        I["gratuity"] = {"actuary": G.get("actuary"), "assumptions": [[a.get("label")] + [_num_or_text(a.get(p)) for p in pk if p != "S"]
                                                                     for a in G.get("assumptions", []) if a.get("label")],
                         "dbo": {k: D.get(k) for k in ("open", "service", "interest", "paid", "actuarial", "closing")},
                         "direct_paid": D.get("direct_paid"), "experience": D.get("experience"), "current": G.get("current", {}),
                         "sensitivity": G.get("sensitivity") if any(any(x is not None for x in r) for r in G.get("sensitivity") or []) else None}
    T = N.get("tax", {})
    itr = {}
    for p_, vals in (T.get("itr") or {}).items():            # amounts typed in the editor come back as text
        itr[p_] = [(_f(str(v).replace(",", "")) if j >= 3 and v not in (None, "") and _f(str(v).replace(",", "")) is not None else v)
                   for j, v in enumerate(vals or [])]
    T = dict(T, itr=itr)
    I["tax"] = {"it_dep": T.get("it_dep", {}), "rows": {r["label"]: {p: r.get(p) for p in pk} for r in T.get("other_rows", []) if r.get("label")},
                "permanent": {r["label"]: {p: r.get(p) for p in pk} for r in T.get("permanent", []) if r.get("label")}, "itr": T.get("itr", {})}
    I["auditor_fees"] = N.get("auditor_fees", {})
    I["foreign"] = N.get("foreign", {})
    if N.get("useful_lives"):
        I["useful_lives"] = [(u["class"], u.get("life"), u.get("rate")) for u in N["useful_lives"] if u.get("class")]
    return I


def _num_or_text(v):
    f = _f(v)
    return f if f is not None else v


def diagnose(calc_path):
    """reasons for a balance sheet difference, read from the check cells of the calculated workbook"""
    wb = openpyxl.load_workbook(calc_path, data_only=True)
    out = []
    bs = wb["BS_1"]
    heads = {c: str(bs.cell(8, c).value or "").replace("\n", " ").replace("As at", "").strip() for c in range(4, 10)}
    diffs = {c: bs.cell(1, c).value for c in range(4, 10) if isinstance(bs.cell(1, c).value, (int, float)) and abs(bs.cell(1, c).value) > 0.01}
    if not diffs:
        return out
    ne = wb["Notes_Equity & Reserves"]
    for r in range(1, ne.max_row + 1):
        if str(ne[f"C{r}"].value or "").startswith("Check: movements in reserves"):
            for c in range(4, 8):
                v = ne.cell(r, c).value
                if isinstance(v, (int, float)) and abs(v) > 0.01:
                    out.append(f"{heads.get(c, '')}: the books moved Rs. {v:,.2f} lakhs directly through Reserves and surplus (not through the profit). "
                               "Restate it in the year it relates to - e.g. gratuity charged to reserves (tab 7, Gratuity: 'charged to reserves in books'), "
                               "PPE aligned with the register through reserves (tab 7, FA register: 'PPE through reserves in books'), or a manual entry in tab 8.")
    for n in wb.sheetnames:
        if n.startswith("TB "):
            ws = wb[n]
            for r in range(ws.max_row, max(1, ws.max_row - 6), -1):
                if str(ws[f"A{r}"].value or "").startswith("Check: TB closing") and isinstance(ws[f"H{r}"].value, (int, float)) and abs(ws[f"H{r}"].value) > 1:
                    out.append(f"{n}: the trial balance itself does not tally (Dr - Cr = Rs. {ws[f'H{r}'].value:,.0f}) - check the TB file / repeated ledgers.")
    if not out:
        out.append("Check the mapping (tab 5) for ledgers marked 'Not considered' and the restatement entries (tab 8): every entry must balance.")
    return out
