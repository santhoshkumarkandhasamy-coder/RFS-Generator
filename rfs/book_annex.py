"""Annexure VI - other notes: RPT, Gratuity and contingencies, Tax Computation, Accounting ratios, Analytical Ratios,
Borrowing Terms, Capitalisation, Additional Disclosures."""
from openpyxl.styles import Font
from .book_core import *          # noqa: F401,F403
from . import book_masters as BM
from .book_notes import NoteSheet, head


def _n(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def rs_f(v, S, rnd=None):
    """typed amount in Rs -> formula in lakhs (blue input)"""
    x = _n(v)
    if x is None:
        return None
    return f"=ROUND({x:g}/{S},{rnd})" if rnd else f"={x:g}/{S}"


def _blue(ws, ref):
    c = ws[ref]
    c.font = Font(name=TN, size=10, color="0000FF", bold=c.font.b)


# ====================================================================== RPT
def rpt(cx):
    I = cx.I
    R = I.get("rpt") or {}
    ns = NoteSheet(cx, "RPT", False, prev_ref=f"'Notes_P&L'!$B${{{{note:EPS}}}}")
    ws = ns.ws
    S = cx.SCALE
    ns.note("RPT", "Related party transactions")
    ns.para(cx.text("rpt_intro", "The Company has entered into the following transactions with related parties for the periods covered under audit, as "
                                 "identified by the management and disclosed in the audited financial statements (except as stated below):"), fI)
    ns.gap()
    ns.sub("(a) List of related parties and nature of relationship")
    r = ns.r
    w(ws, f"C{r}", "Name of the related party", fB, bd=B_TB)
    w(ws, f"D{r}", "Nature of relationship", fB, bd=B_TB)
    ws.merge_cells(f"D{r}:{cx.LASTP}{r}")
    ns.r += 1
    ns.sub("Key Managerial Personnel", fBI)
    people = I.get("people") or []
    nat = R.get("kmp_nature") or []
    for i, p in enumerate(people):
        if not p.get("name"):
            continue
        r = ns.r
        w(ws, f"C{r}", f"=PROPER({MD}!$F${4 + i})")
        dflt = {"Director": "Director and Promoter" if (p.get("category") or "").lower().startswith("promoter") else "Director"}.get(p.get("designation"),
                                                                                                                                     p.get("designation") or "")
        w(ws, f"D{r}", nat[i] if i < len(nat) and nat[i] else dflt)
        ws.merge_cells(f"D{r}:{cx.LASTP}{r}")
        ns.r += 1
    ents = R.get("entities") or []
    if ents:
        ns.sub("Entity in which the directors are interested", fBI)
        for e in ents:
            r = ns.r
            w(ws, f"C{r}", e["name"])
            w(ws, f"D{r}", e.get("relation") or "Entity having common directors")
            ws.merge_cells(f"D{r}:{cx.LASTP}{r}")
            ns.r += 1
    ns.gap()
    ns.sub("(b) Transactions with related parties during the year / period")
    ns.colhead("pl", opening=False)
    # directors' remuneration
    rem = R.get("remuneration") or []
    if cx.nonzero("Directors' remuneration", cx.PK) or rem:
        ns.sub("Directors' remuneration", fBI)
        f0 = ns.r
        for d in rem:
            rr = ns.line(d["name"], lambda p, r, d=d: rs_f(d.get(p), S, 4))
            for p in cx.PK:
                if d.get(p) not in (None, ""):
                    _blue(ws, f"{cx.NC[p]}{rr}")
        miss = [p for p in cx.PK if all(d.get(p) in (None, "") for d in rem)]
        if miss and "pl_dirrem" in cx.refs and "pl_dirrem" not in getattr(cx, "missing", set()):
            lbl = "Not allocated director-wise (" + ", ".join(("period ended " if p == "S" else "year ended ") + long_date(cx.DATE[p]) for p in miss) + ") #"
            ns.line(lbl, lambda p, r: f"='Notes_P&L'!{cx.NC[p]}{{{{pl_dirrem}}}}" if p in miss else None)
        rt = ns.total(f0, ns.r - 1, bd=B_TB)
        if "pl_dirrem" in cx.refs and "pl_dirrem" not in getattr(cx, "missing", set()):
            ns.checks(rt, lambda p: f"=ROUND({cx.NC[p]}{rt}-'Notes_P&L'!{cx.NC[p]}{{{{pl_dirrem}}}},2)")
        ns.gap()
    # loans from directors
    for kind, title in (("acc", "Unsecured loans accepted"), ("rep", "Unsecured loans repaid")):
        mv = cx.DIR_MOV.get(kind) or {}
        extra = [x for x in (R.get("loans_" + kind) or [])]
        if not mv and not extra:
            continue
        ns.sub(title, fBI)
        f0 = ns.r
        ovr = (R.get("loan_overrides") or {}).get(kind) or {}      # figures as disclosed in the audited statements, where they differ
        for m, rr0 in mv.items():
            nm = m.replace(BM.DIRECTOR_PREFIX, "")
            o = ovr.get(nm) or {}
            rr = ns.line(nm, lambda p, r, rr0=rr0, o=o: rs_f(o[p], S, 4) if _n(o.get(p)) is not None else f"=Notes_Liabilities!{cx.NC[p]}{rr0}")
            for p in cx.PK:
                if _n(o.get(p)) is not None:
                    _blue(ws, f"{cx.NC[p]}{rr}")
        for x in extra:
            ns.line(x["name"], lambda p, r, x=x: rs_f(x.get(p), S, 4) or 0)
        ns.total(f0, ns.r - 1, bd=B_TB)
        ns.gap()
    for e in ents:
        for lbl, vals in (e.get("transactions") or {}).items():
            ns.sub(lbl, fBI)
            ns.line(e["name"], lambda p, r, v=vals: (rs_f(v.get(p), S, 4) if _n(v.get(p)) is not None else (v.get(p) or None)))
            ns.gap()
        if e.get("sales"):
            ns.sub("Sale of goods / services", fBI)
            v = e["sales"]
            ns.line(e["name"], lambda p, r, v=v: rs_f(v.get(p), S, 4) if _n(v.get(p)) is not None else "[To be confirmed]")
            ns.gap()
    # balances
    ns.sub("(c) Outstanding balances at the end of the year / period")
    ns.colhead("bs", opening=False)
    if cx.DIR_LOAN_ROW:
        ns.sub("Unsecured loans payable (interest free, repayable on demand)", fBI)
        f0 = ns.r
        for m, rr0 in cx.DIR_LOAN_ROW.items():
            if m.startswith(BM.DIRECTOR_PREFIX):
                ns.line(m.replace(BM.DIRECTOR_PREFIX, ""), lambda p, r, rr0=rr0: f"=Notes_Liabilities!{cx.NC[p]}{rr0}")
        ns.total(f0, ns.r - 1, bd=B_TB)
        ns.gap()
    if cx.nonzero("Directors' remuneration payable"):
        ns.sub("Directors' remuneration payable", fBI)
        ns.line("Remuneration payable to directors", lambda p, r: "=" + cx.master_val(p, "Directors' remuneration payable", nat="Cr"))
        ns.gap()
    for e in ents:
        if e.get("tr_ledger") or e.get("tp_ledger"):
            ns.sub(e["name"], fBI)
            if e.get("tr_ledger"):
                ns.line("Trade receivable", lambda p, r, e=e: "=" + cx.tbl(p, e["tr_ledger"]))
            if e.get("tp_ledger"):
                ns.line("Trade payable", lambda p, r, e=e: "=" + cx.tbl(p, e["tp_ledger"]))
            ns.gap()
    if getattr(cx, "BT_GUAR", None):
        ns.sub("Guarantees", fBI)
        ns.line("Borrowings secured by personal guarantees of the directors (outstanding)",
                lambda p, r: "=" + "+".join(f"'Borrowing Terms'!{cx.BTC[p]}{x}" for x in cx.BT_GUAR))
        ns.gap()
    notes = ["1. Related parties and transactions are as disclosed in the audited financial statements and the books of account."]
    if any(all(d.get(p) in (None, "") for d in rem) for p in cx.PK) and rem:
        notes.append("# Director-wise break-up of remuneration for the period is to be furnished by the management.")
    notes = cx.text_list("rpt_notes", notes, "RPT - notes (one per line)")
    if notes:
        ns.sub("Notes:")
        for t in notes:
            ns.para(t)
    cx.finish_print(ws, ns.r)


# ====================================================================== Gratuity and contingencies
def gratuity(cx):
    I = cx.I
    ns = NoteSheet(cx, "Gratuity and contingencies", True, prev_ref="RPT!$B$8")
    ws = ns.ws
    S = cx.SCALE
    CT = I.get("contingent") or {}
    ns.note("CL", "Contingent liabilities and commitments")
    ns.sub("a) Contingent liabilities")
    ns.colhead("bs", opening=False)
    f0 = ns.r
    for k, lbl in (("claims", "(a) Claims against the Company not acknowledged as debt"), ("guarantees", "(b) Guarantees issued by banks on behalf of the Company"),
                   ("other", "(c) Other money for which the Company is contingently liable")):
        v = CT.get(k) or {}
        rr = ns.line(lbl, lambda p, r, v=v: rs_f(v.get(p), S) or 0, opening=False)
        for p in cx.PK:
            _blue(ws, f"{cx.NC[p]}{rr}")
    ns.total(f0, ns.r - 1, opening=False)
    t = cx.text("cont_note", "")
    if t:
        ns.para(t, fI)
    ns.sub("b) Commitments")
    for lbl, key in (("Estimated amount of contracts remaining to be executed on capital account and not provided for", "capital_commitments"),
                     ("Other commitments", "other_commitments")):
        v = CT.get(key) or {}
        ns.line(lbl, lambda p, r, v=v: rs_f(v.get(p), S) if _n(v.get(p)) else "Nil", opening=False)
    ns.gap()
    # ---- employee benefits
    G = I.get("gratuity") or {}
    ns.note("EB", "Employee benefit plans")
    ns.sub("A. Defined contribution plan")
    ns.colhead("pl", opening=False)
    ns.line("Employer's contribution to provident fund and other funds", lambda p, r: f"='Notes_P&L'!{cx.NC[p]}{{{{pl_pf}}}}"
            if "pl_pf" not in getattr(cx, "missing", set()) else 0, opening=False)
    ns.para(cx.text("eb_dc", "Contributions to provident fund and employees' state insurance are defined contribution plans and are charged to the Statement of "
                             "Profit and Loss in the year in which the services are rendered."), fI, last_col=cx.LASTN)
    ns.gap()
    if getattr(cx, "no_grat", False) and not G:
        cx.setref("gr_cur", ns.r)
        cx.finish_print(ws, ns.r)
        return
    ns.sub("B. Defined benefit plan - Gratuity (unfunded)")
    ns.para(cx.text("eb_grat_intro", "The Company provides for gratuity to eligible employees in accordance with the Payment of Gratuity Act, 1972. An employee "
                                     "who has completed five years of service is entitled to gratuity at 15 days' salary (last drawn basic) for each completed year "
                                     "of service, subject to a ceiling of Rs. 20 lakhs. The scheme is unfunded. The liability is determined by an independent actuary"
                                     + (f" ({G['actuary']})" if G.get("actuary") else "") + " using the Projected Unit Credit method, in accordance with AS 15 "
                                     "(Revised 2005) \"Employee Benefits\"."), last_col=cx.LASTN)
    ns.gap()
    FY = ["Y1", "Y2", "Y3"]
    fcol = {p: cx.NC[p] for p in FY}

    def fyhead(kind, label="Particulars"):
        r = ns.r
        w(ws, f"C{r}", label, fB, bd=B_TB)
        for p in FY:
            w(ws, f"{fcol[p]}{r}", cx.dtxt(p, kind), fB, bd=B_TB, al=A_CW)
        ws.row_dimensions[r].height = 30
        ns.r += 1

    def fyline(label, f, font=fN, fmt=ACC, bd=None, ref=None, lab_font=None):
        r = ns.r
        w(ws, f"C{r}", label, lab_font or font, bd=bd)
        for i, p in enumerate(FY):
            v = f(i, p, r)
            if v is not None:
                w(ws, f"{fcol[p]}{r}", v, font, fmt, bd=bd)
        if ref:
            cx.setref(ref, r)
        ns.r += 1
        return r
    ns.sub("a) Actuarial assumptions")
    fyhead("bs")
    for row in G.get("assumptions") or []:
        lab, *vals = row
        is_pct = all(isinstance(v, (int, float)) and abs(v) < 1 for v in vals if v not in (None, ""))
        fmt = PCT if is_pct and lab and "employees" not in str(lab).lower() else (ACC0 if "employees" in str(lab).lower() else None)
        fyline(lab or "", lambda i, p, r, vals=vals: vals[i] if i < len(vals) else None, fmt=fmt)
    ns.gap()
    D = G.get("dbo") or {}
    ns.sub("b) Change in the present value of defined benefit obligation")
    fyhead("pl")
    b0 = ns.r
    for key, lbl in (("open", "Present value of obligation as at the beginning of the year"), ("service", "Current service cost"), ("interest", "Interest cost"),
                     ("paid", "Benefits paid"), ("actuarial", "Actuarial (gain) / loss on obligations")):
        vals = D.get(key) or [None, None, None]
        if key == "actuarial":
            cl = D.get("closing") or [None, None, None]
            fyline(lbl, lambda i, p, r, vals=vals, cl=cl: (f"={fcol[p]}{r + 1}-SUM({fcol[p]}{b0}:{fcol[p]}{r - 1})" if _n(cl[i]) is not None
                                                          else rs_f(vals[i], S) or 0))
        else:
            fyline(lbl, lambda i, p, r, vals=vals: rs_f(vals[i], S) or 0)
    cl = D.get("closing") or [None, None, None]
    rc = fyline("Present value of obligation as at the end of the year", lambda i, p, r: rs_f(cl[i], S) if _n(cl[i]) is not None
                else f"=SUM({fcol[p]}{b0}:{fcol[p]}{r - 1})", font=fB, bd=B_TB)
    ns.gap()
    ns.sub("c) Actuarial (gain) / loss recognised")
    fyhead("pl")
    ra = fyline("Actuarial (gain) / loss on obligations", lambda i, p, r: f"={fcol[p]}{b0 + 4}")
    rna = fyline("Net actuarial (gain) / loss recognised in the Statement of Profit and Loss", lambda i, p, r: f"={fcol[p]}{ra}", font=fB, bd=B_TB)
    ns.gap()
    ns.sub("d) Expenses recognised in the Statement of Profit and Loss")
    fyhead("pl")
    e0 = ns.r
    fyline("Current service cost", lambda i, p, r: f"={fcol[p]}{b0 + 1}")
    fyline("Interest cost on benefit obligation", lambda i, p, r: f"={fcol[p]}{b0 + 2}")
    fyline("Net actuarial (gain) / loss recognised in the year", lambda i, p, r: f"={fcol[p]}{rna}")
    rex = fyline("Expense recognised as per actuarial valuation", lambda i, p, r: f"=SUM({fcol[p]}{e0}:{fcol[p]}{r - 1})", font=fB, bd=B_TB)
    dp = G.get("direct_paid") or [None, None, None]
    fyline("Add: Gratuity paid directly to employees and charged to the Statement of Profit and Loss", lambda i, p, r: rs_f(dp[i], S) or 0)
    nn = cx.noteno.get("EMP")
    rgx = fyline(f"Gratuity expense as per Note {nn} (Employee benefits expense)" if nn else "Gratuity expense",
                 lambda i, p, r: f"={fcol[p]}{r - 2}+{fcol[p]}{r - 1}", font=fB, bd=B_TB)
    if "pl_grat" not in getattr(cx, "missing", set()):
        for i, p in enumerate(FY):
            w(ws, f"{CL(6 + cx.N + i)}{rgx}", f"=ROUND({fcol[p]}{rgx}-'Notes_P&L'!{cx.NC[p]}{{{{pl_grat}}}},2)", fChk, ACC)
    ns.gap(2)
    ns.sub("e) Balance sheet reconciliation")
    fyhead("bs")
    f0 = ns.r
    fyline("Opening net liability", lambda i, p, r: f"={fcol[p]}{b0}")
    fyline("Expense recognised during the year (as per actuarial valuation)", lambda i, p, r: f"={fcol[p]}{rex}")
    fyline("Benefits paid during the year", lambda i, p, r: f"=-{fcol[p]}{b0 + 3}")
    rbs = fyline("Net liability recognised in the balance sheet", lambda i, p, r: f"=SUM({fcol[p]}{f0}:{fcol[p]}{r - 1})", font=fB, bd=B_TB)
    ns.gap()
    ns.sub("f) Current and non-current classification of the obligation")
    ns.colhead("bs", opening=True)
    CUR = G.get("current") or {}
    nl, ns_ = cx.noteno.get("STP"), cx.noteno.get("LTP")

    def cur(p, r):
        v = CUR.get(p)
        if _n(v) is not None:
            return rs_f(v, S)
        if p == "S":
            return f"={cx.NC['Y1']}{r}"
        return 0
    rcur = ns.line(f"Current (included in short-term provisions - Note {nl})", cur, ref="gr_cur")
    for p in cx.PK + ["OPEN"]:
        if _n(CUR.get(p)) is not None:
            _blue(ws, f"{cx.NC[p]}{rcur}")
    ns.line(f"Non-current (included in long-term provisions - Note {ns_})", lambda p, r: f"={cx.NC[p]}{r + 1}-{cx.NC[p]}{r - 1}")
    PG = "Provision for gratuity"
    rtp = ns.line("Total provision for gratuity", lambda p, r: "=" + cx.master_val(p, PG, qs(PG), "BS", "Cr"), font=fB, bd=B_TB)
    if not getattr(cx, "no_grat", False) and "ltp_grat" in cx.refs:
        for i, k in enumerate(cx.PK + ["OPEN"]):
            w(ws, f"{CL(6 + cx.N + i)}{rtp}", f"=ROUND({cx.NC[k]}{rtp}-Notes_Liabilities!{cx.NC[k]}{{{{ltp_grat}}}}-Notes_Liabilities!{cx.NC[k]}{{{{grat_cur}}}},2)",
              fChk, ACC)
    for i, p in enumerate(FY):           # closing obligation as per valuation = provision restated
        w(ws, f"{CL(6 + cx.N + i)}{rc}", f"=ROUND({fcol[p]}{rc}-{cx.NC[p]}{rtp},2)", fChk, ACC)
    ns.gap()
    ns.sub("g) Experience adjustments")
    fyhead("bs")
    fyline("Defined benefit obligation", lambda i, p, r: f"={fcol[p]}{rc}")
    fyline("Plan assets", lambda i, p, r: 0)
    fyline("Surplus / (deficit)", lambda i, p, r: f"={fcol[p]}{r - 1}-{fcol[p]}{r - 2}")
    ex = G.get("experience") or [None, None, None]
    fyline("Experience adjustments on plan liabilities - loss / (gain)", lambda i, p, r: rs_f(ex[i], S) or 0)
    ns.gap()
    sens = G.get("sensitivity")
    if sens:
        ns.sub(f"h) Sensitivity analysis of the defined benefit obligation as at {long_date(cx.DATE['Y1'])}")
        r = ns.r
        w(ws, f"C{r}", "Particulars", fB, bd=B_TB)
        w(ws, f"{cx.NC['Y1']}{r}", "Increase by 100 basis points", fB, bd=B_TB, al=A_CW)
        w(ws, f"{cx.NC['Y2']}{r}", "Decrease by 100 basis points", fB, bd=B_TB, al=A_CW)
        ns.r += 1
        for lbl, v in zip(["Discount rate", "Salary escalation rate", "Attrition rate"], sens):
            r = ns.r
            w(ws, f"C{r}", lbl)
            w(ws, f"{cx.NC['Y1']}{r}", rs_f(v[0], S), fN, ACC)
            w(ws, f"{cx.NC['Y2']}{r}", rs_f(v[1], S), fN, ACC)
            ns.r += 1
        ns.para(cx.text("eb_sens", "The sensitivity analysis has been determined by the actuary by changing one assumption at a time, keeping the other "
                                   "assumptions constant."), fI, last_col=cx.LASTN)
        ns.gap()
    ns.sub("i) Plan assets" if sens else "h) Plan assets")
    ns.para(cx.text("eb_plan", "The gratuity scheme is unfunded and the Company does not hold any plan assets; accordingly, the disclosures relating to the "
                               "fair value of plan assets and the return on plan assets are not applicable."), last_col=cx.LASTN)
    gn = [f"1. Figures are as per the actuarial valuation reports for the respective years."]
    if cx.stub and _n(CUR.get("S")) is None:
        gn.append(f"2. Actuarial valuation as at {long_date(cx.stub)} is not available; the provision as at {long_date(cx.stub)} is carried at the "
                  f"{long_date(cx.DATE['Y1'])} valuation (including its current / non-current classification).")
    gn = cx.text_list("eb_notes", gn, "Employee benefits - notes (one per line)")
    if gn:
        ns.gap()
        ns.sub("Notes:")
        for t in gn:
            ns.para(t, last_col=cx.LASTN)
    cx.finish_print(ws, ns.r)


# ====================================================================== Tax Computation
def tax_comp(cx):
    I = cx.I
    T = I.get("tax") or {}
    ns = NoteSheet(cx, "Tax Computation", False, prev_ref="'Gratuity and contingencies'!$B${{note:EB}}")
    ws = ns.ws
    S = cx.SCALE
    PL = cx.PL
    ns.note("TAXC", "Statement of tax shelter")
    ns.colhead("pl")
    ra = ns.line("Restated profit before tax (A)", lambda p, r: f"=PL_2!{cx.BC[p]}{PL['pbt']}", font=fB)
    rr = ns.line("Tax rate - normal rate under Section 115BAA (including surcharge and cess)", lambda p, r: f"={cx.TAX}", fmt=PCT3)
    ns.sub("Income considered separately")
    ints = cx.OI_INT_ROWS
    f0 = ns.r
    if ints:
        ns.line('Interest income (taxable under "Income from other sources")', lambda p, r: "=-(" + "+".join(f"'Notes_P&L'!{cx.NC[p]}{x}" for x in ints) + ")")
    else:
        ns.line('Interest income (taxable under "Income from other sources")', lambda p, r: 0)
    rb = ns.total(f0, ns.r - 1, "Total income considered separately (B)", bd=B_TB)
    ns.sub("Permanent differences")
    f0 = ns.r
    if "pl_msme_int" in cx.refs and "pl_msme_int" not in getattr(cx, "missing", set()):
        ns.line("Interest on delayed payments to MSME (Section 23 of the MSMED Act, 2006)", lambda p, r: f"='Notes_P&L'!{cx.NC[p]}{{{{pl_msme_int}}}}")
    for lbl, v in (T.get("permanent") or {}).items():
        ns.line(lbl, lambda p, r, v=v: rs_f(v.get(p), S) or 0)
    if ns.r == f0:
        ns.line("Nil", lambda p, r: 0)
    rc = ns.total(f0, ns.r - 1, "Total permanent differences (C)", bd=B_TB)
    ns.sub("Timing differences")
    r1 = ns.line("Depreciation as per the Companies Act, 2013 (a)", lambda p, r: f"=PL_2!{cx.BC[p]}{PL['dep']}")
    itd = T.get("it_dep") or {}
    r2 = ns.line("Depreciation as per the Income-tax Act, 1961 (b)", lambda p, r: rs_f(-abs(_n(itd.get(p))), S) if _n(itd.get(p)) else 0)
    rd = ns.line("Difference between book depreciation and tax depreciation (a + b)", lambda p, r: f"={cx.NC[p]}{r1}+{cx.NC[p]}{r2}", font=fB, bd=B_TB)
    ns.sub("Other differences", fI)
    o0 = ns.r
    loss_led = (cx.I.get("cfs") or {}).get("ppe_loss_ledgers") or []
    fa_c = getattr(cx, "FBR", None)
    for lbl, v in (T.get("rows") or {}).items():
        if lbl.startswith("Disallowance u/s 40A(7) / 43B - provision for gratuity") and "pl_grat" in cx.refs:
            continue
        if "loss on write off" in lbl.lower() and (loss_led or fa_c):
            # restated loss on deletion of assets = loss in the books - (c) of the FA restatement
            def lf(p, r):
                parts = [cx.tbl(p, l) for l in loss_led]
                if fa_c:
                    parts.append(f"-'FA Restatement'!{cx.FBC[p]}{cx.FBR['c']}")
                return "=" + "+".join(parts).replace("+-", "-")
            ns.line(lbl, lf)
            continue
        ns.line(lbl, lambda p, r, v=v: rs_f(v.get(p), S) or 0)
    if T.get("gratuity_disallow", True) and "pl_grat" not in getattr(cx, "missing", set()):
        ns.line("Disallowance u/s 40A(7) / 43B - provision for gratuity", lambda p, r: f"='Notes_P&L'!{cx.NC[p]}{{{{pl_grat}}}}")
    o1 = ns.r - 1
    rtd = ns.line("Total timing differences (D)", lambda p, r: f"={cx.NC[p]}{rd}" + (f"+SUM({cx.NC[p]}{o0}:{cx.NC[p]}{o1})" if o1 >= o0 else ""), font=fB, bd=B_TB)
    rne = ns.line("Net adjustments (E = B + C + D)", lambda p, r: f"={cx.NC[p]}{rb}+{cx.NC[p]}{rc}+{cx.NC[p]}{rtd}", font=fB, bd=B_TB)
    rf = ns.line("Income from other sources (F)", lambda p, r: f"=-{cx.NC[p]}{rb}")
    rg = ns.line("Taxable income (G = A + E + F)", lambda p, r: f"={cx.NC[p]}{ra}+{cx.NC[p]}{rne}+{cx.NC[p]}{rf}", font=fB, bd=B_TB, ref="tax_taxable")
    rt = ns.line("Tax payable at normal rate (G x tax rate)", lambda p, r: f"=MAX(0,{cx.NC[p]}{rg})*{cx.NC[p]}{rr}", font=fB, bd=B_TB, ref="tax_payable")
    rp = ns.line("Current tax as per Restated Summary Statement of Profit and Loss", lambda p, r: f"=PL_2!{cx.BC[p]}{PL['ct']}")
    ns.line("Difference", lambda p, r: f"=ROUND({cx.NC[p]}{rt}-{cx.NC[p]}{rp},2)")
    ns.gap()
    ns.para(cx.text("tax_notes", "Notes: (1) The Company has opted for the concessional tax regime under Section 115BAA of the Income-tax Act, 1961; "
                                 "accordingly, Minimum Alternate Tax under Section 115JB is not applicable. (2) Depreciation as per the Income-tax Act and the "
                                 "disallowances / allowances are as per the computation of income filed with the returns of income."))
    ns.gap()
    ns.sub("Details as per return of income filed / computation of income")
    ns.colhead("pl")
    ITR = T.get("itr") or {}
    labels = ["Assessment year", "Acknowledgement number", "Date of filing (u/s 139(1))", "Total income as per return / computation (Rs.)",
              "Tax payable as per return / computation (Rs.)", "Taxes paid - advance tax / TDS / TCS (Rs.)", "Refund claimed as per return (Rs.)"]
    ri = {}
    for j, lbl in enumerate(labels):
        def f(p, r, j=j):
            v = (ITR.get(p) or [None] * 7)
            v = v[j] if j < len(v) else None
            if j == 0 and v in (None, ""):
                return "Interim period" if p == "S" else f"{cx.DATE[p].year}-{str(cx.DATE[p].year + 1)[2:]}"
            return v
        ri[j] = ns.line(lbl, f, fmt=ACC0 if j >= 3 else None)
    ns.line("Current tax as per return / computation (Rs. in lakhs)", lambda p, r: f"=IF(ISNUMBER({cx.NC[p]}{ri[4]}),{cx.NC[p]}{ri[4]}/{S},PL_2!{cx.BC[p]}{PL['ct']})",
            font=fB, bd=B_TB)
    ns.line("Difference between total income as per return / computation and taxable income computed above",
            lambda p, r: f"=IF(ISNUMBER({cx.NC[p]}{ri[3]}),ROUND({cx.NC[p]}{ri[3]}/{S}-{cx.NC[p]}{rg},2),\"\")")
    cx.finish_print(ws, ns.r)


# ====================================================================== Accounting ratios
def acc_ratios(cx):
    ns = NoteSheet(cx, "Accounting ratios", False, prev_ref="'Tax Computation'!B8")
    ws = ns.ws
    S = cx.SCALE
    PL, BS = cx.PL, cx.BS
    pl = lambda p, k: f"PL_2!{cx.BC[p]}{PL[k]}"
    bs = lambda p, k: f"BS_1!{cx.BC[p]}{BS[k]}"
    NE = q("Notes_Equity & Reserves")
    ns.note("AR", "Summary of accounting ratios")
    ns.colhead("pl")
    ra = ns.line("Total revenue from operations (A)", lambda p, r: f"={pl(p, 'rev')}", font=fB)
    ns.gap()
    rb = ns.line("Net profit as restated (B)", lambda p, r: f"={pl(p, 'pat')}", font=fB)
    ns.line("a) Tax expense", lambda p, r: f"={pl(p, 'tt')}")
    ns.line("b) Depreciation and amortisation expense", lambda p, r: f"={pl(p, 'dep')}")
    ns.line("c) Interest expense (finance costs)", lambda p, r: f"={pl(p, 'fin')}")
    re_ = ns.line("Earnings before interest, tax, depreciation and amortisation (EBITDA) (C = B + a + b + c)", lambda p, r: f"=SUM({cx.NC[p]}{rb}:{cx.NC[p]}{r - 1})",
                  font=fB)
    ns.line("EBITDA margin (%) (C / A)", lambda p, r: f"=IFERROR({cx.NC[p]}{re_}/{cx.NC[p]}{ra},0)", fmt=PCT)
    ns.gap()
    rd = ns.line("Net worth as restated (D)", lambda p, r: f"={bs(p, 'shf')}", font=fB)
    ns.line("Return on net worth (%) (B / D)", lambda p, r: f"=IFERROR({cx.NC[p]}{rb}/{cx.NC[p]}{rd},0)", fmt=PCT)
    ns.gap()
    rn = ns.line("Number of equity shares outstanding at the end of the year / period (E)", lambda p, r: f"={NE}!{cx.NC[p]}{{{{sc_end}}}}", fmt=ACC0)
    rw = ns.line("Weighted average number of equity shares (F)", lambda p, r: f"=({NE}!{cx.NC[p]}{{{{sc_end}}}}+{NE}!{cx.NC[p]}{{{{sc_begin}}}})/2", fmt=ACC0)
    ra2 = ns.line("Weighted average number of equity shares adjusted for bonus / split (G = F x bonus factor)", lambda p, r: f"={cx.NC[p]}{rw}*{MD}!$B$8",
                  fmt=ACC0, ref="ar_wans_adj")
    ns.line("Basic and diluted earnings per share (Rs.) (B / G)", lambda p, r: f"=IFERROR({cx.NC[p]}{rb}*{S}/{cx.NC[p]}{ra2},0)", font=fB, ref="ar_eps")
    ns.line("Net asset value per equity share (Rs.) (D / E adjusted for bonus / split)",
            lambda p, r: f"=IFERROR({cx.NC[p]}{rd}*{S}/({cx.NC[p]}{rn}*{MD}!$B$8),0)", font=fB)
    ns.gap()
    rca = ns.line("Current assets (H)", lambda p, r: f"={bs(p, 'ca')}")
    rcl = ns.line("Current liabilities (I)", lambda p, r: f"={bs(p, 'cl')}")
    ns.line("Current ratio (in times) (H / I)", lambda p, r: f"=IFERROR({cx.NC[p]}{rca}/{cx.NC[p]}{rcl},0)", font=fB)
    ns.gap()
    ns.sub("Notes:")
    dflt = ["1. Ratios have been calculated as below:",
            "Basic and diluted earnings per share (Rs.) = Restated profit after tax attributable to equity shareholders / Weighted average number of equity "
            "shares outstanding during the year / period (adjusted for bonus / split, if any)",
            "Return on net worth (%) = Restated profit after tax / Restated net worth at the end of the year / period",
            "Net asset value per equity share (Rs.) = Restated net worth at the end of the year / period / Number of equity shares outstanding at the end of "
            "the year / period",
            "Current ratio = Current assets / Current liabilities",
            "EBITDA = Restated profit after tax + tax expense + depreciation and amortisation expense + finance costs",
            "2. \"Net worth\" means the aggregate value of the paid-up share capital and all reserves created out of the profits and securities premium "
            "account, after deducting the aggregate value of the accumulated losses, deferred expenditure and miscellaneous expenditure not written off, as "
            "per the Restated Summary Statement of Assets and Liabilities.",
            "3. There has been no bonus issue, split or consolidation of shares during the years / period presented; accordingly, pre-bonus and post-bonus "
            "figures are the same."]
    if cx.stub:
        dflt.append(f"4. The ratios for the period ended {long_date(cx.stub)} are not annualised.")
    dflt.append(f"{len([x for x in dflt if x[:2] in ('1.', '2.', '3.', '4.')]) + 1}. The figures disclosed above are based on the Restated Summary "
                "Statements of the Company.")
    for t in cx.text_list("ar_notes", dflt, "Accounting ratios - notes (one per line)"):
        ns.para(t)
    cx.finish_print(ws, ns.r)


# ====================================================================== Analytical Ratios
RATIOS = [  # key, name, numerator desc, denominator desc, num row, den row, flow-based (not comparable for a stub)
    ("cr", "Current ratio (in times)", "Current assets", "Current liabilities", "ca", "cl", False),
    ("de", "Debt-equity ratio (in times)", "Total debt", "Shareholders' equity", "debt", "eq", False),
    ("dscr", "Debt service coverage ratio (in times)", "Earnings available for debt service", "Debt service", "eads", "ds", False),
    ("roe", "Return on equity ratio (%)", "Net profit after tax", "Average shareholders' equity", "pat", "aeq", True),
    ("itr", "Inventory turnover ratio (in times)", "Cost of goods sold", "Average inventory", "cogs", "ainv", True),
    ("trt", "Trade receivables turnover ratio (in times)", "Revenue from operations (net credit sales)", "Average trade receivables", "rev", "atr", True),
    ("tpt", "Trade payables turnover ratio (in times)", "Net credit purchases", "Average trade payables", "pur", "atp", True),
    ("nct", "Net capital turnover ratio (in times)", "Revenue from operations", "Average working capital", "rev", "awc", True),
    ("npr", "Net profit ratio (%)", "Net profit after tax", "Revenue from operations", "pat", "rev", False),
    ("roce", "Return on capital employed (%)", "Earnings before interest and tax", "Capital employed", "ebit", "ce", True),
    ("roi", "Return on investment (%)", "Interest income on fixed deposits", "Average fixed deposits", "fdi", "afd", True),
]


def analytical(cx):
    ws = cx.sheet("Analytical Ratios", {"A": 1.0, "B": 5.0, "C": 34.0, "D": 24.0, "E": 24.0, "F": 13.0, "G": 13.0, "H": 11.0, "I": 48.0, "J": 12.0,
                                        "O": 55.0, **{CL(16 + i): 13.0 for i in range(cx.N + 1)}}, print_cols="I")
    head(cx, ws, "ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS", "I")
    w(ws, "B8", "='Accounting ratios'!B8+1", fB, al=A_L)
    w(ws, "C8", "Analytical ratios", fB)
    cx.noteno_counter += 1
    cx.noteno["ANR"] = cx.noteno_counter
    keys = cx.PK + ["OPEN"]
    WC = {k: CL(16 + i) for i, k in enumerate(keys)}
    PL, BS = cx.PL, cx.BS
    pl = lambda p, k: f"PL_2!{cx.BC[p]}{PL[k]}"
    bs = lambda p, k: f"BS_1!{cx.BC[p]}{BS[k]}" if k in BS else "0"
    w(ws, "O10", "Working for analytical ratios (Rs. in lakhs)", fB)
    w(ws, "O11", "Particulars", fB, bd=B_TB)
    for k in keys:
        w(ws, f"{WC[k]}11", cx.dtxt(k, "bs"), fB, bd=B_TB, al=A_CW)
    rows = {}
    r = 12
    nxt = {k: WC[cx.PREV[k]] for k in cx.PK}

    def wl(key, label, f, opening=True):
        nonlocal r
        w(ws, f"O{r}", label)
        for k in keys:
            if k == "OPEN" and not opening:
                w(ws, f"{WC[k]}{r}", "N.A.", fN, al=A_C)
            else:
                w(ws, f"{WC[k]}{r}", f(k), fN, ACC)
        rows[key] = r
        r += 1
    avg = lambda key: (lambda k: f"=({WC[k]}{rows[key]}+{nxt[k]}{rows[key]})/2")
    wl("ca", "Current assets", lambda k: f"={bs(k, 'ca')}")
    wl("cl", "Current liabilities", lambda k: f"={bs(k, 'cl')}")
    wl("debt", "Total debt (long-term + short-term borrowings)", lambda k: f"={bs(k, 'ltb')}+{bs(k, 'stb')}")
    wl("eq", "Shareholders' equity", lambda k: f"={bs(k, 'shf')}")
    wl("pat", "Net profit after tax", lambda k: f"={pl(k, 'pat')}", False)
    wl("dep", "Depreciation and amortisation", lambda k: f"={pl(k, 'dep')}", False)
    wl("fin", "Finance costs", lambda k: f"={pl(k, 'fin')}", False)
    PR = (cx.I.get("analytical") or {}).get("principal_repaid") or {}
    wl("prin", "Principal repayment of long-term borrowings", lambda k: rs_f(PR.get(k), cx.SCALE) or 0, False)
    wl("eads", "Earnings available for debt service (PAT + depreciation + finance costs)",
       lambda k: f"={WC[k]}{rows['pat']}+{WC[k]}{rows['dep']}+{WC[k]}{rows['fin']}", False)
    wl("ds", "Debt service (finance costs + principal repayments)", lambda k: f"={WC[k]}{rows['fin']}+{WC[k]}{rows['prin']}", False)
    wl("aeq", "Average shareholders' equity", avg("eq"), False)
    wl("cogs", "Cost of goods sold (purchases + changes in inventories)", lambda k: f"={pl(k, 'pur')}+{pl(k, 'chinv')}", False)
    wl("inv", "Inventories", lambda k: f"={bs(k, 'inv')}")
    wl("ainv", "Average inventories", avg("inv"), False)
    wl("rev", "Revenue from operations", lambda k: f"={pl(k, 'rev')}", False)
    wl("tr", "Trade receivables", lambda k: f"={bs(k, 'tr')}")
    wl("atr", "Average trade receivables", avg("tr"), False)
    wl("pur", "Purchases of stock-in-trade", lambda k: f"={pl(k, 'pur')}", False)
    wl("tp", "Trade payables", lambda k: f"={bs(k, 'tpm')}+{bs(k, 'tpo')}")
    wl("atp", "Average trade payables", avg("tp"), False)
    wl("wc", "Working capital (current assets - current liabilities)", lambda k: f"={WC[k]}{rows['ca']}-{WC[k]}{rows['cl']}")
    wl("awc", "Average working capital", avg("wc"), False)
    wl("ebit", "Earnings before interest and tax (PBT + finance costs)", lambda k: f"={pl(k, 'pbt')}+{pl(k, 'fin')}", False)
    wl("ce", "Capital employed (net worth + total debt)", lambda k: f"={WC[k]}{rows['eq']}+{WC[k]}{rows['debt']}")
    fdi = cx.OI_FD_ROW
    wl("fdi", "Interest income on fixed deposits", lambda k: f"='Notes_P&L'!{cx.NC[k]}{fdi}" if fdi else 0, False)
    wl("fd", "Fixed deposits with banks", lambda k: f"=Notes_Assets!{cx.NC[k]}{{{{fd_line}}}}" if "fd_line" in cx.refs else 0)
    wl("afd", "Average fixed deposits", avg("fd"), False)
    # ratio blocks
    r = 10
    blocks = [(p, cx.PREV[p]) for p in cx.PK]
    texts = cx.I.get("texts") or {}
    for bi, (p, q_) in enumerate(blocks):
        w(ws, f"C{r}", ("For the period ended " if p == "S" else "For the year ended ") + long_date(cx.DATE[p]), fB)
        r += 1
        for col, h in zip("CDE", ["Ratio", "Numerator", "Denominator"]):
            w(ws, f"{col}{r}", h, fB, bd=B_ALL, al=A_CW)
        w(ws, f"F{r}", f'=TEXT({cx.MDC[p]},"mmmm dd, yyyy")', fB, bd=B_ALL, al=A_CW)
        w(ws, f"G{r}", f'=TEXT({cx.MDC[q_]},"mmmm dd, yyyy")', fB, bd=B_ALL, al=A_CW)
        w(ws, f"H{r}", "% Variance", fB, bd=B_ALL, al=A_CW)
        w(ws, f"I{r}", "Reason for variance (where variance exceeds 25%)", fB, bd=B_ALL, al=A_CW)
        r += 1
        for i, (key, name, nd, dd, nr, dr, flow) in enumerate(RATIOS):
            bd = B_BLR if i == len(RATIOS) - 1 else B_LR
            pc_, qc = WC[p], WC[q_]
            pct = key in ("roe", "npr", "roce", "roi")
            w(ws, f"C{r}", name, fN, bd=bd, al=A_W)
            w(ws, f"D{r}", nd, fN, bd=bd, al=A_W)
            w(ws, f"E{r}", dd, fN, bd=bd, al=A_W)
            w(ws, f"F{r}", f'=IFERROR({pc_}{rows[nr]}/{pc_}{rows[dr]},"N.A.")', fN, PCT if pct else RATIO, bd=bd)
            w(ws, f"G{r}", f'=IFERROR({qc}{rows[nr]}/{qc}{rows[dr]},"N.A.")', fN, PCT if pct else RATIO, bd=bd)
            w(ws, f"H{r}", f'=IFERROR((F{r}-G{r})/ABS(G{r}),"N.A.")', fN, PCT, bd=bd)
            over = texts.get(f"ar_reason:{p}:{key}")
            if over:
                w(ws, f"I{r}", over, fN, bd=bd, al=A_W)
            else:
                ch = lambda row: f'IFERROR(({pc_}{row}-{qc}{row})/ABS({qc}{row}),0)'
                inc = lambda row: f'IF({ch(row)}>=0,"an increase","a decrease")&" of "&TEXT(ABS({ch(row)}),"0%")'
                prior = long_date(cx.DATE[q_])
                normal = (f'IF(H{r}>0,"Increase","Decrease")&" on account of "&{inc(rows[nr])}&" in {nd.lower()} as compared to "&{inc(rows[dr])}&'
                          f'" in {dd.lower()}."')
                if p == "S" and flow:
                    months = round(((cx.stub - cx.DATE["Y1"]).days) / 30.4)
                    normal = (f'"Not comparable - the ratio for the period ended {long_date(cx.stub)} is computed on the basis of {months} months\' figures '
                              f'(not annualised); {nd.lower()} for the period against {dd.lower()} ("&{inc(rows[dr])}&")."')
                fml = (f'=IF(NOT(ISNUMBER(G{r})),"Not applicable - comparative ratio for the {"period" if q_ == "S" else "year"} ended {prior} is not available",'
                       f'IF(NOT(ISNUMBER(H{r})),"-",IF(ABS(H{r})<=0.25,"-",{normal})))')
                w(ws, f"I{r}", fml, fN, bd=bd, al=A_W)
            ws.row_dimensions[r].height = 40
            r += 1
        r += 2
    t = cx.text("anr_notes", "Notes: (1) Average balances are the simple average of the opening and closing balances. (2) Ratios that require items of the "
                             "Statement of Profit and Loss or balances of the year preceding the earliest year presented are not presented (\"N.A.\"). "
                             + (f"(3) Ratios for the period ended {long_date(cx.stub)} are not annualised and hence are not comparable with the ratios for "
                                f"the year ended {long_date(cx.DATE['Y1'])}." if cx.stub else ""))
    merge(ws, f"C{r}:I{r}", t, fN, A_J, est_height(t, 160))
    cx.finish_print(ws, r + 1)


# ====================================================================== Borrowing Terms
def borrowing_terms(cx):
    I = cx.I
    L = I.get("borrowings") or []
    keys = cx.PK
    BTC = {k: CL(10 + i) for i, k in enumerate(keys)}
    cx.BTC = BTC
    last = BTC[keys[-1]]
    ws = cx.sheet("Borrowing Terms", {"A": 2.0, "B": 5.0, "C": 26.0, "D": 34.0, "E": 12.0, "F": 11.0, "G": 26.0, "H": 28.0, "I": 34.0,
                                     **{BTC[k]: 13.0 for k in keys}}, print_cols=last)
    head(cx, ws, "ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS", last)
    w(ws, "B8", "='Analytical Ratios'!B8+1", fB, al=A_L)
    w(ws, "C8", "Statement of principal terms of borrowings (Annexure for terms of borrowings)", fB)
    cx.noteno_counter += 1
    cx.noteno["BT"] = cx.noteno_counter
    for col, h in zip("CDEFGHI", ["Lender", "Nature of facility", "Date of sanction", "Sanctioned limit", "Rate of interest", "Tenor / repayment",
                                   "Security / guarantee"]):
        w(ws, f"{col}10", h, fB, bd=B_TB, al=A_CW)
    for k in keys:
        w(ws, f"{BTC[k]}10", cx.dtxt(k, "Outstanding as at\n"), fB, bd=B_TB, al=A_CW)
    ws.row_dimensions[10].height = 30
    r = 11
    cx.BT_GUAR = []
    TBD = "[Sanction letter to be furnished]"
    for b in L:
        w(ws, f"C{r}", b.get("lender", ""), fN, al=A_W)
        w(ws, f"D{r}", b.get("nature", ""), fN, al=A_W)
        d = b.get("date")
        if d and not str(d).startswith("["):
            w(ws, f"E{r}", ddate(d), fBlue, DMY, al=A_W)
        else:
            w(ws, f"E{r}", d or TBD, fN, al=A_W)
        if _n(b.get("limit")) is not None:
            w(ws, f"F{r}", _n(b["limit"]), fBlue, ACC, al=A_W)
        for col, k in zip("GHI", ["roi", "tenor", "security"]):
            w(ws, f"{col}{r}", b.get(k) or TBD, fN, al=A_W)
        for k in keys:
            led = b.get("ledgers") or []
            w(ws, f"{BTC[k]}{r}", "=" + ("+".join(cx.tbl(k, l) for l in led) if led else "0"), fN, ACC)
        if b.get("guarantee"):
            cx.BT_GUAR.append(r)
        ws.row_dimensions[r].height = max(40, est_height(max((str(b.get(x) or "") for x in ("nature", "security", "roi", "tenor")), key=len), 30))
        r += 1
    if r == 11:                      # no facility entered: one empty row keeps the total off its own cell
        w(ws, f"C{r}", "No borrowings / facilities entered", fGreyI)
        for k in keys:
            w(ws, f"{BTC[k]}{r}", 0, fN, ACC)
        r += 1
    f0, l0 = 11, r - 1
    w(ws, f"C{r}", "Total", fB)
    for k in keys:
        w(ws, f"{BTC[k]}{r}", f"=SUM({BTC[k]}{f0}:{BTC[k]}{l0})", fB, ACC, bd=B_TB)
    rt = r
    r += 1
    nl = cx.noteno.get("STB"), cx.noteno.get("LTB")
    w(ws, f"C{r}", f"Check: agrees with short-term borrowings (Note {nl[0]}) and secured term loans (Note {nl[1]})", fGrey)
    TL = "Term loans from banks and NBFCs (secured)"
    for k in keys:
        sec = f"-Notes_Liabilities!{cx.NC[k]}{{{{ltb_sec_tl}}}}" if "ltb_sec_tl" in cx.refs else ""
        w(ws, f"{BTC[k]}{r}", f"=ROUND({BTC[k]}{rt}-Notes_Liabilities!{cx.NC[k]}{{{{stb_total}}}}{sec},2)", fChk, ACC)
    r += 2
    t = cx.text("bt_notes", "Notes: (i) Sanctioned limits are in Rs. lakhs. (ii) Terms are as per the sanction letters; terms of facilities for which "
                            "sanction letters have not been furnished are to be updated on receipt. (iii) Loans from directors are interest free and repayable "
                            f"on demand (Note {cx.noteno.get('LTB')}).")
    merge(ws, f"C{r}:{last}{r}", t, fI, A_J, est_height(t, 250))
    cx.finish_print(ws, r)


# ====================================================================== Capitalisation
def capitalisation(cx):
    ws = cx.sheet("Capitalisation", {"A": 1.0, "B": 5.0, "C": 60.0, "D": 20.0, "E": 20.0}, print_cols="E")
    head(cx, ws, "ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS", "E")
    w(ws, "B8", "='Borrowing Terms'!B8+1", fB, al=A_L)
    w(ws, "C8", "Capitalisation statement", fB)
    cx.noteno_counter += 1
    cx.noteno["CAP"] = cx.noteno_counter
    p = cx.PK[0]
    c = cx.BC[p]
    BS = cx.BS
    w(ws, "C10", "Particulars", fB, bd=B_TB)
    w(ws, "D10", cx.dtxt(p, "Pre-issue as at\n"), fB, bd=B_TB, al=A_CW)
    w(ws, "E10", "Post-issue *", fB, bd=B_TB, al=A_CW)
    ws.row_dimensions[10].height = 30
    rows = [
        (11, "Borrowings", None, fB),
        (12, "Short-term borrowings (excluding current maturities of long-term borrowings)", f"=BS_1!{c}{BS['stb']}-Notes_Liabilities!{cx.NC[p]}{{{{cur_mat}}}}", fN),
        (13, "Long-term borrowings (including current maturities)", f"=BS_1!{c}{BS['ltb']}+Notes_Liabilities!{cx.NC[p]}{{{{cur_mat}}}}", fN),
        (14, "Total borrowings (A)", "=D12+D13", fB),
        (16, "Shareholders' funds", None, fB),
        (17, "Share capital", f"=BS_1!{c}{BS['sc']}", fN),
        (18, "Reserves and surplus", f"=BS_1!{c}{BS['rs']}", fN),
        (19, "Total shareholders' funds (B)", "=D17+D18", fB),
        (21, "Long-term borrowings / shareholders' funds", "=IFERROR(D13/D19,0)", fN),
        (22, "Total borrowings / shareholders' funds (A / B)", "=IFERROR(D14/D19,0)", fN),
    ]
    for r, lbl, f, fo in rows:
        w(ws, f"C{r}", lbl, fo, bd=B_T if r in (14, 19) else None)
        if f:
            w(ws, f"D{r}", f, fo, RATIO if r in (21, 22) else ACC, bd=B_T if r in (14, 19) else None)
            w(ws, f"E{r}", "[ ]", fN, al=A_C)
    w(ws, "C24", cx.text("cap_note", "* Post-issue figures will be updated on finalisation of the issue price and number of shares."), fI)
    cx.finish_print(ws, 24)


# ====================================================================== Additional Disclosures
ADD_DEFAULT = [
    ("(i)", "(i) The title deeds of immovable properties are held in the name of the Company."),
    ("(ii)", "(ii) The Company has not revalued its property, plant and equipment (including right-of-use assets) or intangible assets during the years / "
             "period presented."),
    ("(iii)", "TABLE_LOANS"),
    ("(iv)", "(iv) No proceedings have been initiated or are pending against the Company for holding any benami property under the Benami Transactions "
             "(Prohibition) Act, 1988 and the rules made thereunder."),
    ("(v)", "(v) The quarterly / monthly returns or statements of current assets filed by the Company with the banks, where working capital limits in excess "
            "of Rs. 5 crores have been sanctioned on the basis of security of current assets, are in agreement with the books of account."),
    ("(vi)", "(vi) The Company has not been declared a wilful defaulter by any bank or financial institution or other lender (or consortium thereof) in "
             "accordance with the guidelines on wilful defaulters issued by the Reserve Bank of India."),
    ("(vii)", "(vii) The Company does not have any transactions with companies struck off under Section 248 of the Companies Act, 2013 or Section 560 of the "
              "Companies Act, 1956."),
    ("(viii)", "(viii) There are no charges or satisfaction of charges yet to be registered with the Registrar of Companies beyond the statutory period."),
    ("(ix)", "(ix) The Company does not have any subsidiary, associate or joint venture; accordingly, compliance with the number of layers prescribed under "
             "clause (87) of Section 2 of the Companies Act, 2013 read with the Companies (Restriction on number of Layers) Rules, 2017 is not applicable."),
    ("(x)", "(x) No scheme of arrangement has been approved by the competent authority in terms of Sections 230 to 237 of the Companies Act, 2013."),
    ("(xi)", "(xi) Utilisation of borrowed funds and share premium:"),
    ("(xi)A", "A. The Company has not advanced or loaned or invested funds (either borrowed funds or share premium or any other sources or kind of funds) to "
              "or in any other person(s) or entity(ies), including foreign entities (Intermediaries), with the understanding (whether recorded in writing or "
              "otherwise) that the Intermediary shall (a) directly or indirectly lend or invest in other persons or entities identified in any manner whatsoever "
              "by or on behalf of the Company (Ultimate Beneficiaries) or (b) provide any guarantee, security or the like to or on behalf of the Ultimate "
              "Beneficiaries."),
    ("(xi)B", "B. The Company has not received any fund from any person(s) or entity(ies), including foreign entities (Funding Party), with the understanding "
              "(whether recorded in writing or otherwise) that the Company shall (a) directly or indirectly lend or invest in other persons or entities "
              "identified in any manner whatsoever by or on behalf of the Funding Party (Ultimate Beneficiaries) or (b) provide any guarantee, security or the "
              "like on behalf of the Ultimate Beneficiaries."),
    ("(xii)", "(xii) The borrowings from banks and financial institutions have been used for the specific purposes for which they were taken."),
    ("(xiii)", "TABLE_AUDITOR"),
    ("(xiv)", "(xiv) There are no transactions that were not recorded in the books of account and that have been surrendered or disclosed as income during "
              "the years in the tax assessments under the Income-tax Act, 1961; there is no previously unrecorded income and related assets to be recorded "
              "in the books of account."),
    ("(xv)", "TABLE_FOREIGN"),
    ("(xvi)", "(xvi) The Company has not traded or invested in crypto currency or virtual currency during the years / period presented."),
    ("(xvii)", "(xvii) Corporate Social Responsibility under Section 135 of the Companies Act, 2013 is not applicable to the Company for the years presented."),
    ("(xviii)", "(xviii) Capital work-in-progress and intangible assets under development - ageing schedule: refer Note {NOTE:PPE} (Notes_PPE). There are "
                "no projects whose completion is overdue or which have exceeded their cost compared to the original plan."),
    ("(xix)", "(xix) Segment reporting: The Company operates in a single reportable business segment as per AS 17 \"Segment Reporting\" and in a single "
              "geographical segment (India)."),
    ("(xx)", "(xx) Audit trail: The Company has used accounting software for maintaining its books of account which has a feature of recording audit trail "
             "(edit log) facility, and the same has operated throughout the year for all relevant transactions recorded in the software, as reported by the "
             "statutory auditors; the audit trail has been preserved by the Company as per the statutory requirements for record retention (Rule 3(1) of the "
             "Companies (Accounts) Rules, 2014)."),
    ("(xxi)", "(xxi) Dividend: The Company has not declared or paid any dividend during the years / period presented."),
    ("(xxii)", "(xxii) Events after the reporting period: There are no material events after the latest balance sheet date which require adjustment to or "
               "disclosure in the Restated Summary Statements."),
    ("(xxiii)", "(xxiii) Labour Codes: The Code on Wages, 2019 and the Code on Social Security, 2020 have been notified by the Government of India. The "
                "Company will assess the impact, if any, on its employee benefit obligations, and account for the same in the period in which the related "
                "rules are notified and become effective (refer Note {NOTE:EB})."),
    ("(xxiv)", "(xxiv) Disclosures required by SEBI ICDR Regulations: details of related party transactions (Note {NOTE:RPT}), contingent liabilities and "
               "commitments (Note {NOTE:CL}), employee benefits (Note {NOTE:EB}), statement of tax shelter (Note {NOTE:TAXC}), accounting ratios (Note "
               "{NOTE:AR}), analytical ratios (Note {NOTE:ANR}), principal terms of borrowings (Note {NOTE:BT}) and capitalisation statement (Note "
               "{NOTE:CAP}) are disclosed separately in this Annexure."),
]


def additional(cx):
    I = cx.I
    ws = cx.sheet("Additional Disclosures", {"A": 1.5, "B": 5.0, "C": 55.0, **{cx.NC[p]: 16.0 for p in cx.PK}, cx.NC["OPEN"]: 16.0}, print_cols=cx.LASTN)
    head(cx, ws, "ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS", cx.LASTN)
    w(ws, "B8", "=Capitalisation!B8+1", fB, al=A_L)
    cx.noteno_counter += 1
    cx.noteno["ADD"] = cx.noteno_counter
    w(ws, "C8", "Statement of additional disclosures with respect to amendments to Schedule III to the Companies Act, 2013 and the SEBI ICDR Regulations", fB)
    S = cx.SCALE
    r = 10
    T = I.get("texts") or {}
    last = cx.LASTN
    width = 55 + 16 * (cx.N + 1)

    def para(t):
        nonlocal r
        merge(ws, f"C{r}:{last}{r}", t, fN, A_J, est_height(t, width))
        r += 1

    def thead(label="Particulars", kind="pl"):
        nonlocal r
        w(ws, f"C{r}", label, fB, bd=B_TB)
        for p in cx.PK:
            w(ws, f"{cx.NC[p]}{r}", cx.dtxt(p, kind), fB, bd=B_TB, al=A_CW)
        ws.row_dimensions[r].height = 30
        r += 1
    for key, dflt in ADD_DEFAULT:
        t = dflt if dflt.startswith("TABLE_") else cx.text("add" + key, dflt, f"Additional disclosures {key}")
        if t == "TABLE_LOANS":
            para(cx.text("add(iii)_text", "(iii) Loans and advances in the nature of loans granted to promoters, directors, KMPs and related parties (repayable on "
                                        "demand or without specifying any terms or period of repayment):"))
            w(ws, f"C{r}", "Type of borrower", fB, bd=B_ALL)
            for p in cx.PK:
                w(ws, f"{cx.NC[p]}{r}", cx.dtxt(p), fB, bd=B_ALL, al=A_CW)
            ws.row_dimensions[r].height = 30
            r += 1
            LN = I.get("loans_to_related") or {}
            for i, lbl in enumerate(["Promoters", "Directors", "KMPs", "Related parties"]):
                bd = B_BLR if i == 3 else B_LR
                w(ws, f"C{r}", lbl, fN, bd=bd)
                for p in cx.PK:
                    v = (LN.get(lbl) or {}).get(p)
                    w(ws, f"{cx.NC[p]}{r}", rs_f(v, S) if _n(v) else "Nil", fN, ACC, bd=bd)
                r += 1
            r += 1
        elif t == "TABLE_AUDITOR":
            para("(xiii) Payments to the auditor:")
            thead()
            TA = (I.get("auditor_fees") or {}).get("tax_audit") or {}
            other = (I.get("auditor_fees") or {}).get("other") or {}
            f0 = r
            aud = "pl_aud_sub" in cx.refs
            nrow = 1 + int(any(_n(TA.get(p)) for p in cx.PK)) + int(any(_n(other.get(p)) for p in cx.PK))
            w(ws, f"C{r}", "As auditor - statutory audit")
            for p in cx.PK:
                sub = "".join(f"-{cx.NC[p]}{r + 1 + j}" for j in range(nrow - 1))
                w(ws, f"{cx.NC[p]}{r}", f"='Notes_P&L'!{cx.NC[p]}{{{{pl_aud_sub}}}}{sub}" if aud else 0, fN, ACC)
            r += 1
            for lbl, src in (("For taxation matters (tax audit)", TA), ("For other services", other)):
                if not any(_n(src.get(p)) for p in cx.PK):
                    continue
                w(ws, f"C{r}", lbl)
                for p in cx.PK:
                    w(ws, f"{cx.NC[p]}{r}", rs_f(src.get(p), S) or 0, fN, ACC)
                r += 1
            w(ws, f"C{r}", f"Total (refer Note {cx.noteno.get('OE')})", fB, bd=B_TB)
            for p in cx.PK:
                w(ws, f"{cx.NC[p]}{r}", f"=SUM({cx.NC[p]}{f0}:{cx.NC[p]}{r - 1})", fB, ACC, bd=B_TB)
            r += 2
        elif t == "TABLE_FOREIGN":
            para("(xv) Value of imports, earnings and expenditure in foreign currency:")
            thead()
            FX = I.get("foreign") or {}
            for i, (lbl, k) in enumerate((("A. CIF value of imports (stock-in-trade)", "imports"), ("B. Earnings in foreign exchange (FOB value of exports)", "exports"),
                                          (FX.get("expenditure_label") or "C. Expenditure in foreign currency", "expenditure"))):
                w(ws, f"C{r}", lbl, fN, bd=B_B if i == 2 else None)
                for p in cx.PK:
                    v = (FX.get(k) or {}).get(p)
                    w(ws, f"{cx.NC[p]}{r}", rs_f(v, S) if _n(v) is not None else (v if isinstance(v, str) else 0), fN, ACC, bd=B_B if i == 2 else None)
                r += 1
            fn = cx.text("add(xv)_note", "", "Additional disclosures (xv) - note under the table")
            if fn:
                para(fn)
            r += 1
        else:
            para(t)
    cx.finish_print(ws, r)
